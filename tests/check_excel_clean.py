#!/usr/bin/env python3
"""E2: 정리 엔진(clean_data.py)이 실습1 원본에서 정답표와 같은 결과를 내고 원본을 바꾸지 않는지 확인한다.

정답표 행 번호는 '번호' 열 값이고, 엑셀 행 번호(src_row)는 번호+1이다(머리글이 1행).
완전 중복으로 지운 행(번호 31·32)의 정규화 기대값은 결과에 없으므로 빼고 비교한다.
양성 대조: 같은 비교 함수를 정리 전 원본에 적용하면 불일치가 나와야 한다.
"""
from __future__ import annotations

import hashlib
import json
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _excel_fixtures import KEY, P1, run_script  # noqa: E402

K = KEY["실습1"]
fails: list[str] = []


def check(cond: bool, msg: str):
    if not cond:
        fails.append(msg)


def rows_by_no(path: Path, sheet: str) -> dict[int, dict]:
    ws = load_workbook(path, data_only=True)[sheet]
    head = [c.value for c in ws[1]]
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        out[int(r[0])] = dict(zip(head, r))
    return out


def mismatches(rows: dict[int, dict], removed_nos: set[int]) -> list[str]:
    bad = []
    for n, v in K["phone_normalize"].items():
        if int(n) in removed_nos:
            continue
        got = rows.get(int(n), {}).get("연락처")
        if got != v:
            bad.append(f"번호 {n} 연락처 {got} ≠ {v}")
    for n, v in K["date_normalize"].items():
        if int(n) in removed_nos:
            continue
        got = rows.get(int(n), {}).get("가입일")
        if got != v:
            bad.append(f"번호 {n} 가입일 {got} ≠ {v}")
    return bad


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    before = hashlib.sha256(P1.read_bytes()).hexdigest()
    orig = rows_by_no(P1, "고객목록")
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as td:
        out = Path(td) / "정리결과.xlsx"
        r = run_script("clean_data.py", P1, "--out", out, "--near-key", "이름,연락처")
        check(r.returncode == 0, f"clean_data 실패: {r.stdout[-400:]} {r.stderr[-400:]}")
        if r.returncode != 0:
            print("\n".join(fails))
            return 1
        rep = json.loads(Path(str(out) + ".report.json").read_text(encoding="utf-8"))
        rows = rows_by_no(out, "정리완료")
        removed_nos = {x["src_row"] - 1 for x in rep["removed"]}

        check(hashlib.sha256(P1.read_bytes()).hexdigest() == before, "원본 파일이 바뀌었다")
        check(rep["rows_in"] == K["rows_total"] and rep["rows_out"] == K["rows_after_exact_dedup"] == len(rows),
              f"행 수 입력 {rep['rows_in']} 결과 {rep['rows_out']}/{len(rows)} 정답 {K['rows_after_exact_dedup']}")
        check(sorted(removed_nos) == K["exact_duplicate_rows"], f"삭제 행 {sorted(removed_nos)} / 정답 {K['exact_duplicate_rows']}")
        dup_of = {str(x["src_row"] - 1): x["duplicate_of"] - 1 for x in rep["removed"]}
        check(dup_of == K["exact_duplicate_of"], f"중복 원본 {dup_of} / 정답 {K['exact_duplicate_of']}")
        bad = mismatches(rows, removed_nos)
        check(not bad, f"정규화 불일치 {len(bad)}건: {bad[:5]}")
        exp_phone = sum(1 for n in K["phone_normalize"] if int(n) not in removed_nos)
        exp_date = sum(1 for n in K["date_normalize"] if int(n) not in removed_nos)
        got_phone = sum(1 for c in rep["changes"] if c["column"] == "연락처")
        got_date = sum(1 for c in rep["changes"] if c["column"] == "가입일")
        check((got_phone, got_date) == (exp_phone, exp_date), f"변경 기록 수 전화 {got_phone}/{exp_phone}, 날짜 {got_date}/{exp_date}")

        cands = {(c["src_row"], c["column"]): c for c in rep["candidates"]}
        for n in K["invalid_phone_rows"]:
            check((n + 1, "연락처") in cands, f"번호 {n} 전화번호 오류가 후보에 없다")
            check(rows[n]["연락처"] == orig[n]["연락처"], f"번호 {n} 잘못된 전화번호를 고쳤다: {rows[n]['연락처']}")
        for n, dom in K["email_typo_rows"].items():
            c = cands.get((int(n) + 1, "이메일"))
            check(c is not None and c["suggestion"] == dom, f"번호 {n} 이메일 오타 후보·제안({dom})이 없다: {c}")
            check(rows[int(n)]["이메일"] == orig[int(n)]["이메일"], f"번호 {n} 이메일을 고쳤다")
        for n in K["outlier_rows"]:
            check((n + 1, "누적구매액") in cands, f"번호 {n} 극단값이 후보에 없다")
            check(rows[n]["누적구매액"] == orig[n]["누적구매액"], f"번호 {n} 극단값을 바꿨다")
        for n in K["ambiguous_date_rows"]:
            check((n + 1, "가입일") in cands, f"번호 {n} 모호한 날짜가 후보에 없다")
            check(rows[n]["가입일"] == orig[n]["가입일"], f"번호 {n} 모호한 날짜를 추측해 바꿨다: {rows[n]['가입일']}")
        blanks = {b["column"]: b["count"] for b in rep["blanks"]}
        for col, nos in K["blank_cells"].items():
            check(blanks.get(col) == len(nos), f"빈 칸 보고 {col}: {blanks.get(col)} / {len(nos)}")
            for n in nos:
                check(rows[n][col] in (None, ""), f"번호 {n} {col} 빈 칸을 채웠다: {rows[n][col]}")
        nd = load_workbook(out, data_only=True)["유사중복_후보"]
        near_rows = sorted(r[1] for r in nd.iter_rows(min_row=2, values_only=True) if isinstance(r[1], int))
        exp_near = sorted(n + 1 for g in K["near_duplicate_groups"] for n in g)
        check(near_rows == exp_near, f"유사 중복 후보 {near_rows} / 정답 {exp_near}")
        check(all(n in rows for g in K["near_duplicate_groups"] for n in g), "유사 중복 행을 지웠다(지우지 말고 보고해야 한다)")
        sheets = load_workbook(out).sheetnames
        check(sheets[:3] == ["정리완료", "변경리포트", "변경내역"], f"결과 시트 {sheets}")

        v = run_script("verify_excel.py", "clean", P1, out)
        check("VERIFY OK" in v.stdout, f"verify_excel clean 실패: {v.stdout[-500:]}")

        # 양성 대조: 정리 전 원본에는 같은 비교가 불일치를 내야 한다
        control = mismatches(orig, removed_nos)
        check(len(control) >= 40, f"양성 대조 실패: 원본에서도 불일치가 {len(control)}건뿐(비교가 실패할 수 없음)")
        # 같은 파일에 덮어쓰기 거부
        same = run_script("clean_data.py", P1, "--out", P1)
        check(same.returncode != 0 and hashlib.sha256(P1.read_bytes()).hexdigest() == before, "원본 덮어쓰기를 거부하지 않았다")

    for f in fails:
        print("FAIL", f)
    if fails:
        print(f"EXCEL CLEAN FAILED ({len(fails)})")
        return 1
    print(f"EXCEL CLEAN OK — 33→31행, 전화 {exp_phone}·날짜 {exp_date}건 정답 일치, 후보 {len(cands)}건 비수정, 양성 대조 {len(control)}건 불일치 확인")
    return 0


if __name__ == "__main__":
    sys.exit(main())
