#!/usr/bin/env python3
"""P3 오라클: verify_deck 양성·음성 대조.

음성(통과해야 함): 출처 숫자 + derived 선언 + 반올림 표기로 만든 덱
양성(실패해야 함):
  a. 출처에 없는 본문 숫자(지어낸 수치)
  b. 출처에 없는 차트 값
  c. 템플릿 안내 문구 잔존
  d. derived 선언을 빼면 계산값(23%)이 걸린다
  e. outline 슬라이드 수 불일치
"""
from __future__ import annotations

import copy
import json
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _deck_fixtures import SAMPLES, SCRIPTS, full_outline, make_sources, write  # noqa: E402


def build_and_verify(td: Path, name: str, outline: dict, src: Path, verify_outline: dict | None = None) -> dict:
    o = write(td / f"{name}.json", outline)
    out = td / name / "deck.pptx"
    b = subprocess.run([sys.executable, str(SCRIPTS / "build_deck.py"), str(o), "--out", str(out), "--sources", str(src)],
                       capture_output=True, text=True, encoding="utf-8")
    if b.returncode != 0:
        raise SystemExit(f"FAIL: {name} 빌드 실패 {b.stdout}")
    vo = write(td / f"{name}.verify.json", verify_outline) if verify_outline is not None else o
    r = subprocess.run([sys.executable, str(SCRIPTS / "verify_deck.py"), str(out), "--sources", str(src),
                        "--outline", str(vo), "--json"], capture_output=True, text=True, encoding="utf-8")
    res = json.loads(r.stdout)
    res["exit"] = r.returncode
    return res


def rules(res: dict) -> set[str]:
    return {e["rule"] for e in res["errors"]}


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    problems: list[str] = []
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        src = make_sources(td / "sources")
        clean = full_outline()
        clean["slides"][7]["bullets"].append("반품률 2.8% 유지(객단가 32,450원)")
        clean["slides"][4]["table"]["rows"].append(["합계", "4,520"])
        res = build_and_verify(td, "clean", clean, src)
        if res["exit"] != 0 or not res["ok"]:
            problems.append(f"깨끗한 덱이 실패: {res['errors']}")
        if res["numbers_checked"] < 8:
            problems.append(f"대조한 숫자가 너무 적음({res['numbers_checked']}) — 검사기가 숫자를 못 읽는 듯")

        a = copy.deepcopy(clean)
        a["slides"][2]["kpis"][1]["value"] = "38,900원"  # 출처에 없는 객단가
        r = build_and_verify(td, "fabricated_text", a, src)
        if "number-provenance" not in rules(r) or not any("38,900" in e["detail"] for e in r["errors"]):
            problems.append("a. 지어낸 본문 숫자를 못 잡음")

        b = copy.deepcopy(clean)
        b["slides"][3]["chart"]["series"][0]["values"] = [1800, 1520, 1350]
        r = build_and_verify(td, "fabricated_chart", b, src)
        if not any("1350" in e["detail"] and "차트" in e["detail"] for e in r["errors"]):
            problems.append("b. 지어낸 차트 값을 못 잡음")

        c = copy.deepcopy(clean)
        c["slides"][7]["bullets"].append("텍스트를 입력하십시오")
        r = build_and_verify(td, "leftover", c, src)
        if "leftover-text" not in rules(r):
            problems.append("c. 템플릿 안내 문구를 못 잡음")

        d = copy.deepcopy(clean)
        d["derived"] = []
        r = build_and_verify(td, "no_derived", d, src)
        if not any("'23'" in e["detail"] for e in r["errors"]):
            problems.append("d. 선언 없는 계산값 23%를 못 잡음")

        e = copy.deepcopy(clean)
        short = copy.deepcopy(clean)
        short["slides"] = short["slides"][:-1]
        r = build_and_verify(td, "count", e, src, verify_outline=short)
        if "slide-count" not in rules(r):
            problems.append("e. 슬라이드 수 불일치를 못 잡음")

        # f. 템플릿 제목 칸이 글자 크기를 상속(명시 없음)하는 긴 제목 → overflow로 잡아야 한다
        from pptx import Presentation
        sys.path.insert(0, str(SCRIPTS))
        import verify_deck
        prs = Presentation(str(SAMPLES / "samples/template_4x3.pptx"))
        sl = prs.slides.add_slide(prs.slide_layouts[5])
        sl.shapes.title.text = "주간 매출 4,520만 원, 쿠팡·스마트스토어·자사몰 3개 채널 중 쿠팡이 성장을 이끌었다는 점을 강조"
        prs.save(str(td / "inherited_title.pptx"))
        rf = verify_deck.verify(td / "inherited_title.pptx", src, None)
        if not any(x["rule"] == "overflow" for x in rf["errors"]):
            problems.append("f. 상속 글자 크기 제목의 넘침을 못 잡음")
    for p in problems:
        print("FAIL:", p)
    if problems:
        return 1
    print(f"clean deck passed ({res['numbers_checked']} numbers checked); controls a–f detected")
    print("VERIFY DECK CONTROL OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
