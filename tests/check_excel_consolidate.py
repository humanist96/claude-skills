#!/usr/bin/env python3
"""E3: 취합 엔진(consolidate.py)이 실습3 원본을 '원본 탭 보존 + 수식 기반 통합관리 시트'로 만들고,
재계산한 값이 원본 탭을 직접 계산한 값(정답표)과 같으며 수식 오류가 0인지 확인한다.

양성 대조(수식 연결이 살아 있는가): 결과 파일의 쿠팡 탭 판매량 한 칸을 바꾸고 다시 계산하면
통합관리 시트의 쿠팡 판매량 합계도 같은 만큼 바뀌어야 한다. 값을 붙여 넣었다면 바뀌지 않는다.

재계산 엔진(Excel 또는 LibreOffice)이 필요하다. 엔진이 없는 CI에서는 --allow-no-engine으로
수식 구조만 검사하고 'EXCEL CONSOLIDATE STRUCTURE OK'를 출력한다(값 일치는 증명하지 않음).
"""
from __future__ import annotations

import hashlib
import json
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _excel_fixtures import CONSOLIDATE_SPEC, KEY, P3, engine_available, run_script  # noqa: E402

K = KEY["실습3"]
fails: list[str] = []


def check(cond: bool, msg: str):
    if not cond:
        fails.append(msg)


def column_values(path: Path, sheet: str, name: str) -> list:
    ws = load_workbook(path, data_only=True)[sheet]
    head = [c.value for c in ws[1]]
    j = head.index(name)
    return [r[j] for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]


def main(argv: list[str]) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    has_engine = engine_available()
    if not has_engine and "--allow-no-engine" not in argv:
        print("EXCEL CONSOLIDATE FAILED — 재계산 엔진(Excel·LibreOffice)이 없다. CI에서는 --allow-no-engine")
        return 1
    before = hashlib.sha256(P3.read_bytes()).hexdigest()
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as td:
        spec = Path(td) / "spec.json"
        spec.write_text(json.dumps(CONSOLIDATE_SPEC, ensure_ascii=False), encoding="utf-8")
        out = Path(td) / "취합결과.xlsx"
        args = [P3, "--spec", spec, "--out", out] + ([] if has_engine else ["--no-recalc"])
        r = run_script("consolidate.py", *args)
        check(r.returncode == 0, f"consolidate 실패: {r.stdout[-400:]} {r.stderr[-400:]}")
        if r.returncode != 0:
            print("\n".join(fails))
            return 1
        check(hashlib.sha256(P3.read_bytes()).hexdigest() == before, "원본 파일이 바뀌었다")
        wb = load_workbook(out)
        check(wb.sheetnames == ["통합관리", "취합리포트"] + K["sheets"], f"시트 구성 {wb.sheetnames}")
        t = wb["통합관리"]
        keys = [t.cell(r, 1).value for r in range(2, t.max_row + 1) if t.cell(r, 1).value]
        check(len(keys) == len(set(keys)) == K["union_codes"], f"통합 키 {len(keys)}개(고유 {len(set(keys))}) / 정답 {K['union_codes']}")
        formulas = [c.value for row in t.iter_rows(min_row=2, min_col=2) for c in row if isinstance(c.value, str) and c.value.startswith("=")]
        check(len(formulas) == (t.max_row - 1) * (t.max_column - 1), f"수식이 아닌 칸이 있다: 수식 {len(formulas)}칸")
        check(all("XLOOKUP" not in f.upper() and "FILTER(" not in f.upper() for f in formulas), "호환되지 않는 함수 사용")
        check(any("INDEX(" in f and "MATCH(" in f and "IFERROR(" in f for f in formulas), "INDEX/MATCH 수식이 없다")
        check(any("'" in f or "!" in f for f in formulas), "시트 참조가 없다")

        if has_engine:
            for ch, exp in K["channels"].items():
                vals = column_values(out, "통합관리", f"{ch}_판매량")
                nums = [v for v in vals if isinstance(v, (int, float))]
                check(len(nums) == exp["products"], f"{ch} 등록 상품 {len(nums)} / 정답 {exp['products']}")
                check(sum(nums) == exp["qty_sum"], f"{ch} 판매량 합 {sum(nums)} / 정답 {exp['qty_sum']}")
            tot = sum(v for v in column_values(out, "통합관리", "전체_판매량") if isinstance(v, (int, float)))
            check(tot == sum(c["qty_sum"] for c in K["channels"].values()), f"전체 판매량 {tot}")
            reg = sum(v for v in column_values(out, "통합관리", "등록채널수") if isinstance(v, (int, float)))
            check(reg == sum(c["products"] for c in K["channels"].values()), f"등록채널수 합 {reg}")
            names = column_values(out, "통합관리", "상품명")
            check(all(isinstance(n, str) and n for n in names), "상품명이 빈 행이 있다(상품마스터 조회 실패)")
            v = run_script("verify_excel.py", "consolidate", P3, out)
            check("VERIFY OK" in v.stdout, f"verify_excel consolidate 실패: {v.stdout[-600:]}")
            mm = [ln for ln in v.stdout.splitlines() if "lookup-checked" in ln]
            check(bool(mm) and "불일치 0" in mm[0], f"INDEX/MATCH 대조 결과: {mm}")

            # 양성 대조: 원본 탭(결과 파일 안의 쿠팡 탭) 한 칸을 바꾸면 통합관리 값이 따라 바뀌어야 한다
            live = Path(td) / "연결확인.xlsx"
            wb2 = load_workbook(out)
            ws = wb2["쿠팡"]
            head = [c.value for c in ws[1]]
            qc = head.index("3월판매량") + 1
            old = ws.cell(2, qc).value
            ws.cell(2, qc, old + 1000)
            wb2.save(live)
            rr = run_script("recalc.py", live)
            check("RECALC OK" in rr.stdout, f"재계산 실패: {rr.stdout[-300:]}")
            new_sum = sum(v for v in column_values(live, "통합관리", "쿠팡_판매량") if isinstance(v, (int, float)))
            check(new_sum == K["channels"]["쿠팡"]["qty_sum"] + 1000, f"양성 대조 실패: 원본 탭을 바꿨는데 통합 합계가 {new_sum}(수식 연결 안 됨)")
        else:
            v = run_script("verify_excel.py", "consolidate", P3, out, "--allow-uncached")
            check("VERIFY OK" in v.stdout, f"verify_excel consolidate(구조) 실패: {v.stdout[-600:]}")

        same = run_script("consolidate.py", P3, "--spec", spec, "--out", P3)
        check(same.returncode != 0 and hashlib.sha256(P3.read_bytes()).hexdigest() == before, "원본 덮어쓰기를 거부하지 않았다")

    for f in fails:
        print("FAIL", f)
    if fails:
        print(f"EXCEL CONSOLIDATE FAILED ({len(fails)})")
        return 1
    if has_engine:
        print(f"EXCEL CONSOLIDATE OK — 원본 탭 7개 보존, 키 {K['union_codes']}개, 채널별 합계 정답 일치, 수식 오류 0, 원본 수정 시 통합 시트 연동 확인")
    else:
        print("EXCEL CONSOLIDATE STRUCTURE OK — 재계산 엔진 없음: 수식 구조만 확인(값 일치는 미증명)")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
