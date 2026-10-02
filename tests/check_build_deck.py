#!/usr/bin/env python3
"""P2 오라클: build_deck.py 렌더링 검증.

- 기본 디자인: 9종 슬라이드, 16:9, 네이티브 차트(값 일치)·표, 노트에 파일명 출처, 한글 East Asian 글꼴
- 회사 템플릿(4:3): 크기 10×7.5 유지, 템플릿 예시 슬라이드 제거, 제목이 템플릿 제목 칸에 들어감, 빈 칸 없음
- 잘못된 outline(알 수 없는 type, headline 누락)은 실패 코드
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _deck_fixtures import SAMPLES, SCRIPTS, full_outline, make_sources, write  # noqa: E402

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import PP_PLACEHOLDER  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402


def build(outline: Path, out: Path, *extra) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(SCRIPTS / "build_deck.py"), str(outline), "--out", str(out), *map(str, extra)],
                          capture_output=True, text=True, encoding="utf-8")


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    problems: list[str] = []
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        src = make_sources(td / "sources")
        logo = SAMPLES / "example/example_3_comany-to-ppt/7. 로고.png"
        o = write(td / "outline.json", full_outline(str(logo)))
        out = td / "deck.pptx"
        r = build(o, out, "--sources", src)
        if r.returncode != 0:
            print("FAIL: 기본 렌더링 실패", r.stdout, r.stderr)
            return 1
        prs = Presentation(str(out))
        if len(prs.slides) != 9:
            problems.append(f"슬라이드 수 {len(prs.slides)} != 9")
        if (round(prs.slide_width / 914400, 2), round(prs.slide_height / 914400, 2)) != (13.33, 7.5):
            problems.append("기본 크기가 16:9가 아님")
        charts = [sh for s in prs.slides for sh in s.shapes if getattr(sh, "has_chart", False) and sh.has_chart]
        if len(charts) != 1 or list(charts[0].chart.plots[0].series[0].values) != [1800.0, 1520.0, 1200.0]:
            problems.append("네이티브 차트 없음 또는 값 불일치")
        tables = [sh for s in prs.slides for sh in s.shapes if getattr(sh, "has_table", False) and sh.has_table]
        if len(tables) != 1 or len(tables[0].table.rows) != 4:
            problems.append("표 없음 또는 행 수 불일치")
        notes = prs.slides[2].notes_slide.notes_text_frame.text if prs.slides[2].has_notes_slide else ""
        if "주간매출.csv" not in notes:
            problems.append("노트 출처가 파일명으로 표시되지 않음")
        pics = [sh for sh in prs.slides[0].shapes if sh.shape_type == 13]
        if not pics:
            problems.append("표지에 로고 없음")
        ea = prs.slides[2].shapes[0]._element.xpath(".//a:ea/@typeface")
        if not ea or ea[0] != "맑은 고딕":
            problems.append(f"East Asian 글꼴 미설정: {ea[:1]}")
        # 한글 run은 lang=ko-KR이어야 단어 단위로 줄바꿈된다(없으면 '반품/과'처럼 글자 단위로 잘림)
        bad_lang = [r.text for s in prs.slides for sh in s.shapes if sh.has_text_frame
                    for p in sh.text_frame.paragraphs for r in p.runs
                    if any("가" <= ch <= "힣" for ch in r.text) and r._r.get_or_add_rPr().get("lang") != "ko-KR"]
        if bad_lang:
            problems.append(f"한글 run에 ko-KR 언어 태그 없음 {len(bad_lang)}개: {bad_lang[:2]}")
        report = json.loads((td / "build_report.json").read_text(encoding="utf-8"))
        if report["slides"] != 9:
            problems.append("build_report 슬라이드 수 불일치")

        # 회사 템플릿 4:3 — 예시 슬라이드가 든 사본을 만든다
        tpl_src = SAMPLES / "samples/template_4x3.pptx"
        t = Presentation(str(tpl_src))
        s = t.slides.add_slide(t.slide_layouts[1])
        s.shapes.title.text = "SAMPLE 예시 슬라이드"
        tpl = td / "company.pptx"
        t.save(str(tpl))
        out2 = td / "deck_tpl.pptx"
        r = build(o, out2, "--template", tpl, "--sources", src)
        if r.returncode != 0:
            problems.append(f"템플릿 렌더링 실패: {r.stdout[-300:]}")
        else:
            p2 = Presentation(str(out2))
            if (round(p2.slide_width / 914400, 2), round(p2.slide_height / 914400, 2)) != (10.0, 7.5):
                problems.append("템플릿 4:3 크기를 따르지 않음")
            alltext = " ".join(sh.text_frame.text for sl in p2.slides for sh in sl.shapes if sh.has_text_frame)
            if "SAMPLE" in alltext:
                problems.append("템플릿 예시 슬라이드가 남아 있음")
            if len(p2.slides) != 9:
                problems.append(f"템플릿 덱 슬라이드 수 {len(p2.slides)}")
            kpi_slide = p2.slides[2]
            titles = [ph.text_frame.text for ph in kpi_slide.placeholders
                      if ph.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)]
            if titles != ["매출이 전주 대비 23% 늘었다"]:
                problems.append(f"헤드라인이 템플릿 제목 칸에 없음: {titles}")
            for n, sl in enumerate(p2.slides, 1):
                for ph in sl.placeholders:
                    if ph.has_text_frame and not ph.text_frame.text.strip() and ph.placeholder_format.type not in (
                            PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER):
                        problems.append(f"템플릿 덱 {n}번 슬라이드에 빈 칸 {ph.name}")
        # 잘못된 outline
        bad = full_outline()
        bad["slides"].append({"type": "mystery", "headline": "x"})
        if build(write(td / "bad1.json", bad), td / "bad1.pptx").returncode == 0:
            problems.append("알 수 없는 type이 통과함")
        bad2 = full_outline()
        bad2["slides"][3].pop("headline")
        if build(write(td / "bad2.json", bad2), td / "bad2.pptx").returncode == 0:
            problems.append("headline 누락이 통과함")
    for p in problems:
        print("FAIL:", p)
    if problems:
        return 1
    print("default 16:9 9 slides, native chart+table, notes sources, ea font; template 4:3 kept, sample slide removed, title placeholder used")
    print("BUILD DECK OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
