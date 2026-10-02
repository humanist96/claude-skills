#!/usr/bin/env python3
"""E4: verify_excel.py가 각 결함을 실제로 잡아내는지 양성 대조로 확인한다.

먼저 세 기능의 정상 결과가 VERIFY OK인지 본다(검사가 항상 실패하는 것은 아닌지).
그다음 정상 결과를 하나씩 망가뜨린 사본마다 해당 검사 이름이 오류로 나와야 한다.
"""
from __future__ import annotations

import json
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _excel_fixtures import CONSOLIDATE_SPEC, P1, P2, P3, engine_available, run_script  # noqa: E402

results: list[tuple[str, bool, str]] = []


def verify(mode: str, src: Path, out: Path, *extra) -> dict:
    r = run_script("verify_excel.py", mode, src, out, "--json", *extra)
    try:
        body = r.stdout[: r.stdout.rindex("VERIFY ")]
        return json.loads(body)
    except ValueError:
        return {"errors": [{"check": "verifier-crashed", "detail": r.stdout[-300:] + r.stderr[-300:]}], "warnings": []}


def expect(name: str, rep: dict, check: str):
    got = {e["check"] for e in rep["errors"]}
    results.append((name, check in got, f"기대 {check}, 실제 {sorted(got)}"))


def expect_ok(name: str, rep: dict):
    results.append((name, not rep["errors"], f"오류 {[e['check'] + ':' + e.get('detail', '')[:60] for e in rep['errors']]}"))


def copy(src: Path, dst: Path) -> Path:
    shutil.copy2(src, dst)
    rj = Path(str(src) + ".report.json")
    if rj.is_file():
        shutil.copy2(rj, Path(str(dst) + ".report.json"))
    return dst


def edit_cached_value(path: Path, sheet: str, cell: str, new: str):
    """수식은 그대로 두고 저장된 계산값만 바꾼다(엔진이 틀린 값을 냈거나 누가 값을 덮은 경우 흉내)."""
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    wbx = files["xl/workbook.xml"].decode("utf-8")
    rid = re.search(r'<sheet [^>]*name="' + re.escape(sheet) + r'"[^>]*r:id="([^"]+)"', wbx).group(1)
    rels = files["xl/_rels/workbook.xml.rels"].decode("utf-8")
    target = re.search(r'Id="' + rid + r'"[^>]*Target="([^"]+)"|Target="([^"]+)"[^>]*Id="' + rid + '"', rels)
    tpath = "xl/" + (target.group(1) or target.group(2)).lstrip("/").replace("xl/", "")
    xml = files[tpath].decode("utf-8")
    pat = re.compile(r'(<c r="' + cell + r'"[^>]*>.*?<f>[^<]*</f>\s*<v>)[^<]*(</v>)', re.S)
    assert pat.search(xml), f"{cell} 수식 셀을 찾지 못함"
    files[tpath] = pat.sub(lambda m: m.group(1) + new + m.group(2), xml, count=1).encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in files.items():
            z.writestr(n, b)


def main(argv: list[str]) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    has_engine = engine_available()
    if not has_engine and "--allow-no-engine" not in argv:
        print("VERIFY EXCEL CONTROL FAILED — 재계산 엔진이 없다. CI에서는 --allow-no-engine")
        return 1
    unc = [] if has_engine else ["--allow-uncached"]
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as td:
        T = Path(td)
        # ── 정상 결과 3종 ──
        src1 = copy(P1, T / "고객_원본.xlsx")
        clean = T / "정리.xlsx"
        run_script("clean_data.py", src1, "--out", clean, "--near-key", "이름,연락처")
        expect_ok("정리 결과 정상", verify("clean", src1, clean))

        spec = T / "spec.json"
        spec.write_text(json.dumps(CONSOLIDATE_SPEC, ensure_ascii=False), encoding="utf-8")
        cons = T / "취합.xlsx"
        run_script("consolidate.py", P3, "--spec", spec, "--out", cons, *([] if has_engine else ["--no-recalc"]))
        expect_ok("취합 결과 정상", verify("consolidate", P3, cons, *unc))

        snap = T / "snap.json"
        run_script("verify_excel.py", "snapshot", P2, "--out", snap)
        ana = T / "분석.xlsx"
        run_script("analysis_tools.py", "report", P2, "--out", ana, "--date-col", "주문일", "--value-col", "결제금액",
                   "--category-col", "카테고리", "--item-col", "상품명")
        expect_ok("분석 결과 정상", verify("analysis", P2, ana, "--snapshot", snap))

        # ── 정리(clean) 대조 ──
        wb = load_workbook(src1)
        wb.active.cell(2, 9, "작업 중 메모")  # 원본 비고 칸을 바꿈
        wb.save(src1)
        expect("원본 변경", verify("clean", src1, clean), "input-changed")
        shutil.copy2(P1, src1)

        def clean_variant(name, fn, check):
            p = copy(clean, T / f"정리_{name}.xlsx")
            w = load_workbook(p)
            fn(w["정리완료"])
            w.save(p)
            expect(name, verify("clean", src1, p), check)

        def row_of(ws, no):
            return next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == no)

        clean_variant("행 수 불일치", lambda ws: ws.delete_rows(ws.max_row, 1), "row-count")
        clean_variant("후보 셀 수정", lambda ws: ws.cell(row_of(ws, 5), 3, "010-7777-0888"), "candidate-modified")
        clean_variant("기록 없는 변경", lambda ws: ws.cell(row_of(ws, 2), 2, "다른이름"), "unlogged-change")
        clean_variant("빈 칸 채움", lambda ws: ws.cell(row_of(ws, 9), 8, "서울 강남구"), "blank-filled")
        p = copy(clean, T / "정리_중복오판.xlsx")
        rj = Path(str(p) + ".report.json")
        rep = json.loads(rj.read_text(encoding="utf-8"))
        rep["removed"][0]["duplicate_of"] = 2
        rj.write_text(json.dumps(rep, ensure_ascii=False), encoding="utf-8")
        expect("잘못된 중복 삭제", verify("clean", src1, p), "bad-dedup")
        p = T / "정리_보고서없음.xlsx"
        shutil.copy2(clean, p)
        expect("변경 기록 없음", verify("clean", src1, p), "missing-report")

        # ── 취합(consolidate) 대조 ──
        def cons_variant(name, fn, check, recalc=False):
            p = copy(cons, T / f"취합_{name}.xlsx")
            w = load_workbook(p)
            fn(w)
            w.save(p)
            if recalc and has_engine:
                run_script("recalc.py", p)
            expect(name, verify("consolidate", P3, p, *unc), check)

        cons_variant("원본 시트 누락", lambda w: w.remove(w["쿠팡"]), "original-sheet-missing")
        cons_variant("원본 시트 변경", lambda w: w["자사몰"].cell(2, 1, "변경됨"), "original-sheet-changed")
        cons_variant("통합 키 누락", lambda w: w["통합관리"].delete_rows(5, 1), "missing-keys")
        cons_variant("값 붙여넣기", lambda w: w["통합관리"].cell(3, 2, "그냥 값"), "values-pasted")
        cons_variant("호환 안 되는 함수", lambda w: w["통합관리"].cell(2, w["통합관리"].max_column + 1, "=XLOOKUP(A2,쿠팡!A:A,쿠팡!B:B)"),
                     "forbidden-function")
        if has_engine:
            cons_variant("수식 오류", lambda w: w["통합관리"].cell(2, w["통합관리"].max_column + 1, "=1/0"), "formula-error", recalc=True)
            p = copy(cons, T / "취합_계산값불일치.xlsx")
            edit_cached_value(p, "통합관리", "B2", "엉뚱한 상품")
            expect("조회값 불일치", verify("consolidate", P3, p), "lookup-mismatch")
            p = copy(cons, T / "취합_미계산.xlsx")
            w = load_workbook(p)
            w.save(p)  # openpyxl로 다시 저장하면 계산값이 사라진다
            expect("재계산 안 함", verify("consolidate", P3, p), "not-recalculated")
        else:
            cons_variant("수식 오류", lambda w: w["통합관리"].cell(2, w["통합관리"].max_column + 1, "#DIV/0!"), "formula-error")

        # ── 분석(analysis) 대조 ──
        p = T / "분석_출처없는숫자.xlsx"
        shutil.copy2(ana, p)
        w = load_workbook(p)
        ws = w["인사이트"]
        ws.cell(2, 2, ws.cell(2, 2).value.replace("88,097,750", "91,250,000"))
        w.save(p)
        expect("출처 없는 인사이트 숫자", verify("analysis", P2, p, "--snapshot", snap), "insight-number-provenance")
        p = T / "분석_만단위.xlsx"
        shutil.copy2(ana, p)
        w = load_workbook(p)
        ws = w["인사이트"]
        ws.cell(2, 2, ws.cell(2, 2).value.replace("88,097,750원", "약 8,810만 원"))
        w.save(p)
        expect_ok("만 단위 반올림 표기는 통과", verify("analysis", P2, p, "--snapshot", snap))
        p = T / "분석_차트없음.xlsx"
        shutil.copy2(ana, p)
        w = load_workbook(p)
        for s in w.worksheets:
            s._charts = []
        w.save(p)
        expect("엑셀 차트 없음", verify("analysis", P2, p, "--snapshot", snap), "no-native-chart")
        p = T / "분석_인사이트없음.xlsx"
        shutil.copy2(ana, p)
        w = load_workbook(p)
        w.remove(w["인사이트"])
        w.save(p)
        expect("인사이트 없음", verify("analysis", P2, p, "--snapshot", snap), "no-insights")
        tampered = T / "매출_원본.xlsx"
        shutil.copy2(P2, tampered)
        w = load_workbook(tampered)
        w.active.cell(2, 8, 1)
        w.save(tampered)
        expect("분석 중 원본 변경", verify("analysis", tampered, ana, "--snapshot", snap), "input-changed")
        expect("원본에 덮어쓰기", verify("clean", P1, P1), "input-changed")

    bad = [r for r in results if not r[1]]
    for name, ok, why in results:
        print(f"{'PASS' if ok else 'FAIL'}  {name}" + ("" if ok else f" — {why}"))
    if bad:
        print(f"VERIFY EXCEL CONTROL FAILED ({len(bad)}/{len(results)})")
        return 1
    print(f"VERIFY EXCEL CONTROL OK — 정상 3종 통과, 결함 대조 {len(results) - 4}종 모두 검출" + ("" if has_engine else " (엔진 없음: 계산값 대조 생략)"))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
