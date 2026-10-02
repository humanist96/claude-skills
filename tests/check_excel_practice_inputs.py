#!/usr/bin/env python3
"""E1: 엑셀 실습 원본 3종이 재현 가능하고, 정답표가 파일에서 독립적으로 다시 계산한 값과 같으며,
v1.6.1 결과물은 완성 예시로 재분류되었는지 확인한다.

정답표 숫자를 그대로 믿지 않고 원본 파일에서 직접 다시 센다.
양성 대조: 값을 하나 바꾼 사본에서는 같은 계산이 정답표와 달라져야 한다(검사가 실제로 실패할 수 있는지).
"""
from __future__ import annotations

import fnmatch
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _excel_fixtures import INPUTS, KEY, P1, P2, P3, ROOT  # noqa: E402

CATALOG = ROOT / "plugins/kevin-skills-practice/skills/practice-samples/samples/catalog.json"
fails: list[str] = []


def check(cond: bool, msg: str):
    if not cond:
        fails.append(msg)


def customers_facts(p: Path) -> dict:
    df = pd.read_excel(p, dtype=object)
    cmp = df.drop(columns=["번호"])
    dup = cmp.duplicated(keep="first")
    return {"rows": len(df), "dups": sorted(int(x) for x in df.loc[dup, "번호"]), "outlier": int(df["누적구매액"].max()),
            "blank_region": sorted(int(x) for x in df.loc[df["지역"].isna(), "번호"]),
            "blank_grade": sorted(int(x) for x in df.loc[df["등급"].isna(), "번호"])}


def sales_facts(p: Path) -> dict:
    df = pd.read_excel(p)
    df["월"] = pd.to_datetime(df["주문일"]).dt.strftime("%Y-%m")
    m = df.groupby("월")["결제금액"].sum()
    prod = df.groupby("상품명")["결제금액"].sum().sort_values(ascending=False)
    return {"rows": len(df), "total": int(df["결제금액"].sum()), "monthly": {k: int(v) for k, v in m.items()},
            "category": {k: int(v) for k, v in df.groupby("카테고리")["결제금액"].sum().items()},
            "top": prod.index[0], "top_amount": int(prod.iloc[0]), "best": m.idxmax(), "worst": m.idxmin()}


def multichannel_facts(p: Path) -> dict:
    wb = load_workbook(p, data_only=True)
    out = {"sheets": wb.sheetnames, "channels": {}}
    union = set()
    for name in wb.sheetnames:
        df = pd.read_excel(p, sheet_name=name)
        if "상품코드" in df.columns:
            union |= set(df["상품코드"].dropna())
        if "3월판매량" in df.columns:
            out["channels"][name] = {"products": int(df["상품코드"].notna().sum()), "qty_sum": int(df["3월판매량"].sum())}
    out["union"] = len(union)
    return out


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    r = subprocess.run([sys.executable, str(ROOT / "tools/dev/make_excel_practice_inputs.py"), "--check"], capture_output=True,
                       text=True, encoding="utf-8", errors="replace")
    check("REPRODUCIBLE" in r.stdout, f"재현 불가: {r.stdout[-300:]} {r.stderr[-300:]}")

    k1, k2, k3 = KEY["실습1"], KEY["실습2"], KEY["실습3"]
    f1 = customers_facts(P1)
    check(f1["rows"] == k1["rows_total"] == 33, f"실습1 행 수 {f1['rows']} / 정답표 {k1['rows_total']}")
    check(f1["dups"] == k1["exact_duplicate_rows"], f"실습1 완전 중복 {f1['dups']} / 정답표 {k1['exact_duplicate_rows']}")
    check(f1["outlier"] == 98_000_000, f"실습1 극단값 {f1['outlier']}")
    check(f1["blank_region"] == k1["blank_cells"]["지역"] and f1["blank_grade"] == k1["blank_cells"]["등급"], f"실습1 빈 칸 {f1}")
    wb1 = load_workbook(P1)
    check(wb1.sheetnames == ["고객목록"], f"실습1은 원본 시트 하나여야 한다: {wb1.sheetnames}")

    f2 = sales_facts(P2)
    check(f2["total"] == k2["total_amount"], f"실습2 합계 {f2['total']} / 정답표 {k2['total_amount']}")
    check(f2["monthly"] == k2["monthly_amount"], "실습2 월별 합계가 정답표와 다르다")
    check(f2["category"] == k2["category_amount"], "실습2 카테고리 합계가 정답표와 다르다")
    check((f2["top"], f2["top_amount"]) == (k2["top_product"], k2["top_product_amount"]), f"실습2 1위 상품 {f2['top']}")
    check((f2["best"], f2["worst"]) == (k2["best_month"], k2["worst_month"]), f"실습2 최고·최저 월 {f2['best']}, {f2['worst']}")
    check(len(load_workbook(P2).sheetnames) == 1, "실습2는 원본 시트 하나여야 한다(분석 결과 시트 없음)")

    f3 = multichannel_facts(P3)
    check(f3["sheets"] == k3["sheets"] and "통합관리" not in f3["sheets"], f"실습3 시트 {f3['sheets']}")
    check(f3["union"] == k3["union_codes"] == 108, f"실습3 상품코드 합집합 {f3['union']}")
    check(f3["channels"] == k3["channels"], f"실습3 채널별 값 {f3['channels']} / 정답표 {k3['channels']}")

    # 양성 대조: 매출 한 칸을 바꾸면 합계 검사가 정답표와 달라져야 한다
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as td:
        bad = Path(td) / "tampered.xlsx"
        shutil.copy2(P2, bad)
        wb = load_workbook(bad)
        ws = wb.active
        col = [c.value for c in ws[1]].index("결제금액") + 1
        ws.cell(2, col, ws.cell(2, col).value + 1000)
        wb.save(bad)
        check(sales_facts(bad)["total"] != k2["total_amount"], "양성 대조 실패: 바꾼 사본의 합계가 정답표와 같다(검사가 실패할 수 없음)")

    # 카탈로그: 실습 입력은 inputs/의 원본, v1.6.1 결과물은 완성 예시
    cat = json.loads(CATALOG.read_text(encoding="utf-8"))
    sets = {s["id"]: s for s in cat["sets"]}
    ip = sets.get("excel-practice", {})
    check(ip.get("kind") == "input" and all(f.startswith("excel-automation/inputs/") and f.endswith("_원본.xlsx") for f in ip.get("files", []))
          and len(ip.get("files", [])) == 3, f"excel-practice 세트가 원본 3종을 가리키지 않는다: {ip.get('files')}")
    outs = [s for s in cat["sets"] if s["skill"] == "excel-automation" and s["kind"] == "reference-output"]
    old = ["excel-automation/excel-automation-outputs/" + n for n in
           ("실습1_고객관리데이터.xlsx", "실습2_쇼핑몰매출데이터.xlsx", "실습3_멀티채널판매관리.xlsx")]
    for o in old:
        check(any(fnmatch.fnmatch(o, pat) for s in outs for pat in s["files"]), f"v1.6.1 결과물이 완성 예시로 분류되지 않았다: {o}")
        check(not any(fnmatch.fnmatch(o, pat) for pat in ip.get("files", [])), f"v1.6.1 결과물이 아직 실습 입력에 있다: {o}")
    check(any("answer_key.json" in f for s in outs for f in s["files"]), "정답표가 카탈로그에 없다")
    check((INPUTS.parent / "prompt" / "2.md").stat().st_size > 200, "prompt/2.md가 비어 있다")

    for f in fails:
        print("FAIL", f)
    if fails:
        print(f"PRACTICE INPUTS FAILED ({len(fails)})")
        return 1
    print("PRACTICE INPUTS OK — 원본 3종 재현, 정답표 독립 재계산 일치, 양성 대조 통과, 카탈로그 재분류")
    return 0


if __name__ == "__main__":
    sys.exit(main())
