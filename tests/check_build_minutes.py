#!/usr/bin/env python3
"""M4: build_minutes.py가 용도 4종 × 형식 4종을 렌더링하고, 결정·할 일이 하나도 빠지지 않으며,
용도별 구성(보고용은 결론 먼저, 이메일은 인사·맺음, 노션은 체크박스)과 빈 섹션 표시, 영어 머리말, 워드 한글 글꼴을 지키는지 확인한다.
"""
from __future__ import annotations

import copy
import json
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _meeting_fixtures import GOOD_MINUTES, TEXT, run_script  # noqa: E402

PURPOSES = ["team", "report", "personal", "email"]
FORMATS = {"md": ".md", "docx": ".docx", "notion": ".md", "text": ".txt"}


def text_of(p: Path) -> str:
    if p.suffix == ".docx":
        from docx import Document
        d = Document(str(p))
        parts = [x.text for x in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                parts += [c.text for c in r.cells]
        return "\n".join(parts)
    return p.read_text(encoding="utf-8")


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    fails: list[str] = []
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as td:
        T = Path(td)
        run_script("prepare_transcript.py", TEXT, "--out-dir", T)
        mp = T / "minutes.json"
        mp.write_text(json.dumps(GOOD_MINUTES, ensure_ascii=False), encoding="utf-8")
        outs = []
        for pur in PURPOSES:
            for fmt, ext in FORMATS.items():
                out = T / f"{pur}_{fmt}{ext}"
                r = run_script("build_minutes.py", mp, "--out", out, "--format", fmt, "--purpose", pur, "--transcript", T / "transcript.json")
                if "BUILD DONE" not in r.stdout:
                    fails.append(f"{pur}/{fmt} 실패: {r.stdout[-200:]} {r.stderr[-300:]}")
                    continue
                outs.append(out)
                body = text_of(out)
                if "TODO" in body or "{" in body:
                    fails.append(f"{pur}/{fmt}: 자리표시가 남았다")
                if pur == "email" and ("안녕하세요" not in body or "회신" not in body or "T0008" in body):
                    fails.append(f"{pur}/{fmt}: 이메일 인사·맺음이 없거나 근거 번호가 노출됨")
                if pur == "report":
                    if body.find("결론") < 0 or body.find("결론") > body.find("결정 사항"):
                        fails.append(f"{pur}/{fmt}: 결론이 결정 사항보다 먼저 오지 않는다")
                if pur == "personal" and "전환율 2.8%" not in body:
                    fails.append(f"{pur}/{fmt}: 개인 기록에 안건별 상세 논의가 없다")
                if fmt == "notion" and pur in ("team", "report", "personal") and "- [ ] 물류팀" not in body:
                    fails.append(f"{pur}/{fmt}: 노션용 할 일 체크박스가 없다")
                if "리스크 및 이슈" not in body:
                    fails.append(f"{pur}/{fmt}: 역할별 추가 섹션이 없다")
                if fmt in ("md", "docx") and pur != "email" and "근거 발언" not in body:
                    fails.append(f"{pur}/{fmt}: 근거 대조표가 없다")
        r = run_script("verify_minutes.py", mp, "--transcript", T / "transcript.json", "--rendered", *outs)
        if "MINUTES VERIFY OK" not in r.stdout:
            fails.append(f"렌더링 16종 대조 실패: {r.stdout[-600:]}")
        # 워드 한글 글꼴
        from docx import Document
        from docx.oxml.ns import qn
        d = Document(str(T / "report_docx.docx"))
        ea = d.styles["Normal"].element.rPr.find(qn("w:rFonts")).get(qn("w:eastAsia"))
        if ea != "맑은 고딕":
            fails.append(f"워드 한글 글꼴 {ea}")
        if len(d.tables) < 2:
            fails.append("워드에 참석자·할 일 표가 없다")
        # 빈 섹션은 '없음'
        empty = copy.deepcopy(GOOD_MINUTES)
        empty["decisions"], empty["action_items"], empty["open_questions"] = [], [], []
        ep = T / "empty.json"
        ep.write_text(json.dumps(empty, ensure_ascii=False), encoding="utf-8")
        run_script("build_minutes.py", ep, "--out", T / "empty.md", "--purpose", "team")
        eb = (T / "empty.md").read_text(encoding="utf-8")
        if eb.count("- 없음") < 3 or "## 결정 사항" not in eb or "## 할 일" not in eb:
            fails.append("빈 결정·할 일·미결을 '없음'으로 표시하지 않았다")
        # 영어 회의
        en = copy.deepcopy(GOOD_MINUTES)
        en["meta"]["language"] = "en"
        enp = T / "en.json"
        enp.write_text(json.dumps(en, ensure_ascii=False), encoding="utf-8")
        run_script("build_minutes.py", enp, "--out", T / "en.md", "--purpose", "team")
        enb = (T / "en.md").read_text(encoding="utf-8")
        if "## Decisions" not in enb or "## Action items" not in enb or "결정 사항" in enb:
            fails.append("영어 회의록 머리말이 영어가 아니다")
        # 날짜를 모르는 회의의 이메일 제목에 '(미상)'·'(Unknown)'이 찍히지 않는다(평가 중 발견)
        nod = copy.deepcopy(en)
        nod["meta"]["date"] = "Unknown"
        ndp = T / "nodate.json"
        ndp.write_text(json.dumps(nod, ensure_ascii=False), encoding="utf-8")
        run_script("build_minutes.py", ndp, "--out", T / "nodate.txt", "--purpose", "email", "--format", "text")
        subj = (T / "nodate.txt").read_text(encoding="utf-8").splitlines()[0]
        if "Unknown" in subj or "미상" in subj:
            fails.append(f"날짜 미상인 이메일 제목: {subj}")
        # 양성 대조: 할 일 한 줄을 지운 문서는 대조에서 실패해야 한다
        cut = T / "cut.md"
        cut.write_text("\n".join(l for l in (T / "team_md.md").read_text(encoding="utf-8").splitlines() if "광고 소재" not in l), encoding="utf-8")
        r = run_script("verify_minutes.py", mp, "--transcript", T / "transcript.json", "--rendered", cut)
        if "rendered-missing" not in r.stdout:
            fails.append("양성 대조 실패: 할 일이 빠진 문서를 잡지 못했다")
    for f in fails:
        print("FAIL", f)
    if fails:
        print(f"BUILD MINUTES FAILED ({len(fails)})")
        return 1
    print(f"BUILD MINUTES OK — 용도 {len(PURPOSES)} × 형식 {len(FORMATS)} = {len(outs)}개 렌더링, 결정·할 일 누락 0, 빈 섹션·영어·워드 글꼴 확인, 양성 대조 통과")
    return 0


if __name__ == "__main__":
    sys.exit(main())
