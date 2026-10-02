#!/usr/bin/env python3
"""meeting-minutes 평가 실행 결과 채점 (skill-creator grading.json 형식).

객관 항목은 결과 회의록 텍스트(md·txt·docx, 표는 행 단위)를 정답표와 대조한다. 두 구성을 같은 기준으로 채점하려고
- 회의록 파일은 이름이 아니라 확장자로 고르고, 녹취 사본(transcript·녹취·stt·raw·json)은 제외한다
- 할 일은 '한 줄(표의 한 행) 안에 담당자·할 일 핵심어·기한'이 함께 있으면 찾은 것으로 본다(담당자는 위 3줄 머리말도 인정)
[판단] 항목은 judgments.json으로 채점자가 넣는다.

사용법: python tools/grade_meeting_runs.py <iteration 디렉터리> [--judgments judgments.json]
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

PROMO = re.compile(r"(이 스킬은|이 스킬로|회의록 외에도 활용|강의 노트 정리, 인터뷰|다음에 녹음 파일이나 텍스트를 주시면)")
EXCLUDE = re.compile(r"(transcript|녹취|stt|raw|whisper)", re.I)


def units_of(p: Path) -> list[str]:
    if p.suffix.lower() == ".docx":
        from docx import Document
        d = Document(str(p))
        out = [x.text for x in d.paragraphs if x.text.strip()]
        for t in d.tables:
            for r in t.rows:
                cells = []
                for c in r.cells:
                    if c.text not in cells:
                        cells.append(c.text)
                out.append(" | ".join(cells))
        return out
    return [l for l in p.read_text(encoding="utf-8", errors="replace").splitlines() if l.strip()]


def minutes_files(run: Path) -> list[Path]:
    return sorted(p for p in (run / "outputs").rglob("*") if p.is_file() and p.suffix.lower() in (".md", ".txt", ".docx")
                  and p.name != "final_response.md" and not EXCLUDE.search(p.name) and not p.name.startswith("~$"))


def final_text(run: Path) -> str:
    p = run / "outputs" / "final_response.md"
    return p.read_text(encoding="utf-8", errors="replace") if p.is_file() else ""


def has(text: str, *alts) -> bool:
    return any(a.lower() in text.lower() for a in alts)


DUE_ALIASES = {"월요일": ["(월)", "Monday"], "화요일": ["(화)"], "수요일": ["(수)", "Wednesday"], "목요일": ["(목)", "Thursday"],
               "금요일": ["(금)", "Friday"], "내일": ["tomorrow"], "오늘": ["today"]}


def expand(dues):
    out = list(dues)
    for d in dues:
        out += DUE_ALIASES.get(d, [])
    return out


def body(units: list[str]) -> list[str]:
    """오인식 '부재' 검사용 본문: 녹취 원문을 인용하는 근거 대조표 구간과, 오인식을 알리는 줄(→, '인식', '들렸')은 뺀다."""
    out, skip = [], False
    for u in units:
        s = u.strip()
        if s.startswith("#"):
            skip = any(w in s for w in ("근거", "Evidence", "STT", "음성 인식", "변환 메모"))
        if skip or "→" in s or "->" in s or has(s, "인식", "들렸", "들림", "misheard", "transcrib"):
            continue
        out.append(u)
    return out


def find_action(units: list[str], owners, tasks, dues) -> bool:
    dues = expand(dues)
    for i, u in enumerate(units):
        ctx = " ".join(units[max(0, i - 3):i + 1])
        if has(u, *tasks) and (not dues or has(u, *dues)) and (has(u, *owners) or has(ctx, *owners)):
            return True
    return False


def section(units: list[str], title_words) -> list[str]:
    """머리말(#, **…**, 단독 줄)이 title_words를 담은 구간의 줄"""
    out, on = [], False
    for u in units:
        s = u.strip()
        is_head = s.startswith("#") or re.fullmatch(r"\*\*[^*]{1,30}\*\*:?", s) or re.fullmatch(r"\[[^\]]{1,20}\]", s)
        if is_head:
            on = any(w in s for w in title_words)
            continue
        if on:
            out.append(s)
    return out


def grade(run: Path, meta: dict) -> list[dict]:
    A = meta["assertions"]
    name = meta["eval_name"]
    files = minutes_files(run)
    units = [u for f in files for u in units_of(f)]
    doc = "\n".join(units)
    fin = final_text(run)
    res: dict[str, tuple] = {t: (None, "판단 대기") for t in A if t.startswith("[판단]")}
    ev = f"파일 {[f.name for f in files]}"

    def put(i, ok, why):
        res[A[i]] = (bool(ok), why)

    promo_ok = not PROMO.search(doc)
    if name == "sprint-audio-team-md":
        put(0, any(f.suffix == ".md" for f in files), ev)
        names = ["김PM", "이개발", "박디자인", "정기획"]
        bdoc = "\n".join(body(units))
        bad = [w for w in ("김 피암", "피암", "정기회") if w in bdoc]
        put(1, all(has(doc, n, n.replace("PM", " PM")) for n in names) and not bad, f"이름 {[n for n in names if not has(doc, n, n.replace('PM', ' PM'))]} 없음, 오인식 {bad}")
        soc = [u for u in units if "소셜" in u]
        put(2, any(has(u, "다음 스프린트", "이월", "넘기", "넘김", "연기") for u in soc) and "유월" not in bdoc, f"소셜 줄 {soc[:2]}")
        acts = [(["이개발"], ["API", "백엔드", "백앤드"], ["금요일"]), (["이개발"], ["메일"], ["내일"]),
                (["박디자인", "박 디자인"], ["UI", "퍼블리싱", "로그인"], ["수요일"]), (["박디자인", "박 디자인"], ["회원"], ["목요일"]),
                (["정기획"], ["기획서"], ["내일"])]
        miss = [a[1][0] for a in acts if not find_action(units, *a)]
        put(3, not miss, f"못 찾은 할 일 {miss}")
        m4 = re.search(r"매일\s*(발송|서비스|옵션)", bdoc)
        put(4, not m4, "매일 오인식 없음(근거 인용·오인식 안내 줄 제외)" if not m4 else f"본문에 '{m4.group(0)}'")
        nx = [" ".join(units[i:i + 3]) for i, u in enumerate(units) if has(u, "다음 회의", "다음 미팅", "next meeting")]
        put(5, any("금요일" in u for u in nx) and not any("52시" in u for u in body(units)), f"다음 회의 줄 {[x[:60] for x in nx[:2]]}")
        put(6, promo_ok, "홍보 문구 없음" if promo_ok else "홍보 문구 있음")
    elif name == "planning-audio-report-docx":
        put(0, any(f.suffix == ".docx" for f in files), ev)
        # 가격(A[1])은 모델마다 인식이 달라 [판단] 항목으로 채점한다(정답표 note 참고)
        put(2, has(doc, "4월 28일", "4/28", "04-28", "4.28"), "베타 4월 28일")
        put(3, any("규칙" in u and has(u, "정식", "AI 모델") for u in units), "규칙 기반 → 정식 AI")
        acts = [(["김과장", "김 과장"], ["가격"], []), (["이대리", "이 대리"], ["결제", "로드맵"], []),
                (["박사원", "박 사원"], ["인플루언서", "인플로언서", "예산"], []), (["최부장", "최 부장", "부장"], ["이사회"], [])]
        miss = [a[1][0] for a in acts if not find_action(units, *a)]
        put(4, not miss, f"못 찾은 할 일 {miss}")
        b1 = has(doc, "3억")
        b2 = has(doc, "1억 5천", "1억5천", "1.5억", "1억 5,000", "1억5,000", "150,000,000")
        b3 = has(doc, "2,000만", "2000만", "2천만", "20,000,000")
        put(5, b1 and b2 and b3, f"3억 {b1}, 1억5천 {b2}, 2,000만 {b3}")
        put(6, has(doc, "3월 22일", "3/22", "03-22") and has(doc, "10시", "10:00"), "다음 회의 3월 22일 10시")
        put(7, "2024" in doc and has(doc, "3월 15일", "03-15", "3/15", "03.15", ".3.15"), "회의일 2024-03-15")
        put(8, promo_ok, "홍보 문구 없음" if promo_ok else "홍보 문구 있음")
    elif name == "english-audio-email":
        allu = units + [l for l in fin.splitlines() if l.strip()]
        alltext = doc + "\n" + fin
        greet = has(alltext, "Hi ", "Hi,", "Hello", "Dear", "안녕하세요")
        close = has(alltext, "Thanks", "Thank you", "Best", "Regards", "감사합니다", "드림")
        put(0, greet and close, f"인사 {greet}, 맺음 {close}")
        put(1, all(n in alltext for n in ("Sarah", "Mike", "Jenny")), "참석자 3명")
        m1 = find_action(allu, ["Mike", "마이크"], ["bug", "버그"], ["tomorrow", "내일"])
        m2 = find_action(allu, ["Mike", "마이크"], ["load", "부하"], ["Friday", "금요일"])
        put(2, m1 and m2, f"버그 수정 {m1}, 부하 테스트 {m2}")
        j1 = find_action(allu, ["Jenny", "제니"], ["Figma", "onboarding", "온보딩"], ["today", "오늘", "end of day", "EOD"])
        j2 = find_action(allu, ["Jenny", "제니"], ["dark", "다크"], ["Monday", "월요일"])
        put(3, j1 and j2, f"Figma {j1}, 다크 모드 추정 {j2}")
        dk = [u for u in allu if has(u, "dark", "다크")]
        put(4, any(has(u, "next week", "다음 주", "table", "보류", "미루", "미뤘", "postpone", "defer", "estimate first", "추정 후") for u in dk), f"다크 모드 줄 {dk[:2]}")
        put(5, "12%" in alltext and has(alltext, "10,000", "10000", "1만 명", "1만명"), "12%, 10,000")
        nx = [u for u in allu if has(u, "next meeting", "다음 회의", "next sync")]
        put(6, any(has(u, "Friday", "금요일") and "10" in u for u in nx), f"다음 회의 줄 {nx[:2]}")
        pm = PROMO.search(alltext)
        put(7, not pm, "홍보 문구 없음" if not pm else f"홍보 문구: '{pm.group(0)}'(파일 또는 최종 응답)")
    elif name == "clova-text-notion":
        put(0, bool(files), ev)
        dec_units = section(units, ("결정",)) or units
        c23 = [u for u in dec_units if re.search(r"(3월\s*)?23일", u)]
        ok23 = all(has(u, "취소", "변경", "연기", "미루", "미뤄", "번복", "대신", "→", "->", "에서", "밀려", "기존", "당초", "원래") for u in c23)
        put(1, "3월 30일" in doc and ok23, f"결정 구간의 23일 언급 {len(c23)}줄, 모두 변경 맥락 {ok23}")
        put(2, all(has(doc, a, a.replace(",", "")) for a in ("2,000만", "1,500만", "1,000만")), "예산 배분")
        put(3, any("챗봇" in u and has(u, "자주 묻는", "FAQ", "세 가지", "3가지") for u in units), "챗봇 결정")
        acts = [(["이준호", "준호"], ["재고", "물류"], ["수요일"]), (["최동훈", "동훈"], ["소재"], ["금요일"]),
                (["박서연", "서연"], ["챗봇 시나리오", "시나리오"], []), (["최동훈", "동훈"], ["인플루언서"], [])]  # 기한은 녹취에 명시되지 않음(다음 회의 전은 추론) — 담당·할 일만 본다
        miss = [a[1][0] for a in acts if not find_action(units, *a)]
        landing = [u for u in units if "검수" in u and has(u, "랜딩", "문구")]
        land_ok = any(has(u, "미정", "TBD", "담당자 없음", "정해지지") for u in landing)
        put(4, not miss and land_ok, f"못 찾은 할 일 {miss}, 랜딩 검수 담당 미정 {land_ok}")
        sc = [u for u in units if "시나리오" in u]
        put(5, any(has(u, "미정", "기한 없", "없음", "TBD", "정해지지", "추후") for u in sc), f"챗봇 시나리오 줄 {sc[:2]}")
        dec = section(units, ("결정",))
        bad_dec = [u for u in dec if has(u, "SNS 비중", "비중 확대", "비중을 늘", "인플루언서")]
        put(6, bool(dec) and not bad_dec, f"결정 구간 {len(dec)}줄, 잘못 들어간 것 {bad_dec[:2]}")
        put(7, sum(1 for u in units if u.strip().startswith("- [ ]")) >= 5, f"체크박스 {sum(1 for u in units if u.strip().startswith('- [ ]'))}개")
        att = [u for u in units if "참석" in u and "불참" not in u]
        put(8, not any("하늘" in u and not has(u, "불참", "휴가", "absent") for u in att), f"참석자 줄 {att[:2]}")
        put(9, not has(doc, "점심", "국숫집"), "잡담 없음")
        put(10, promo_ok, "홍보 문구 없음" if promo_ok else "홍보 문구 있음")
    (run / "_grade").mkdir(exist_ok=True)
    (run / "_grade" / "minutes_text.md").write_text("\n".join(units) + "\n\n## 최종 응답\n" + fin, encoding="utf-8")
    return [{"text": t, "passed": res.get(t, (False, "채점 규칙 없음"))[0], "evidence": res.get(t, (False, "채점 규칙 없음"))[1]} for t in A]


def main(argv: list[str]) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("iteration")
    ap.add_argument("--judgments")
    a = ap.parse_args(argv)
    it = Path(a.iteration)
    judg = json.loads(Path(a.judgments).read_text(encoding="utf-8")) if a.judgments else {}
    pending = 0
    for meta_file in sorted(it.glob("eval-*/eval_metadata.json")):
        meta = json.loads(meta_file.read_text(encoding="utf-8"))
        for cfg_dir in sorted(p for p in meta_file.parent.iterdir() if p.is_dir()):
            for run_dir in sorted(cfg_dir.glob("run-*")):
                key = f"{meta['eval_name']}/{cfg_dir.name}"
                exps = grade(run_dir, meta)
                for x in exps:
                    j = judg.get(key, {}).get(x["text"])
                    if j:
                        x["passed"], x["evidence"] = bool(j["passed"]), j["evidence"]
                passed = sum(1 for x in exps if x["passed"] is True)
                decided = sum(1 for x in exps if x["passed"] is not None)
                pending += len(exps) - decided
                (run_dir / "grading.json").write_text(json.dumps({"expectations": exps, "summary": {
                    "passed": passed, "failed": decided - passed, "total": len(exps),
                    "pass_rate": round(passed / len(exps), 4) if exps else 0}}, ensure_ascii=False, indent=2), encoding="utf-8")
                print(f"{key:<40} {passed}/{len(exps)}  (판단 대기 {len(exps) - decided})")
    print(f"PENDING JUDGMENTS {pending}" if pending else "GRADED ALL")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
