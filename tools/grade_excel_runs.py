#!/usr/bin/env python3
"""excel-automation 평가 실행 결과 채점 (skill-creator grading.json 형식).

객관 항목은 이 스크립트가 결과 파일을 직접 열어 정답표(answer_key.json)와 대조한다. [판단] 항목은
judgments.json으로 채점자가 판정을 넣는다. 두 구성(with_skill, old_skill)을 같은 기준으로 채점하려고
- 결과 파일은 시트 이름·열 위치에 기대지 않고 머리글 내용으로 찾는다
- 수식 값은 결과 파일 '사본'을 Excel(또는 LibreOffice)로 재계산해 읽는다(실행이 재계산했는지와 무관)

산출물(각 run 디렉터리)
  grading.json          expectations[{text, passed, evidence}], summary
  _grade/*.xlsx         재계산한 사본
  _grade/facts.md       판단 항목 채점용 요약(시트 목록, 인사이트·후보 문구)

사용법
  python tools/grade_excel_runs.py <iteration 디렉터리> [--judgments judgments.json]
judgments.json: {"<eval-name>/<config>": {"<기대 항목 원문>": {"passed": true, "evidence": "..."}}}
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import zipfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "plugins/kevin-claude-skills-book/skills/excel-automation/scripts"
INPUTS = ROOT / "plugins/kevin-claude-skills-practice/skills/practice-samples/samples/excel-automation/inputs"
KEY = json.loads((INPUTS / "answer_key.json").read_text(encoding="utf-8"))
sys.path.insert(0, str(SCRIPTS))
from verify_excel import ERRORS, FORBIDDEN, exempt, has_number, numbers_with_units  # noqa: E402

MARKER = re.compile(r"(미입력|미기재|없음|확인|빈\s?칸|누락|N/?A|^-$|null|공란)", re.I)
CHANNELS = {"스마트스토어": ["스마트스토어"], "쿠팡": ["쿠팡"], "카카오선물하기": ["카카오"], "자사몰": ["자사몰"]}


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def outputs(run: Path) -> list[Path]:
    return sorted(p for p in (run / "outputs").glob("*.xlsx") if not p.name.startswith("~$"))


def recalc_copy(src: Path, run: Path) -> Path:
    g = run / "_grade"
    g.mkdir(exist_ok=True)
    dst = g / src.name
    shutil.copy2(src, dst)
    subprocess.run([sys.executable, str(SCRIPTS / "recalc.py"), str(dst)], capture_output=True, text=True,
                   encoding="utf-8", errors="replace", timeout=600)
    return dst


def input_unchanged(run: Path, name: str) -> tuple[bool, str]:
    w = run / "work" / name
    if not w.is_file():
        return False, f"작업 폴더의 원본 {name}이 없다(이동·삭제)"
    same = sha(w) == sha(INPUTS / name)
    return same, "작업 폴더 원본 해시가 배포본과 같다" if same else "작업 폴더 원본이 바뀌었다"


def merged_value(ws, r: int, c: int):
    v = ws.cell(r, c).value
    if v is not None:
        return v
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= r <= rng.max_row and rng.min_col <= c <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col).value
    return None


def header_rows(ws, need: list[str], scan: int = 6) -> int | None:
    """need의 모든 단어가 들어 있는 첫 행(머리글)."""
    for r in range(1, min(ws.max_row, scan) + 1):
        vals = " ".join(str(merged_value(ws, r, c) or "") for c in range(1, ws.max_column + 1))
        if all(n in vals for n in need):
            return r
    return None


def all_cells(wb) -> list:
    return [c for ws in wb.worksheets for row in ws.iter_rows() for c in row]


def numeric_values(wb) -> list[float]:
    return [float(c.value) for c in all_cells(wb) if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]


def as_date_str(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return None if v is None else str(v).strip()


def blank(v) -> bool:
    return v is None or str(v).strip() == ""


def final_text(run: Path) -> str:
    p = run / "outputs" / "final_response.md"
    return p.read_text(encoding="utf-8", errors="replace") if p.is_file() else ""


# ───────────────────────── 1. 정리 ─────────────────────────
def find_clean_sheet(paths: list[Path]):
    best = None
    for p in paths:
        wb = load_workbook(p, data_only=True)
        for ws in wb.worksheets:
            hr = header_rows(ws, ["연락처", "가입일"])
            if hr is None:
                continue
            score = 0
            t = ws.title
            if re.search(r"정리|클린|완료|clean|결과", t, re.I):
                score += 2
            if re.search(r"원본|original|raw|before", t, re.I):
                score -= 3
            if best is None or score > best[0]:
                best = (score, p, ws.title, hr)
    return best


def grade_clean(run: Path, meta: dict) -> list[dict]:
    K = KEY["실습1"]
    exps = {t: (None, "") for t in meta["assertions"]}
    A = meta["assertions"]
    outs = outputs(run)
    ok_in, why = input_unchanged(run, "실습1_고객관리_원본.xlsx")
    exps[A[0]] = (bool(outs) and ok_in, f"결과 {[p.name for p in outs]}; {why}")
    found = find_clean_sheet(outs)
    facts = [f"결과 파일: {[p.name for p in outs]}"]
    if not found:
        for t in A[1:9]:
            exps[t] = (False, "정리된 데이터 시트(연락처·가입일 열)를 찾지 못함")
    else:
        _, p, title, hr = found
        wb = load_workbook(p, data_only=True)
        ws = wb[title]
        head = [str(merged_value(ws, hr, c) or "").strip() for c in range(1, ws.max_column + 1)]
        col = {h: i for i, h in enumerate(head) if h}
        rows = []
        for r in ws.iter_rows(min_row=hr + 1, values_only=True):  # 첫 빈 행에서 데이터가 끝난다(아래 범례·메모 제외)
            if not any(not blank(v) for v in r):
                break
            rows.append(r)
        facts.append(f"정리 시트: {p.name}!{title} (머리글 {hr}행) 열 {head} 데이터 {len(rows)}행")
        orig = {int(r[0]): r for r in load_workbook(INPUTS / "실습1_고객관리_원본.xlsx").active.iter_rows(min_row=2, values_only=True)}
        ohead = ["번호", "이름", "연락처", "이메일", "가입일", "등급", "누적구매액", "지역", "비고"]
        if "번호" in col:
            by_no = {}
            for r in rows:
                try:
                    by_no[int(r[col["번호"]])] = r
                except (TypeError, ValueError):
                    pass
        else:  # 번호 열을 지웠다면 이름+이메일로 원본 번호를 찾는다
            idx = {(o[1], o[3]): n for n, o in orig.items() if n <= 30 or n == 33}
            by_no = {idx.get((r[col.get("이름", 0)], r[col.get("이메일", 0)])): r for r in rows}
            by_no.pop(None, None)
        removed = set(K["exact_duplicate_rows"])
        exp_nos = set(range(1, 34)) - removed
        exps[A[1]] = (len(rows) == 31 and set(by_no) == exp_nos, f"데이터 {len(rows)}행, 번호 {sorted(set(range(1, 34)) - set(by_no))} 없음")

        def get(n, name):
            r = by_no.get(n)
            return None if r is None or name not in col else r[col[name]]
        pm = [f"{n}:{get(int(n), '연락처')}≠{v}" for n, v in K["phone_normalize"].items() if int(n) not in removed and get(int(n), "연락처") != v]
        exps[A[2]] = (not pm, f"불일치 {len(pm)}건 {pm[:4]}")
        dm = [f"{n}:{as_date_str(get(int(n), '가입일'))}≠{v}" for n, v in K["date_normalize"].items()
              if int(n) not in removed and as_date_str(get(int(n), "가입일")) != v]
        exps[A[3]] = (not dm, f"불일치 {len(dm)}건 {dm[:4]}")
        bad = [f"{n}:{orig[n][2]}→{get(n, '연락처')}" for n in K["invalid_phone_rows"] if str(get(n, "연락처")) != str(orig[n][2])]
        exps[A[4]] = (not bad, f"바뀐 값 {bad}" if bad else "원래 값 유지")
        ov = get(14, "누적구매액")
        exps[A[5]] = (ov == 98_000_000, f"번호 14 누적구매액 = {ov}")
        filled = []
        for c, nos in K["blank_cells"].items():
            for n in nos:
                v = get(n, c)
                if not blank(v) and not MARKER.search(str(v)):
                    filled.append(f"{n}/{c}={v}")
        exps[A[6]] = (not filled, f"채운 값 {filled}" if filled else "빈 칸 유지(표시 문구는 허용)")
        # 이메일 오타 표시: 다른 시트 문구, 메모, 채우기 색 중 하나
        allwb = [load_workbook(q) for q in outs]
        texts = " ".join(str(c.value) for w in allwb for s in w.worksheets if s.title != title for row in s.iter_rows() for c in row if c.value)
        marked = []
        for n in K["email_typo_rows"]:
            email = orig[int(n)][3]
            dom = email.split("@")[1]
            hit = email in texts or dom in texts
            wsf = load_workbook(p)[title]
            for row in wsf.iter_rows(min_row=hr + 1):
                for cell in row:
                    if cell.value == email and (cell.comment is not None or (cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb not in (None, "00000000"))):
                        hit = True
            marked.append((n, hit))
        exps[A[7]] = (all(h for _, h in marked), f"표시 여부 {marked}")
        # 변경 내역: 같은 행에 원래 값과 바뀐 값이 함께 있는 기록
        pairs = []
        for n, v in list(K["phone_normalize"].items())[:25]:
            pairs.append((str(orig[int(n)][2]), v))
        for n, v in list(K["date_normalize"].items())[:25]:
            pairs.append((str(orig[int(n)][4]), v))
        logged = 0
        for w in allwb:
            for s in w.worksheets:
                if s.title == title:
                    continue
                for row in s.iter_rows(values_only=True):
                    cells = {str(as_date_str(x)) for x in row if x is not None}
                    logged += sum(1 for b, a in pairs if b in cells and a in cells)
        exps[A[8]] = (logged >= 10, f"이전·이후 값이 같은 행에 기록된 변경 {logged}건")
        sheets = [f"{q.name}: {load_workbook(q, read_only=True).sheetnames}" for q in outs]
        facts += sheets
    (run / "_grade").mkdir(exist_ok=True)
    (run / "_grade" / "facts.md").write_text("\n".join(facts) + "\n\n## 최종 응답\n" + final_text(run), encoding="utf-8")
    return [{"text": t, "passed": exps[t][0], "evidence": exps[t][1] or "판단 대기"} for t in A]


# ───────────────────────── 2. 분석 ─────────────────────────
CELL_REF = re.compile(r"^[^!\s]+![A-Z]+\d+")
NUMBERED = re.compile(r"^\s*([①-⑩]|\d{1,2}\s*[.)]|\d{1,2}$)")


def row_text(row) -> str:
    """한 행의 문장(근거 셀 참조 제외). 제목 칸 + 설명 칸으로 나뉜 인사이트는 합쳐서 하나로 본다."""
    return " ".join(str(v).strip() for v in row if isinstance(v, str) and v.strip() and not CELL_REF.match(v.strip()))


def insight_texts(wb) -> tuple[list[str], str]:
    for ws in wb.worksheets:
        if "인사이트" in ws.title or "insight" in ws.title.lower():
            rows = [r for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]
            numbered = []
            for r in rows:
                first = next((v for v in r if v is not None and str(v).strip()), None)
                if first is not None and NUMBERED.match(str(first)) and len(row_text(r)) >= 20:
                    numbered.append(row_text(r))
            if numbered:  # 번호 붙은 행(1~5, ①~⑤)이 인사이트다
                return numbered, ws.title
            t = [row_text(r) for r in rows if len(row_text(r)) >= 20]
            if t:
                return t, ws.title
    for ws in wb.worksheets:  # 시트 안 '인사이트' 제목 아래 문장
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "인사이트" in c.value and len(c.value) < 30:
                    out = []
                    for r in range(c.row + 1, min(c.row + 12, ws.max_row) + 1):
                        for cc in range(1, ws.max_column + 1):
                            v = ws.cell(r, cc).value
                            if isinstance(v, str) and len(v.strip()) >= 20:
                                out.append(v.strip())
                    if out:
                        return out, f"{ws.title}!{c.coordinate} 아래"
    return [], ""


def grade_analysis(run: Path, meta: dict) -> list[dict]:
    K = KEY["실습2"]
    A = meta["assertions"]
    exps = {t: (None, "") for t in A}
    outs = outputs(run)
    ok_in, why = input_unchanged(run, "실습2_쇼핑몰매출_원본.xlsx")
    exps[A[0]] = (bool(outs) and ok_in, f"결과 {[p.name for p in outs]}; {why}")
    charts, images = 0, 0
    for p in outs:
        with zipfile.ZipFile(p) as z:
            charts += sum(1 for n in z.namelist() if re.match(r"xl/charts/chart\d+\.xml$", n))
            images += sum(1 for n in z.namelist() if n.startswith("xl/media/"))
    exps[A[1]] = (charts >= 2, f"엑셀 차트 {charts}개, 그림 {images}개")
    rec = [load_workbook(recalc_copy(p, run), data_only=True) for p in outs]
    nums = [v for w in rec for v in numeric_values(w)]
    exact = {round(abs(v), 6) for v in nums}
    miss_m = [m for m, v in K["monthly_amount"].items() if round(float(v), 6) not in exact]
    exps[A[2]] = (bool(outs) and not miss_m, f"없는 월 {miss_m}" if miss_m else "12개월 값 모두 셀에 있음")
    miss_c = [c for c, v in K["category_amount"].items() if round(float(v), 6) not in exact]
    exps[A[3]] = (bool(outs) and not miss_c, f"없는 카테고리 {miss_c}" if miss_c else "4개 값 모두 셀에 있음")
    exps[A[4]] = (88097750.0 in exact, "합계 셀 있음" if 88097750.0 in exact else "합계 88,097,750 셀 없음")
    texts, where = [], ""
    for w in rec:
        texts, where = insight_texts(w)
        if texts:
            break
    exps[A[5]] = (3 <= len(texts) <= 5, f"인사이트 {len(texts)}개({where})")
    bad = []
    for t in texts:
        for tok, mult in numbers_with_units(t):
            if not exempt(tok) and not has_number(nums, exact, tok, mult):
                bad.append(tok)
    exps[A[6]] = (bool(texts) and not bad, f"셀에 없는 숫자 {bad[:6]}" if bad else (f"숫자 대조 통과({len(texts)}문장)" if texts else "인사이트 없음"))
    alltext = " ".join(str(c.value) for w in rec for c in all_cells(w) if isinstance(c.value, str)) + final_text(run)
    exps[A[7]] = ("스마트 워치" in alltext or "스마트워치" in alltext, "스마트 워치 언급" if "스마트 워치" in alltext else "언급 없음")
    (run / "_grade").mkdir(exist_ok=True)
    (run / "_grade" / "facts.md").write_text(
        f"결과: {[p.name for p in outs]}\n시트: {[w.sheetnames for w in rec]}\n차트 {charts}, 그림 {images}\n\n## 인사이트({where})\n"
        + "\n".join(f"- {t}" for t in texts) + "\n\n## 최종 응답\n" + final_text(run), encoding="utf-8")
    return [{"text": t, "passed": exps[t][0], "evidence": exps[t][1] or "판단 대기"} for t in A]


# ───────────────────────── 3. 취합 ─────────────────────────
def norm(v):
    if isinstance(v, str):
        s = v.strip()
        return s.replace(" ", "").upper() if s.startswith("=") else s
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if hasattr(v, "isoformat"):
        return v.isoformat()[:19]
    return v


def grade_consolidate(run: Path, meta: dict) -> list[dict]:
    K = KEY["실습3"]
    A = meta["assertions"]
    exps = {t: (None, "") for t in A}
    outs = outputs(run)
    src = INPUTS / "실습3_멀티채널판매_원본.xlsx"
    ok_in, why = input_unchanged(run, src.name)
    exps[A[0]] = (bool(outs) and ok_in, f"결과 {[p.name for p in outs]}; {why}")
    iwb = load_workbook(src)
    union = set()
    for ws in load_workbook(src, data_only=True).worksheets:
        hr = header_rows(ws, ["상품코드"])
        if hr:
            j = [str(ws.cell(hr, c).value) for c in range(1, ws.max_column + 1)].index("상품코드") + 1
            union |= {ws.cell(r, j).value for r in range(hr + 1, ws.max_row + 1) if ws.cell(r, j).value}
    best = None
    for p in outs:
        wb = load_workbook(p)
        missing = [s for s in iwb.sheetnames if s not in wb.sheetnames]
        changed = []
        for ws in iwb.worksheets:
            if ws.title in wb.sheetnames:
                o = wb[ws.title]
                d = sum(1 for row in ws.iter_rows() for c in row if norm(c.value) != norm(o[c.coordinate].value))
                if d:
                    changed.append(f"{ws.title}({d})")
        cand = [s for s in wb.sheetnames if s not in iwb.sheetnames and header_rows(wb[s], ["상품코드"])]
        score = (not missing) + (not changed) + bool(cand)
        if best is None or score > best[0]:
            best = (score, p, missing, changed, cand)
    if best is None:
        for t in A[1:8]:
            exps[t] = (False, "결과 파일 없음")
        return [{"text": t, "passed": exps[t][0], "evidence": exps[t][1] or "판단 대기"} for t in A]
    _, p, missing, changed, cand = best
    exps[A[1]] = (not missing and not changed, f"누락 {missing}, 변경 {changed}" if (missing or changed) else "원본 탭 7개 내용 동일")
    wbf = load_workbook(p)
    rec = load_workbook(recalc_copy(p, run), data_only=True)
    facts = [f"결과 {p.name} 시트 {wbf.sheetnames}", f"통합 후보 시트 {cand}"]
    # 통합 시트: 상품코드 행이 가장 많은 새 시트
    tbest = None
    for s in cand:
        ws = rec[s]
        hr = header_rows(ws, ["상품코드"])
        j = next((c for c in range(1, ws.max_column + 1) if str(merged_value(ws, hr, c) or "").strip() == "상품코드"), None) if hr else None
        if not j:
            continue
        keys = [ws.cell(r, j).value for r in range(hr + 1, ws.max_row + 1) if ws.cell(r, j).value in union]
        if tbest is None or len(keys) > len(tbest[3]):
            tbest = (s, hr, j, keys)
    if not tbest:
        for t in A[2:8]:
            exps[t] = (False, "상품코드가 있는 통합 시트를 찾지 못함")
    else:
        s, hr, j, keys = tbest
        ws, wf = rec[s], wbf[s]
        exps[A[2]] = (len(keys) == len(set(keys)) == K["union_codes"] and set(keys) == union, f"'{s}' 상품코드 {len(keys)}개(고유 {len(set(keys))})")
        data_rows = [r for r in range(hr + 1, ws.max_row + 1) if ws.cell(r, j).value in union]
        heads = {}
        for c in range(1, ws.max_column + 1):
            heads[c] = " ".join(str(merged_value(ws, r, c) or "") for r in range(1, hr + 1))
        sums, counts, fcols = {}, {}, {}
        for ch, aliases in CHANNELS.items():
            cols = [c for c, h in heads.items() if any(a in h for a in aliases) and re.search(r"판매량|수량|판매수", h) and not re.search(r"합계|전체|총", h.replace(ch, ""))]
            if not cols:
                continue
            c = cols[0]
            vals = [ws.cell(r, c).value for r in data_rows]
            sums[ch] = sum(v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool))
            counts[ch] = sum(1 for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool))
            fs = [wf.cell(r, c).value for r in data_rows]
            fcols[ch] = sum(1 for f in fs if isinstance(f, str) and f.startswith("=") and "!" in f) / max(len(fs), 1)
        exp_s = {k: v["qty_sum"] for k, v in K["channels"].items()}
        exp_c = {k: v["products"] for k, v in K["channels"].items()}
        exps[A[3]] = (sums == exp_s, f"판매량 합계 {sums} / 정답 {exp_s}")
        exps[A[4]] = (counts == exp_c, f"값이 있는 칸 {counts} / 정답 {exp_c}")
        exps[A[5]] = (len(fcols) == 4 and all(v >= 0.9 for v in fcols.values()), f"수식 비율 {({k: round(v, 2) for k, v in fcols.items()})}")
        facts.append(f"통합 시트 '{s}' 머리글 {hr}행, 열 {[heads[c].strip() for c in heads]}")
    errs = [f"{w.title}!{c.coordinate}={c.value}" for w in rec.worksheets for row in w.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.strip() in ERRORS]
    exps[A[6]] = (not errs, f"오류 셀 {len(errs)}개 {errs[:4]}")
    forb = [f"{w.title}!{c.coordinate}" for w in wbf.worksheets for row in w.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("=") and FORBIDDEN.search(c.value)]
    exps[A[7]] = (not forb, f"호환 안 되는 함수 {len(forb)}칸 {forb[:4]}" if forb else "없음")
    (run / "_grade").mkdir(exist_ok=True)
    (run / "_grade" / "facts.md").write_text("\n".join(facts) + "\n\n## 최종 응답\n" + final_text(run), encoding="utf-8")
    return [{"text": t, "passed": exps[t][0], "evidence": exps[t][1] or "판단 대기"} for t in A]


GRADERS = {"clean-customer-list": grade_clean, "analyze-shop-sales": grade_analysis, "consolidate-multichannel": grade_consolidate}


def main(argv: list[str]) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("iteration")
    ap.add_argument("--judgments")
    a = ap.parse_args(argv)
    it = Path(a.iteration)
    judg = json.loads(Path(a.judgments).read_text(encoding="utf-8")) if a.judgments else {}
    pending = 0
    for meta_file in sorted(it.glob("eval-*/eval_metadata.json")):
        meta = json.loads(meta_file.read_text(encoding="utf-8"))
        for cfg_dir in sorted(p for p in meta_file.parent.iterdir() if p.is_dir()):
            for run_dir in sorted(cfg_dir.glob("run-*")):
                key = f"{meta['eval_name']}/{cfg_dir.name}"
                gfile = run_dir / "grading.json"
                if a.judgments and gfile.is_file():
                    exps = json.loads(gfile.read_text(encoding="utf-8"))["expectations"]
                else:
                    exps = GRADERS[meta["eval_name"]](run_dir, meta)
                for x in exps:
                    j = judg.get(key, {}).get(x["text"])
                    if j:
                        x["passed"], x["evidence"] = bool(j["passed"]), j["evidence"]
                passed = sum(1 for x in exps if x["passed"] is True)
                decided = sum(1 for x in exps if x["passed"] is not None)
                pending += len(exps) - decided
                gfile.write_text(json.dumps({"expectations": exps, "summary": {
                    "passed": passed, "failed": decided - passed, "total": len(exps),
                    "pass_rate": round(passed / len(exps), 4) if exps else 0}}, ensure_ascii=False, indent=2), encoding="utf-8")
                print(f"{key:<40} {passed}/{len(exps)}  (판단 대기 {len(exps) - decided})")
    print(f"PENDING JUDGMENTS {pending}" if pending else "GRADED ALL")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
