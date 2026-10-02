#!/usr/bin/env python3
"""doc-automation / hwpx-editor 평가 실행 결과 채점 (skill-creator grading.json 형식).

객관 항목은 이 스크립트가 직접 측정하고, [판단] 항목은 judgments.json으로 채점자가 판정을 넣는다.
두 구성(with_skill, old_skill)을 같은 기준으로 채점하기 위해 출처는 각 실행의 원래 입력 파일을
이 저장소의 extract_sources.py로 독립 추출한다(실행이 만든 sources는 쓰지 않는다).

산출물(각 run 디렉터리)
  grading.json            expectations[{text, passed, evidence}], summary
  _grade/deck_text.md     슬라이드별 텍스트·제목·노트(판단 항목 채점용)
  outputs/preview-slide-NN.png  슬라이드 이미지(사람 검토용, 렌더 가능 시)

사용법
  python tools/grade_doc_runs.py <iteration 디렉터리> [--judgments judgments.json] [--no-render]
judgments.json: {"<eval-name>/<config>": {"<기대 항목 원문>": {"passed": true, "evidence": "..."}}}
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DOC = ROOT / "plugins/claude-skills-book/skills/doc-automation/scripts"
HWPX = ROOT / "plugins/claude-skills-book/skills/hwpx-editor/scripts"
PY = sys.executable

sys.path.insert(0, str(HWPX))


def run(*args, timeout=600) -> subprocess.CompletedProcess:
    return subprocess.run([PY, *map(str, args)], capture_output=True, text=True, encoding="utf-8",
                          errors="replace", timeout=timeout)


def main_pptx(out: Path) -> Path | None:
    cands = [p for p in out.glob("*.pptx") if not p.name.startswith("~$")]
    return max(cands, key=lambda p: p.stat().st_mtime) if cands else None


def deck_facts(pptx: Path) -> dict:
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER
    prs = Presentation(str(pptx))
    slides = []
    for i, s in enumerate(prs.slides, 1):
        title = ""
        in_ph = False
        for ph in s.placeholders:
            if ph.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE) and ph.text_frame.text.strip():
                title, in_ph = ph.text_frame.text.strip(), True
        texts = []
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text.strip())
                if not title and sh.name == "Headline":
                    title = sh.text_frame.text.strip()
            if getattr(sh, "has_table", False) and sh.has_table:
                texts.append("\n".join(" | ".join(c.text for c in r.cells) for r in sh.table.rows))
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                texts.append(f"[차트] " + "; ".join(f"{ser.name}: {list(ser.values)}" for ser in ch.plots[0].series))
        if not title and texts:
            title = texts[0].split("\n")[0]
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        slides.append({"n": i, "title": title, "title_in_placeholder": in_ph, "texts": texts, "notes": notes,
                       "pictures": sum(1 for sh in s.shapes if sh.shape_type == 13),
                       "charts": sum(1 for sh in s.shapes if getattr(sh, "has_chart", False) and sh.has_chart)})
    return {"size": (round(prs.slide_width / 914400, 2), round(prs.slide_height / 914400, 2)), "slides": slides}


def grade_doc(run_dir: Path, meta: dict, render: bool) -> list[dict]:
    out, work, g = run_dir / "outputs", run_dir / "work", run_dir / "_grade"
    if g.exists():
        shutil.rmtree(g)
    g.mkdir()
    exps = meta["assertions"]
    res = {}
    pptx = main_pptx(out)
    inputs = [p for p in work.iterdir() if "template" not in p.name.lower()] if work.is_dir() else []
    # 실행 중 생긴 산출물 폴더는 제외하고 원래 입력만 추출(파일 이름 기준)
    original = json.loads((run_dir.parents[2] / "run_inputs.json").read_text(encoding="utf-8")).get(
        f"{run_dir.parent.parent.name}", None) if (run_dir.parents[2] / "run_inputs.json").is_file() else None
    if original:
        inputs = [work / n for n in original if "template" not in n.lower()]
    src = g / "sources"
    run(DOC / "extract_sources.py", *inputs, "--out", src)
    facts, ver = None, None
    if pptx:
        facts = deck_facts(pptx)
        outline = None
        for cand in list(work.rglob("outline.json")) + list(out.rglob("outline.json")):
            outline = cand
            break
        vargs = [DOC / "verify_deck.py", pptx, "--sources", src, "--json"]
        if outline:
            # 실행이 선언한 derived/given만 인정하고 슬라이드 수 비교는 하지 않는다(outline 사본에서 slides 제거)
            o = json.loads(outline.read_text(encoding="utf-8"))
            slim = {"derived": o.get("derived", []), "given": o.get("given", [])}
            (g / "declared.json").write_text(json.dumps(slim, ensure_ascii=False), encoding="utf-8")
            vargs += ["--outline", g / "declared.json"]
        r = run(*vargs)
        try:
            ver = json.loads(r.stdout)
        except json.JSONDecodeError:
            ver = {"errors": [{"rule": "verify-crash", "detail": r.stdout[-300:] + r.stderr[-300:]}], "numbers_unverified": -1,
                   "numbers_checked": 0}
        lines = [f"# {pptx.name}  size={facts['size']}  slides={len(facts['slides'])}"]
        for s in facts["slides"]:
            lines.append(f"\n## {s['n']}. {s['title']}  (제목칸={s['title_in_placeholder']}, 그림={s['pictures']}, 차트={s['charts']})")
            lines += ["  " + t.replace("\n", "\n  ") for t in s["texts"]]
            if s["notes"]:
                lines.append("  (노트) " + s["notes"].replace("\n", " ")[:300])
        lines.append("\n# verify_deck")
        lines += [f"- {e['rule']} #{e.get('slide')}: {e['detail']}" for e in ver.get("errors", [])]
        (g / "deck_text.md").write_text("\n".join(lines), encoding="utf-8")
        if render:
            pv = g / "png"
            rr = run(DOC / "render_slides.py", pptx, "--out", pv, timeout=600)
            if "RENDER OK" in rr.stdout:
                for f in sorted(pv.glob("slide-*.png")):
                    shutil.copy2(f, out / f"preview-{f.name}")
    for e in exps:
        if e.startswith("[판단]"):
            continue
        ok, ev = None, ""
        if ".pptx 보고서가 있다" in e:
            ok, ev = pptx is not None, (pptx.name if pptx else "pptx 없음")
        elif not pptx and "메일" not in e:
            ok, ev = False, "pptx 없음"
        elif "숫자" in e and "확인된다" in e:
            u = ver.get("numbers_unverified", -1)
            ok = u == 0
            bad = [x["detail"] for x in ver.get("errors", []) if x["rule"] == "number-provenance"][:5]
            ev = f"대조 {ver.get('numbers_checked')}개, 출처 미확인 {u}개" + (f": {bad}" if bad else "")
        elif "빈 플레이스홀더" in e:
            bad = [x for x in ver.get("errors", []) if x["rule"] in ("empty-placeholder", "leftover-text")]
            ok, ev = not bad, (f"{len(bad)}건: " + "; ".join(f"#{x['slide']} {x['detail']}" for x in bad[:4])) if bad else "없음"
        elif "4장 이상 12장 이하" in e:
            n = len(facts["slides"]); ok, ev = 4 <= n <= 12, f"{n}장"
        elif "10장 이하" in e:
            n = len(facts["slides"]); ok, ev = n <= 10, f"{n}장"
        elif "로고" in e:
            p1 = facts["slides"][0]["pictures"]; ok, ev = p1 > 0, f"표지 그림 {p1}개"
        elif "메일 파일" in e:
            mails = [p for p in out.iterdir() if p.suffix in (".md", ".txt") and p.name != "final_response.md"
                     and ("mail" in p.name.lower() or "메일" in p.name or "제목" in p.read_text(encoding="utf-8", errors="ignore")[:500])]
            ok, ev = bool(mails), ", ".join(p.name for p in mails) or "메일 파일 없음"
        elif "네이티브 차트" in e:
            n = sum(s["charts"] for s in facts["slides"]); ok, ev = n > 0, f"차트 {n}개"
        elif "4:3" in e:
            ok, ev = tuple(facts["size"]) == (10.0, 7.5), f"크기 {facts['size']}"
        elif "제목 칸" in e:
            content = facts["slides"][1:]
            k = sum(1 for s in content if s["title_in_placeholder"])
            ok, ev = bool(content) and k * 2 >= len(content), f"내용 슬라이드 {len(content)}장 중 {k}장"
        res[e] = {"passed": ok, "evidence": ev}
    return [{"text": e, "passed": res[e]["passed"], "evidence": res[e]["evidence"]} if e in res
            else {"text": e, "passed": None, "evidence": "판단 대기"} for e in exps]


def grade_hwpx(run_dir: Path, meta: dict) -> list[dict]:
    import hwpx_template as ht
    import verify_hwpx as vh
    out, work = run_dir / "outputs", run_dir / "work"
    originals = {p.name: p for p in work.glob("*.hwpx")}
    results = [p for p in list(out.glob("*.hwpx")) if p.name not in originals]
    src = next(iter(originals.values()), None)
    res = {}
    dst = max(results, key=lambda p: p.stat().st_mtime) if results else None
    text = "\n".join(ht.extract_texts(str(dst))) if dst else ""
    for e in meta["assertions"]:
        if e.startswith("[판단]"):
            continue
        if "새 .hwpx" in e:
            res[e] = (dst is not None, dst.name if dst else "결과 hwpx 없음")
        elif "verify_hwpx" in e:
            v = vh.verify(str(src), str(dst)) if dst and src else {"ok": False, "errors": ["파일 없음"]}
            res[e] = (v["ok"], "; ".join(v["errors"][:3]) or "통과")
        elif "보고서 작성 서식 안내" in e:
            res[e] = ("보고서 작성 서식 안내" not in text, "남아 있음" if "보고서 작성 서식 안내" in text else "없음")
        elif "'불꽃'" in e:
            res[e] = ("불꽃" in text, "있음" if "불꽃" in text else "없음")
        elif "모두 있다" in e:
            miss = [w for w in ("한강사업본부", "김한강", "박여의") if w not in text]
            res[e] = (not miss, f"누락 {miss}" if miss else "모두 있음")
        elif "남아 있지 않다" in e and "서울대공원" in e:
            left = [w for w in ("서울대공원", "강준민") if w in text]
            res[e] = (not left, f"잔존 {left}" if left else "없음")
        elif "{{" in e:
            left = re.findall(r"\{\{\w+\}\}", text)
            res[e] = (not left, f"잔존 {left[:5]}" if left else "없음")
        elif "'한강사업본부'와 '김한강'" in e:
            ok = "한강사업본부" in text and "김한강" in text
            res[e] = (ok, "있음" if ok else "없음")
    return [{"text": e, "passed": res[e][0], "evidence": res[e][1]} if e in res
            else {"text": e, "passed": None, "evidence": "판단 대기"} for e in meta["assertions"]]


def main(argv: list[str]) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("iteration")
    ap.add_argument("--judgments")
    ap.add_argument("--no-render", action="store_true")
    a = ap.parse_args(argv)
    it = Path(a.iteration)
    judg = json.loads(Path(a.judgments).read_text(encoding="utf-8")) if a.judgments else {}
    pending = 0
    for meta_file in sorted(it.glob("eval-*/eval_metadata.json")):
        meta = json.loads(meta_file.read_text(encoding="utf-8"))
        for cfg_dir in sorted(p for p in meta_file.parent.iterdir() if p.is_dir()):
            for run_dir in sorted(cfg_dir.glob("run-*")):
                key = f"{meta['eval_name']}/{cfg_dir.name}"
                gfile = run_dir / "grading.json"
                if a.judgments and gfile.is_file():
                    exps = json.loads(gfile.read_text(encoding="utf-8"))["expectations"]
                else:
                    is_hwpx = any(run_dir.joinpath("work").glob("*.hwpx")) and not any(run_dir.joinpath("work").glob("*.pdf")) \
                        and "hwpx" in json.dumps(meta, ensure_ascii=False) and "PPT" not in meta["prompt"]
                    exps = grade_hwpx(run_dir, meta) if is_hwpx else grade_doc(run_dir, meta, not a.no_render)
                for x in exps:
                    j = judg.get(key, {}).get(x["text"])
                    if j:
                        x["passed"], x["evidence"] = bool(j["passed"]), j["evidence"]
                passed = sum(1 for x in exps if x["passed"] is True)
                decided = sum(1 for x in exps if x["passed"] is not None)
                pending += len(exps) - decided
                gfile.write_text(json.dumps({"expectations": exps, "summary": {
                    "passed": passed, "failed": decided - passed, "total": len(exps),
                    "pass_rate": round(passed / len(exps), 4) if exps else 0}}, ensure_ascii=False, indent=2), encoding="utf-8")
                print(f"{key:<45} {passed}/{len(exps)}  (판단 대기 {len(exps) - decided})")
    print(f"PENDING JUDGMENTS {pending}" if pending else "GRADED ALL")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
