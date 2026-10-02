#!/usr/bin/env python3
"""excel-automation 실습 원본과 정답표를 만든다(시드 고정 → 매번 같은 파일).

왜: v1.6.1 실습 폴더의 '실습1~3' 파일은 원본이 아니라 이미 처리된 결과물이었다(정리완료·인사이트·통합관리 시트 포함).
수강생이 정리·분석·취합을 해 볼 원본이 없었으므로 다시 만든다. 정답표(answer_key)는 평가 채점에 쓴다.

출력(실습 플러그인 samples/excel-automation/inputs/)
  실습1_고객관리_원본.xlsx      고객 명단 — 중복·전화번호·날짜 형식·이메일 오타·빈 칸·이상값을 의도적으로 심음
  실습2_쇼핑몰매출_원본.xlsx    12개월 주문 내역(2024-04 ~ 2025-03)
  실습3_멀티채널판매_원본.xlsx  탭 7개(전체현황 + 상품마스터·재고현황·채널 4개) — v3 완료본에서 통합관리 시트만 제거
  answer_key.json             세 파일의 정답(개수·정규화 기대값·합계)

사용법: python tools/dev/make_excel_practice_inputs.py [--check]   (--check: 다시 만들어 기존 파일과 같은지 비교)
"""
from __future__ import annotations

import hashlib
import json
import random
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
SAMPLES = ROOT / "plugins/claude-skills-practice/skills/practice-samples/samples/excel-automation"
OUT = SAMPLES / "inputs"
V3 = SAMPLES / "excel-automation-outputs" / "실습3_통합관리_완료_v3.xlsx"

SURNAMES = list("김이박최정강조윤장임한오서신권황안송류홍")
GIVEN = ["민수", "서연", "지훈", "하은", "도윤", "수빈", "현우", "지민", "예준", "유진", "시우", "채원", "준서", "다은",
         "건우", "소율", "우진", "하린", "선우", "나연", "태윤", "가은", "은호", "서윤", "민재", "지아", "승현", "예린"]
REGIONS = ["서울 강남구", "서울 마포구", "경기 성남시", "경기 수원시", "부산 해운대구", "인천 연수구", "대구 수성구", "광주 서구"]
DOMAINS = ["gmail.com", "naver.com", "daum.net", "kakao.com", "hanmail.net"]
GRADES = ["VIP", "골드", "일반"]


def fixed_meta(wb: Workbook) -> None:
    """파일 바이트를 재현 가능하게: 생성·수정 시각을 고정한다."""
    from datetime import datetime
    wb.properties.created = datetime(2026, 10, 1)
    wb.properties.modified = datetime(2026, 10, 1)
    wb.properties.creator = "claude-skills practice generator"


def make_customers(rng: random.Random) -> tuple[Workbook, dict]:
    base = []
    used = set()
    while len(base) < 30:
        name = rng.choice(SURNAMES) + rng.choice(GIVEN)
        if name in used:
            continue
        used.add(name)
        mid, last = rng.randint(1000, 9999), rng.randint(1000, 9999)
        d = date(2024, 1, 1) + timedelta(days=rng.randint(0, 420))
        uid = f"{name}{rng.randint(10, 99)}"
        base.append({"이름": name, "phone": f"010-{mid}-{last}", "email": f"user{len(base)+1:02d}@{rng.choice(DOMAINS)}",
                     "date": d, "등급": rng.choice(GRADES), "누적구매액": rng.randrange(50_000, 3_000_000, 10_000),
                     "지역": rng.choice(REGIONS), "uid": uid})
    key = {"phone_normalize": {}, "date_normalize": {}, "invalid_phone_rows": [], "email_typo_rows": {},
           "blank_cells": {}, "outlier_rows": [], "exact_duplicate_rows": [], "near_duplicate_groups": []}
    rows = []
    phone_styles = ["{a}-{b}-{c}", "{a}{b}{c}", "{a} {b} {c}", "{a}.{b}.{c}", "+82-10-{b}-{c}"]
    date_styles = ["{y}-{m:02d}-{d:02d}", "{y}/{m:02d}/{d:02d}", "{y}.{m}.{d}", "{y}{m:02d}{d:02d}", "{y}년 {m}월 {d}일"]
    for i, b in enumerate(base, 1):
        a, bb, c = b["phone"].split("-")
        ps = phone_styles[i % len(phone_styles)]
        phone = ps.format(a=a, b=bb, c=c)
        ds = date_styles[i % len(date_styles)]
        dd = b["date"]
        dt = ds.format(y=dd.year, m=dd.month, d=dd.day)
        rows.append({"번호": i, "이름": b["이름"], "연락처": phone, "이메일": b["email"], "가입일": dt, "등급": b["등급"],
                     "누적구매액": b["누적구매액"], "지역": b["지역"], "비고": None})
    # 이상값·오류를 특정 행에 심는다(행 번호는 '번호' 열 기준)
    rows[4]["연락처"] = "010-777-888"          # 자릿수 오류
    rows[11]["연락처"] = "010-7788-99OO"        # 영문 O 혼입
    rows[6]["이메일"] = rows[6]["이메일"].split("@")[0] + "@gamil.com"
    rows[15]["이메일"] = rows[15]["이메일"].split("@")[0] + "@naver.con"
    rows[8]["지역"] = None
    rows[19]["지역"] = None
    rows[22]["등급"] = None
    rows[13]["누적구매액"] = 98_000_000         # 극단값
    rows[17]["가입일"] = "03/04/2024"           # 월/일 모호
    # 완전 중복 2건(행 전체 같음), 유사 중복 1쌍(이름·연락처 같고 등급만 다름)
    dup1 = dict(rows[2]); dup1["번호"] = 31
    dup2 = dict(rows[9]); dup2["번호"] = 32
    near = dict(rows[20]); near["번호"] = 33; near["등급"] = "VIP" if rows[20]["등급"] != "VIP" else "일반"
    rows += [dup1, dup2, near]
    # 정답 계산
    for r in rows:
        n = r["번호"]
        digits = "".join(ch for ch in str(r["연락처"]) if ch.isdigit())
        if digits.startswith("8210"):
            digits = "0" + digits[2:]
        if r["번호"] in (5, 12):
            key["invalid_phone_rows"].append(n)
        elif len(digits) == 11 and digits.startswith("010"):
            norm = f"{digits[:3]}-{digits[3:7]}-{digits[7:]}"
            if norm != r["연락처"]:
                key["phone_normalize"][str(n)] = norm
    for r in rows:
        n, v = r["번호"], str(r["가입일"])
        if n == 18:
            continue  # 모호 → 바꾸지 않고 후보로
        b = base[(n - 1) if n <= 30 else {31: 2, 32: 9, 33: 20}[n]]
        norm = b["date"].isoformat()
        if v != norm:
            key["date_normalize"][str(n)] = norm
    key["ambiguous_date_rows"] = [18]
    key["email_typo_rows"] = {"7": "gmail.com", "16": "naver.com"}
    key["blank_cells"] = {"지역": [9, 20], "등급": [23]}
    key["outlier_rows"] = [14]
    key["exact_duplicate_rows"] = [31, 32]          # 3번·10번 행의 완전 중복(번호 열만 다름 → 번호 제외 비교)
    key["exact_duplicate_of"] = {"31": 3, "32": 10}
    key["near_duplicate_groups"] = [[21, 33]]
    key["rows_total"] = len(rows)
    key["rows_after_exact_dedup"] = len(rows) - 2
    key["note"] = "완전 중복은 '번호' 열을 제외한 모든 열이 같은 행이다. 번호는 일련번호라 비교에서 뺀다."
    wb = Workbook()
    ws = wb.active
    ws.title = "고객목록"
    cols = ["번호", "이름", "연락처", "이메일", "가입일", "등급", "누적구매액", "지역", "비고"]
    ws.append(cols)
    for r in rows:
        ws.append([r[c] for c in cols])
    fixed_meta(wb)
    return wb, key


def make_sales(rng: random.Random) -> tuple[Workbook, dict]:
    catalog = {
        "전자기기": [("무선 이어폰", 89000), ("블루투스 스피커", 59000), ("보조배터리", 32000), ("스마트 워치", 199000)],
        "뷰티": [("수분 크림", 28000), ("선크림", 19000), ("립밤 세트", 15000), ("클렌징폼", 12000)],
        "생활용품": [("텀블러", 22000), ("수건 세트", 18000), ("디퓨저", 25000), ("밀폐용기", 16000)],
        "식품": [("견과류 세트", 24000), ("그래놀라", 9800), ("콜드브루", 14000), ("단백질바", 21000)],
    }
    regions = ["서울", "경기", "부산", "인천", "대구", "광주", "대전"]
    pays = ["카드", "간편결제", "계좌이체"]
    wb = Workbook()
    ws = wb.active
    ws.title = "주문내역"
    cols = ["주문번호", "주문일", "카테고리", "상품명", "수량", "단가", "할인율", "결제금액", "지역", "결제수단"]
    ws.append(cols)
    monthly: dict[str, int] = {}
    by_cat: dict[str, int] = {}
    by_prod: dict[str, int] = {}
    start = date(2024, 4, 1)
    n = 0
    for day in range(365):
        d = start + timedelta(days=day)
        season = 1.4 if d.month in (11, 12, 1) else (0.8 if d.month in (4, 5) else 1.0)
        for _ in range(max(1, int(rng.gauss(4 * season, 1.5)))):
            n += 1
            cat = rng.choice(list(catalog))
            prod, price = rng.choice(catalog[cat])
            qty = rng.choice([1, 1, 1, 2, 2, 3])
            disc = rng.choice([0, 0, 0.05, 0.1, 0.15])
            amount = int(round(price * qty * (1 - disc), -1))
            ws.append([f"ORD-{n:05d}", d.isoformat(), cat, prod, qty, price, disc, amount, rng.choice(regions), rng.choice(pays)])
            ym = d.strftime("%Y-%m")
            monthly[ym] = monthly.get(ym, 0) + amount
            by_cat[cat] = by_cat.get(cat, 0) + amount
            by_prod[prod] = by_prod.get(prod, 0) + amount
    fixed_meta(wb)
    key = {"rows_total": n, "total_amount": sum(monthly.values()), "monthly_amount": monthly, "category_amount": by_cat,
           "top_product": max(by_prod, key=by_prod.get), "top_product_amount": max(by_prod.values()),
           "best_month": max(monthly, key=monthly.get), "worst_month": min(monthly, key=monthly.get)}
    return wb, key


def make_multichannel() -> tuple[Workbook, dict]:
    wb = load_workbook(V3)  # 수식 그대로 보존(원본 탭에는 수식이 없다)
    wb.remove(wb["통합관리"])
    fixed_meta(wb)
    import pandas as pd
    sheets = {ws.title: pd.DataFrame(list(ws.values)[1:], columns=list(ws.values)[0]) for ws in wb.worksheets if ws.title != "전체현황"}
    codes = set()
    for name, df in sheets.items():
        codes |= set(df["상품코드"].dropna())
    per_channel = {}
    for ch, qty_col in (("스마트스토어", "3월판매량"), ("쿠팡", "3월판매량"), ("카카오선물하기", "3월판매량"), ("자사몰", "3월판매량")):
        df = sheets[ch]
        per_channel[ch] = {"products": int(df["상품코드"].nunique()), "qty_sum": int(df[qty_col].fillna(0).sum())}
    key = {"sheets": [ws.title for ws in wb.worksheets], "key_column": "상품코드", "union_codes": len(codes),
           "master_codes": int(sheets["상품마스터"]["상품코드"].nunique()), "channels": per_channel,
           "channel_registration_counts": {ch: per_channel[ch]["products"] for ch in per_channel}}
    return wb, key


def digest(p: Path) -> str:
    """xlsx는 zip 시각이 들어가므로 셀 값으로 비교한다."""
    wb = load_workbook(p)
    h = hashlib.sha256()
    for ws in wb.worksheets:
        h.update(ws.title.encode())
        for row in ws.iter_rows(values_only=True):
            h.update(repr(row).encode())
    return h.hexdigest()


def build(target: Path) -> dict:
    target.mkdir(parents=True, exist_ok=True)
    rng = random.Random(20261001)
    w1, k1 = make_customers(rng)
    w1.save(target / "실습1_고객관리_원본.xlsx")
    w2, k2 = make_sales(random.Random(20261002))
    w2.save(target / "실습2_쇼핑몰매출_원본.xlsx")
    w3, k3 = make_multichannel()
    w3.save(target / "실습3_멀티채널판매_원본.xlsx")
    key = {"generated_by": "tools/dev/make_excel_practice_inputs.py", "실습1": k1, "실습2": k2, "실습3": k3}
    (target / "answer_key.json").write_text(json.dumps(key, ensure_ascii=False, indent=2), encoding="utf-8")
    return key


def main(argv: list[str]) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    if "--check" in argv:
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            build(Path(td))
            bad = [f.name for f in Path(td).glob("*.xlsx") if digest(f) != digest(OUT / f.name)]
            same_key = json.loads((Path(td) / "answer_key.json").read_text(encoding="utf-8")) == \
                json.loads((OUT / "answer_key.json").read_text(encoding="utf-8"))
        if bad or not same_key:
            print("DIFF:", bad, "answer_key 동일" if same_key else "answer_key 다름")
            return 1
        print("REPRODUCIBLE")
        return 0
    key = build(OUT)
    print(json.dumps({k: {kk: vv for kk, vv in v.items() if not isinstance(vv, dict)} for k, v in key.items() if isinstance(v, dict)},
                     ensure_ascii=False, indent=1)[:1500])
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
