#!/usr/bin/env python3
"""G12 오라클: 예제 docx의 내장 글꼴 제거가 안전하게 되었는지 검증한다.

v1.6.1 태그의 원본과 비교해 다음을 모두 확인한다.
- word/fonts/* 파트가 없다
- 모든 XML 파트가 well-formed다
- fontTable·rels·Content_Types에 글꼴 잔재 참조가 없다
- python-docx로 열리고, 문단·표 텍스트가 원본과 정확히 같다
- 글꼴 외 파트 목록이 원본과 같다
"""
from __future__ import annotations

import io
import sys
import zipfile
import xml.etree.ElementTree as ET

sys.path.insert(0, __import__("os").path.dirname(__file__))
from _lib import BASELINE_TAG, ROOT, force_utf8_stdout, git  # noqa: E402

PAIRS = {
    # 현재 경로: v1.6.1 경로
    "plugins/kevin-skills-practice/skills/practice-samples/samples/doc-automation/example/example_3_comany-to-ppt":
        "skills/doc-automation/assets/example/example_3_comany-to-ppt",
    "chapter02-doc-automation/example/example_3_comany-to-ppt":
        "chapter02-doc-automation/example/example_3_comany-to-ppt",
}


def doc_text(data: bytes) -> list[str]:
    import docx  # python-docx

    d = docx.Document(io.BytesIO(data))
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append("\t".join(c.text for c in row.cells))
    return out


def main() -> int:
    force_utf8_stdout()
    problems: list[str] = []
    checked = 0
    for cur_dir, old_dir in PAIRS.items():
        files = sorted((ROOT / cur_dir).glob("*.docx"))
        if not files:
            problems.append(f"{cur_dir}: docx 없음")
        for f in files:
            checked += 1
            rel_old = f"{old_dir}/{f.name}"
            orig = git("show", f"{BASELINE_TAG}:{rel_old}", binary=True)
            cur = f.read_bytes()
            zc = zipfile.ZipFile(io.BytesIO(cur))
            zo = zipfile.ZipFile(io.BytesIO(orig))
            names = zc.namelist()
            if any(n.startswith("word/fonts/") for n in names):
                problems.append(f"{f.name}: word/fonts 잔존")
            expected = [n for n in zo.namelist() if not n.startswith("word/fonts/")]
            if sorted(names) != sorted(expected):
                problems.append(f"{f.name}: 글꼴 외 파트 목록 불일치")
            for n in names:
                if n.endswith((".xml", ".rels")):
                    try:
                        ET.fromstring(zc.read(n))
                    except ET.ParseError as e:
                        problems.append(f"{f.name}:{n}: XML 오류 {e}")
            ft = zc.read("word/fontTable.xml").decode("utf-8") if "word/fontTable.xml" in names else ""
            if "w:embed" in ft:
                problems.append(f"{f.name}: fontTable에 embed 잔존")
            rels = "word/_rels/fontTable.xml.rels"
            if rels in names and "fonts/" in zc.read(rels).decode("utf-8"):
                problems.append(f"{f.name}: fontTable rels에 fonts/ 잔존")
            if "/word/fonts/" in zc.read("[Content_Types].xml").decode("utf-8"):
                problems.append(f"{f.name}: Content_Types에 글꼴 Override 잔존")
            try:
                if doc_text(cur) != doc_text(orig):
                    problems.append(f"{f.name}: 본문 텍스트가 원본과 다름")
            except Exception as e:  # noqa: BLE001
                problems.append(f"{f.name}: python-docx 열기 실패 {e}")
            if len(cur) > 200 * 1024:
                problems.append(f"{f.name}: 여전히 {len(cur)//1024}KB")
    if checked != 12:
        problems.append(f"검사한 docx 수 {checked} (기대 12)")
    if problems:
        for p in problems:
            print("FAIL:", p)
        return 1
    print(f"checked {checked} docx")
    print("DOCX SLIM OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
