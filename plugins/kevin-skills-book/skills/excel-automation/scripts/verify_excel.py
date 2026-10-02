#!/usr/bin/env python3
"""엑셀 결과 품질 게이트 — 사용자에게 전달하기 전에 반드시 통과시킨다.

공통(모든 모드)
- input-changed: 원본 파일이 작업 전과 다르다(정리·취합은 report.json, 분석은 snapshot의 해시와 비교)
- formula-error: 계산된 값 중 #N/A·#REF!·#VALUE!·#DIV/0!·#NAME? 등이 있다
- forbidden-function: XLOOKUP·FILTER·UNIQUE 등 Excel 2019 이하·LibreOffice에서 깨지는 함수를 썼다

clean(정리) — clean_data.py의 <결과>.report.json과 대조
- row-count: 입력 행 수 = 정리 후 행 수 + 삭제 행 수가 아니거나, 결과 시트 행 수가 기록과 다르다
- bad-dedup: '완전 중복'으로 지운 행이 실제로는 다른 값을 가진다
- unlogged-change: 변경내역에 없는 셀 변경이 있다(조용한 변경)
- candidate-modified: '확인 필요 후보'로 보고한 셀을 고쳤다(사람이 판단할 값)
- blank-filled: 원본 빈 칸을 임의로 채웠다
- missing-report-sheet: 변경리포트·변경내역 시트가 없다

consolidate(취합)
- original-sheet-missing / original-sheet-changed: 원본 탭이 빠졌거나 내용이 바뀌었다
- missing-keys / extra-keys / duplicate-keys: 통합 시트 키가 원본 탭 키의 합집합과 다르다
- values-pasted: 통합 시트 값이 수식이 아니라 붙여 넣은 값이다(원본을 고쳐도 반영되지 않음)
- lookup-mismatch: INDEX/MATCH 수식의 계산값이 원본 탭을 직접 찾아본 값과 다르다
- wrong-column: 수식이 report의 열 대응과 다른 원본 열을 참조한다
- not-recalculated: 수식 값이 비어 있다(재계산 엔진 없음 → --allow-uncached로 경고 처리 가능)

analysis(분석)
- no-native-chart: 엑셀 기본 차트가 없다(그림 차트만 있거나 차트 없음)
- no-insights: '인사이트' 시트가 없거나 비었다
- insight-number-provenance: 인사이트 문장 속 숫자가 통합 문서의 어떤 셀 값과도 맞지 않는다

사용법
  python verify_excel.py snapshot <원본.xlsx> --out snap.json            # 작업 전에 원본 지문 저장
  python verify_excel.py clean <원본.xlsx> <결과.xlsx> [--report r.json]
  python verify_excel.py consolidate <원본.xlsx> <결과.xlsx> [--report r.json] [--sheet 통합관리 --key 상품코드] [--allow-uncached]
  python verify_excel.py analysis <원본.xlsx> <결과.xlsx> [--snapshot snap.json]
  공통 옵션: --json
종료 코드: 오류 0개면 0. 마지막 줄: VERIFY OK 또는 VERIFY FAILED
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

ERRORS = {"#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}
FORBIDDEN = re.compile(r"\b(XLOOKUP|XMATCH|FILTER|UNIQUE|SORT|SORTBY|SEQUENCE|LET|LAMBDA|RANDARRAY|VSTACK|HSTACK|TAKE|DROP|CHOOSECOLS|TEXTSPLIT)\s*\(", re.I)
INDEX_MATCH = re.compile(
    r"^=IFERROR\(INDEX\((?P<sh>'(?:[^']|'')+'|[^!]+)!\$?(?P<vc>[A-Z]+):\$?[A-Z]+,MATCH\(\$?A\$?(?P<row>\d+),"
    r"(?P<sh2>'(?:[^']|'')+'|[^!]+)!\$?(?P<kc>[A-Z]+):\$?[A-Z]+,0\)\),(?P<miss>\"[^\"]*\")\)$", re.I)

BLANK_FALLBACK = re.compile(r',\s*""\s*\)\s*$')

# ───── 숫자 대조(doc-automation verify_deck과 같은 규칙, 날짜·월 표기 추가) ─────
NUM = re.compile(r"(?<![A-Za-z0-9.])[-+−]?\d{1,3}(?:,\d{3})+(?:\.\d+)?(?![0-9])|(?<![A-Za-z0-9.,])[-+−]?\d+(?:\.\d+)?")
DATE_PATTERNS = [
    re.compile(r"\b(19|20)\d{2}[-./]\s?\d{1,2}[-./]\s?\d{1,2}\.?"),
    re.compile(r"\b(19|20)\d{2}[-./]\d{1,2}\b"),
    re.compile(r"\b(19|20)\d{2}\s?년\s?\d{1,2}\s?월(\s?\d{1,2}\s?일)?"),
    re.compile(r"\d{1,2}\s?월\s?\d{1,2}\s?일"),
    re.compile(r"\d{1,2}\s?월"),
    re.compile(r"\b\d{1,2}/\d{1,2}\b"),
    re.compile(r"\bQ[1-4]\b|\b[1-4]Q\b|[1-4]분기"),
]
UNIT = re.compile(r"\s?(조|억|천만|백만|만|천)")
UNIT_MULT = {"조": 1e12, "억": 1e8, "천만": 1e7, "백만": 1e6, "만": 1e4, "천": 1e3}


def sha256(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def strip_dates(text: str) -> str:
    for p in DATE_PATTERNS:
        text = p.sub(" ", text)
    return text


def to_value(tok: str) -> float | None:
    try:
        return float(tok.replace(",", "").replace("−", "-").replace("+", ""))
    except ValueError:
        return None


def decimals(tok: str) -> int:
    t = tok.replace(",", "")
    return len(t.split(".")[1]) if "." in t else 0


def numbers_with_units(text: str) -> list[tuple[str, float]]:
    t = strip_dates(text)
    out = []
    for m in NUM.finditer(t):
        u = UNIT.match(t, m.end())
        out.append((m.group(0), UNIT_MULT[u.group(1)] if u else 1.0))
    return out


def exempt(tok: str) -> bool:
    v = to_value(tok)
    if v is None:
        return True
    if decimals(tok) == 0 and abs(v) <= 9:
        return True
    return decimals(tok) == 0 and 1990 <= v <= 2100 and "," not in tok


def has_number(values: list[float], exact: set, tok: str, mult: float) -> bool:
    v = to_value(tok)
    if v is None:
        return True
    if round(abs(v), 6) in exact:
        return True
    d = decimals(tok)
    for s in values:
        if round(abs(s), d) == round(abs(v), d):
            return True
        if mult != 1.0 and round(abs(s) / mult, d) == round(abs(v), d):
            return True
        if round(abs(s) * 100, d) == round(abs(v), d):  # 0.519 ↔ 51.9(%)
            return True
    return False


# ───── 공통 도구 ─────
def norm(v):
    if isinstance(v, str):
        s = v.strip()
        return s.replace(" ", "").upper() if s.startswith("=") else s
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if hasattr(v, "isoformat"):
        return v.isoformat()[:19]
    return v


def header_map(ws) -> tuple[int, dict[str, int]]:
    best, best_n = 1, -1
    for r in range(1, min(ws.max_row, 10) + 1):
        n = sum(isinstance(ws.cell(r, c).value, str) for c in range(1, ws.max_column + 1))
        if n > best_n:
            best, best_n = r, n
    return best, {str(ws.cell(best, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(best, c).value is not None}


def common_checks(out: Path, issues: list, warns: list) -> tuple:
    wf = load_workbook(out)
    wv = load_workbook(out, data_only=True)
    uncached, cached, blank_ok = 0, 0, 0
    for ws in wf.worksheets:
        vs = wv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    if FORBIDDEN.search(v):
                        issues.append({"check": "forbidden-function", "where": f"{ws.title}!{c.coordinate}", "detail": v[:120]})
                    cv = vs[c.coordinate].value
                    if cv is None:
                        # Excel은 빈 문자열 결과(IFERROR(...,""))를 값 없이 저장한다
                        if BLANK_FALLBACK.search(v):
                            blank_ok += 1
                        else:
                            uncached += 1
                        continue
                    cached += 1
                    if isinstance(cv, str) and cv.strip() in ERRORS:
                        issues.append({"check": "formula-error", "where": f"{ws.title}!{c.coordinate}", "detail": cv})
                elif isinstance(v, str) and v.strip() in ERRORS:
                    issues.append({"check": "formula-error", "where": f"{ws.title}!{c.coordinate}", "detail": v})
    if cached == 0:
        uncached += blank_ok  # 한 번도 계산되지 않은 파일
    return wf, wv, uncached


def check_input(src: Path, expected_sha: str | None, issues: list, warns: list):
    if not expected_sha:
        warns.append({"check": "no-snapshot", "detail": "원본 지문(report 또는 snapshot)이 없어 원본 변경 여부를 확인하지 못했다"})
        return
    if sha256(src) != expected_sha:
        issues.append({"check": "input-changed", "where": str(src), "detail": "원본 파일이 작업 전과 다르다"})


def load_report(out: Path, given: str | None) -> dict | None:
    p = Path(given) if given else Path(str(out) + ".report.json")
    return json.loads(p.read_text(encoding="utf-8")) if p.is_file() else None


# ───── clean ─────
def verify_clean(src: Path, out: Path, rep: dict | None, issues: list, warns: list):
    if rep is None:
        issues.append({"check": "missing-report", "detail": "clean_data.py가 만든 <결과>.report.json이 없다 — 변경 기록 없이 정리하면 검증할 수 없다"})
        return
    check_input(src, rep.get("input_sha256_before"), issues, warns)
    common_checks(out, issues, warns)
    owb = load_workbook(out, data_only=True)
    for name in ("정리완료", "변경리포트", "변경내역"):
        if name not in owb.sheetnames:
            issues.append({"check": "missing-report-sheet", "detail": f"'{name}' 시트가 없다"})
    if "정리완료" not in owb.sheetnames:
        return
    iws = load_workbook(src, data_only=True)[rep["sheet"]]
    header = rep["header"]
    hrow = rep["header_row"]
    inp = {r: [iws.cell(r, j + 1).value for j in range(len(header))] for r in range(hrow + 1, iws.max_row + 1)}
    inp = {r: v for r, v in inp.items() if any(x is not None and str(x).strip() != "" for x in v)}
    ows = owb["정리완료"]
    outrows = [list(r)[:len(header)] for r in ows.iter_rows(min_row=2, values_only=True)
               if any(x is not None and str(x).strip() != "" for x in r)]
    kept = rep["kept_src_rows"]
    if not (len(outrows) == rep["rows_out"] == len(kept)) or rep["rows_in"] != len(inp) or rep["rows_in"] != rep["rows_out"] + len(rep["removed"]):
        issues.append({"check": "row-count", "detail": f"입력 {len(inp)}행(기록 {rep['rows_in']}), 결과 {len(outrows)}행(기록 {rep['rows_out']}), "
                                                      f"삭제 기록 {len(rep['removed'])}행"})
    seq = set(rep.get("sequence_columns", []))
    cmp_idx = [j for j, n in enumerate(header) if n not in seq]
    for r in rep["removed"]:
        a, b = inp.get(r["src_row"]), inp.get(r["duplicate_of"])
        if a is None or b is None or [norm(a[j]) for j in cmp_idx] != [norm(b[j]) for j in cmp_idx]:
            issues.append({"check": "bad-dedup", "where": f"{r['src_row']}행", "detail": f"{r['duplicate_of']}행과 값이 다른데 중복으로 삭제했다"})
    logged = {(c["src_row"], c["column"]): c for c in rep["changes"]}
    cand = {(c["src_row"], c["column"]) for c in rep["candidates"] if isinstance(c["src_row"], int)}
    for i, src_row in enumerate(kept[:len(outrows)]):
        before = inp.get(src_row)
        if before is None:
            issues.append({"check": "row-count", "where": f"{src_row}행", "detail": "기록된 원본 행이 원본에 없다"})
            continue
        for j, name in enumerate(header):
            b, a = before[j], outrows[i][j] if j < len(outrows[i]) else None
            blank_b = b is None or str(b).strip() == ""
            blank_a = a is None or str(a).strip() == ""
            if blank_b and not blank_a:
                issues.append({"check": "blank-filled", "where": f"{src_row}행 {name}", "detail": f"빈 칸 → {a}"})
                continue
            if norm(b) == norm(a) or (blank_a and blank_b):
                continue
            if (src_row, name) in cand:
                issues.append({"check": "candidate-modified", "where": f"{src_row}행 {name}", "detail": f"{b} → {a} (확인 필요 후보를 고쳤다)"})
                continue
            log = logged.get((src_row, name))
            if log is None or str(log["after"]) != str(a):
                issues.append({"check": "unlogged-change", "where": f"{src_row}행 {name}", "detail": f"{b} → {a} (변경내역에 없음)"})


# ───── consolidate ─────
def unq(sh: str) -> str:
    return sh[1:-1].replace("''", "'") if sh.startswith("'") else sh


def verify_consolidate(src: Path, out: Path, rep: dict | None, sheet: str | None, key: str | None, allow_uncached: bool,
                       issues: list, warns: list):
    if rep:
        check_input(src, rep.get("input_sha256_before"), issues, warns)
        sheet = sheet or rep.get("sheet")
        key = key or rep.get("key")
    else:
        warns.append({"check": "no-report", "detail": "report.json 없음 — --sheet·--key 기준으로만 검사"})
    if not sheet or not key:
        issues.append({"check": "missing-args", "detail": "통합 시트 이름과 키 열(--sheet, --key)이 필요하다"})
        return
    wf, wv, uncached = common_checks(out, issues, warns)
    iwb = load_workbook(src)
    iwv = load_workbook(src, data_only=True)
    for ws in iwb.worksheets:
        if ws.title not in wf.sheetnames:
            issues.append({"check": "original-sheet-missing", "detail": f"원본 탭 '{ws.title}'이 결과에 없다"})
            continue
        o = wf[ws.title]
        diffs = [c.coordinate for row in ws.iter_rows() for c in row if norm(c.value) != norm(o[c.coordinate].value)]
        extra = o.max_row > ws.max_row and any(o.cell(r, c).value is not None for r in range(ws.max_row + 1, o.max_row + 1)
                                               for c in range(1, o.max_column + 1))
        if diffs or extra:
            issues.append({"check": "original-sheet-changed", "where": ws.title, "detail": f"바뀐 셀 {len(diffs)}개 예: {diffs[:5]}"})
    if sheet not in wf.sheetnames:
        issues.append({"check": "missing-sheet", "detail": f"통합 시트 '{sheet}'가 없다"})
        return
    # 원본 탭의 키 합집합과 조회표
    lookup: dict[str, dict] = {}
    union: set = set()
    used = [s["sheet"] for s in rep["spec"]["sheets"]] if rep else [ws.title for ws in iwv.worksheets if key in header_map(ws)[1]]
    for name in used:
        if name not in iwv.sheetnames:
            continue
        ws = iwv[name]
        hr, hm = header_map(ws)
        if key not in hm:
            continue
        kc = hm[key]
        letters = {}
        for r in range(hr + 1, ws.max_row + 1):
            k = ws.cell(r, kc).value
            if k is None or str(k).strip() == "":
                continue
            union.add(k)
            letters.setdefault(k, r)
        lookup[name] = {"ws": ws, "first_row": letters, "hm": hm, "kc": kc}
    t, tv = wf[sheet], wv[sheet]
    _, thm = header_map(t)
    kcol = thm.get(key, 1)
    keys = [t.cell(r, kcol).value for r in range(2, t.max_row + 1) if t.cell(r, kcol).value not in (None, "")]
    if len(keys) != len(set(keys)):
        issues.append({"check": "duplicate-keys", "detail": f"통합 시트 키 중복 {len(keys) - len(set(keys))}개"})
    miss, extra = union - set(keys), set(keys) - union
    if miss:
        issues.append({"check": "missing-keys", "detail": f"원본 탭에 있으나 통합 시트에 없는 키 {len(miss)}개 예: {sorted(map(str, miss))[:5]}"})
    if extra:
        issues.append({"check": "extra-keys", "detail": f"원본 탭에 없는 키 {len(extra)}개 예: {sorted(map(str, extra))[:5]}"})
    col_src = {}
    if rep:
        for s in rep["spec"]["sheets"]:
            for sc, oc in s["columns"].items():
                col_src[oc] = (s["sheet"], sc)
    pasted, checked, mismatched, wrong = 0, 0, 0, 0
    for cname, cj in thm.items():
        if cj == kcol:
            continue
        for r in range(2, t.max_row + 1):
            f = t.cell(r, cj).value
            if f is None:
                continue
            if not (isinstance(f, str) and f.startswith("=")):
                if cname in col_src or not rep:
                    pasted += 1
                continue
            m = INDEX_MATCH.match(f.replace(" ", ""))
            if not m:
                continue
            sh = unq(m.group("sh"))
            if sh not in lookup:
                continue
            L = lookup[sh]
            src_col_name = str(L["ws"].cell(header_map(L["ws"])[0], column_index_from_string(m.group("vc"))).value or "").strip()
            if column_index_from_string(m.group("kc")) != L["kc"]:
                wrong += 1
            if cname in col_src and col_src[cname] != (sh, src_col_name):
                wrong += 1
                if wrong <= 3:
                    issues.append({"check": "wrong-column", "where": f"{sheet}!{t.cell(r, cj).coordinate}",
                                   "detail": f"'{cname}'은 {col_src[cname]}를 가리켜야 하는데 {sh}!{src_col_name}를 참조"})
            k = t.cell(int(m.group("row")), kcol).value
            fr = L["first_row"].get(k)
            exp = L["ws"].cell(fr, column_index_from_string(m.group("vc"))).value if fr else m.group("miss").strip('"')
            exp = "" if exp is None else exp
            got = tv.cell(r, cj).value
            if got is None:
                continue
            checked += 1
            if norm(got) != norm(exp) and not (exp == "" and got in ("", 0)):
                mismatched += 1
                if mismatched <= 3:
                    issues.append({"check": "lookup-mismatch", "where": f"{sheet}!{t.cell(r, cj).coordinate}", "detail": f"계산값 {got} ≠ 원본 {exp}"})
    if pasted:
        issues.append({"check": "values-pasted", "detail": f"수식이 아닌 붙여 넣은 값 {pasted}칸 — 원본 탭을 고쳐도 통합 시트가 따라오지 않는다"})
    if mismatched > 3:
        issues.append({"check": "lookup-mismatch", "detail": f"총 {mismatched}칸 불일치"})
    if uncached:
        (warns if allow_uncached else issues).append({"check": "not-recalculated", "detail": f"계산값이 없는 수식 {uncached}개 — recalc.py로 재계산하거나 Excel에서 열어 저장"})
    warns.append({"check": "lookup-checked", "detail": f"INDEX/MATCH 계산값 대조 {checked}칸, 불일치 {mismatched}"})


# ───── analysis ─────
def verify_analysis(src: Path, out: Path, snap: dict | None, insight_sheet: str, issues: list, warns: list):
    check_input(src, snap.get("sha256") if snap else None, issues, warns)
    wf, wv, uncached = common_checks(out, issues, warns)
    if uncached:
        warns.append({"check": "not-recalculated", "detail": f"계산값이 없는 수식 {uncached}개"})
    iwb = load_workbook(src, read_only=True)
    names = list(iwb.sheetnames)
    iwb.close()
    missing = [s for s in names if s not in wf.sheetnames]
    if missing:
        warns.append({"check": "original-sheets-not-included", "detail": f"결과에 원본 탭이 없다: {missing} (새 파일로 분석만 저장한 경우)"})
    with zipfile.ZipFile(out) as z:
        charts = [n for n in z.namelist() if re.match(r"xl/charts/chart\d+\.xml$", n)]
        images = [n for n in z.namelist() if n.startswith("xl/media/")]
    if not charts:
        issues.append({"check": "no-native-chart", "detail": f"엑셀 기본 차트가 없다(그림 {len(images)}개)"})
    if insight_sheet not in wv.sheetnames:
        issues.append({"check": "no-insights", "detail": f"'{insight_sheet}' 시트가 없다"})
        return
    ws = wv[insight_sheet]
    texts = []
    for row in ws.iter_rows(min_row=1, values_only=True):
        cells = [c for c in row if isinstance(c, str) and len(c) >= 12]
        texts.extend(cells[:1])
    if not texts:
        issues.append({"check": "no-insights", "detail": "인사이트 문장이 없다"})
        return
    if not 3 <= len(texts) <= 5:
        warns.append({"check": "insight-count", "detail": f"인사이트 {len(texts)}개(권장 3~5개)"})
    values = []
    for s in wv.worksheets:
        if s.title == insight_sheet:
            continue
        for row in s.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    values.append(float(v))
    exact = {round(abs(v), 6) for v in values}
    for t in texts:
        for tok, mult in numbers_with_units(t):
            if exempt(tok):
                continue
            if not has_number(values, exact, tok, mult):
                issues.append({"check": "insight-number-provenance", "where": insight_sheet, "detail": f"'{tok}'이(가) 통합 문서 어느 셀에도 없다: {t[:80]}"})


def main(argv: list[str]) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    ap = argparse.ArgumentParser(description="엑셀 결과 품질 게이트")
    ap.add_argument("mode", choices=["snapshot", "clean", "consolidate", "analysis"])
    ap.add_argument("input")
    ap.add_argument("output", nargs="?")
    ap.add_argument("--out")
    ap.add_argument("--report")
    ap.add_argument("--snapshot")
    ap.add_argument("--sheet")
    ap.add_argument("--key")
    ap.add_argument("--insight-sheet", default="인사이트")
    ap.add_argument("--allow-uncached", action="store_true")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args(argv)
    src = Path(a.input)
    if a.mode == "snapshot":
        snap = {"path": str(src.resolve()), "sha256": sha256(src)}
        if a.out:
            Path(a.out).write_text(json.dumps(snap, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"SNAPSHOT {snap['sha256'][:16]} {src.name}")
        return 0
    if not a.output:
        print("오류: 결과 파일 경로가 필요하다")
        return 2
    out = Path(a.output)
    if out.resolve() == src.resolve():
        issues = [{"check": "input-changed", "detail": "결과 파일이 원본과 같은 경로다(원본 덮어쓰기)"}]
        warns: list = []
    else:
        issues, warns = [], []
        if a.mode == "clean":
            verify_clean(src, out, load_report(out, a.report), issues, warns)
        elif a.mode == "consolidate":
            verify_consolidate(src, out, load_report(out, a.report), a.sheet, a.key, a.allow_uncached, issues, warns)
        else:
            snap = json.loads(Path(a.snapshot).read_text(encoding="utf-8")) if a.snapshot else None
            verify_analysis(src, out, snap, a.insight_sheet, issues, warns)
    if a.json:
        print(json.dumps({"mode": a.mode, "errors": issues, "warnings": warns}, ensure_ascii=False, indent=2))
    else:
        for i in issues:
            print(f"ERROR [{i['check']}] {i.get('where', '')} {i['detail']}")
        for w in warns:
            print(f"WARN  [{w['check']}] {w.get('where', '')} {w['detail']}")
    print(f"VERIFY {'OK' if not issues else 'FAILED'} mode={a.mode} errors={len(issues)} warnings={len(warns)}")
    return 0 if not issues else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
