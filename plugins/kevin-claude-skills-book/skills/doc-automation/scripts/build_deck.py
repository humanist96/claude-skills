#!/usr/bin/env python3
"""outline.json(보고 스토리라인) → 편집 가능한 .pptx

내용은 Claude가 outline.json에 설계하고, 이 스크립트는 일관된 디자인으로 그리기만 한다.
차트·표는 그림이 아니라 PowerPoint 네이티브 개체로 만들어 받는 사람이 직접 고칠 수 있다.

슬라이드 유형: title, agenda, section, bullets, kpi, chart, table, two_column, closing
  (스키마: references/outline-schema.md)

회사 템플릿(--template 또는 outline.meta.template)
- 슬라이드 크기·테마 색·글꼴을 따른다
- 표지는 템플릿의 표지 레이아웃, 나머지는 '제목만' 레이아웃(없으면 제목+내용, 없으면 빈 화면)을 쓴다
- 제목은 템플릿의 제목 칸에 넣고, 쓰지 않은 빈 칸은 지운다
- 템플릿에 들어 있던 예시 슬라이드는 지운다

출력: <out>.pptx 와 같은 폴더의 build_report.json (슬라이드별 경고: 글자 넘침 위험, 항목 과다 등)

사용법
  python build_deck.py outline.json --out <출력폴더>/report.pptx [--template 회사.pptx] [--sources <sources 폴더>]
"""
from __future__ import annotations

import argparse
import json
import math
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent / "_vendor"))

from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

from style_config import StyleConfig, default_style, from_template_info  # noqa: E402

DEFAULT_FONT = "맑은 고딕"
SLIDE_TYPES = {"title", "agenda", "section", "bullets", "kpi", "chart", "table", "two_column", "closing"}
CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
}
PALETTE = ["#2F6FDB", "#F28C28", "#3BA272", "#9A60B4", "#E15759", "#5B6B7F"]
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}


# ──────────────────────────── 글꼴·텍스트 유틸 ────────────────────────────
def _hex(c: str) -> RGBColor:
    c = c.lstrip("#")
    return RGBColor(int(c[:2], 16), int(c[2:4], 16), int(c[4:6], 16))


KOREAN_FONTS = {"맑은 고딕", "malgun gothic", "나눔고딕", "nanumgothic", "pretendard", "noto sans kr", "noto sans cjk kr",
                "굴림", "gulim", "돋움", "dotum", "바탕", "batang", "apple sd gothic neo", "applegothic", "spoqa han sans",
                "spoqa han sans neo", "본고딕", "source han sans kr", "나눔스퀘어", "nanumsquare", "에스코어 드림", "s-core dream"}


def font_pair(name: str) -> tuple[str, str]:
    """(라틴 글꼴, 한글 글꼴). 템플릿 글꼴이 Calibri처럼 한글이 없으면 한글은 기본 한글 글꼴로 둔다."""
    has_ko = any("가" <= ch <= "힣" for ch in name) or name.strip().lower() in KOREAN_FONTS
    return (name, name if has_ko else DEFAULT_FONT)


def _set_font(run, family, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None):
    latin, ea_font = family if isinstance(family, tuple) else (family, family)
    f = run.font
    f.name = latin
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color is not None:
        f.color.rgb = color
    # 한글은 East Asian 글꼴로 지정해야 Office가 대체 글꼴로 바꾸지 않는다
    rpr = run._r.get_or_add_rPr()
    # 언어 태그가 없거나 en-US이면 PowerPoint가 한글을 글자 단위로 잘라 "반품/과"처럼 줄바꿈한다.
    # ko-KR로 지정하면 단어(띄어쓰기) 단위로 줄바꿈한다(PowerPoint 렌더링으로 확인).
    if any("가" <= ch <= "힣" for ch in run.text or ""):
        rpr.set("lang", "ko-KR")
        rpr.set("altLang", "en-US")
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", ea_font)


def _text_units(text: str) -> float:
    """대략적 글자 폭(전각=1, 반각=0.55). 넘침 추정용."""
    return sum(1.0 if ord(ch) > 0x2E7F else 0.55 for ch in text)


def _lines_needed(paragraphs: list[str], width_in: float, size_pt: float) -> int:
    per_line = max(1.0, width_in * 72 / size_pt)
    return sum(max(1, math.ceil(_text_units(p) / per_line)) for p in paragraphs)


def _fit_size(paragraphs: list[str], width_in: float, height_in: float, start: float, minimum: float = 11,
              spacing: float = 1.35) -> tuple[float, bool]:
    """박스에 들어가는 가장 큰 글자 크기와 넘침 여부."""
    size = start
    while size >= minimum:
        if _lines_needed(paragraphs, width_in, size) * size * spacing / 72 <= height_in:
            return size, False
        size -= 1
    return minimum, True


class Deck:
    def __init__(self, outline: dict, template: str | None, sources: dict[str, str]):
        self.o = outline
        self.meta = outline.get("meta", {})
        self.sources = sources
        self.warnings: list[dict] = []
        if template:
            from template_analyzer import analyze_template
            self.prs = Presentation(template)
            info = analyze_template(template)
            self.style: StyleConfig = from_template_info(info)
            self._remove_slides()
            self.templated = True
        else:
            self.prs = Presentation()
            self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
            self.style = default_style()
            self.templated = False
        if self.meta.get("accent"):
            self.style.colors["accent"] = self.meta["accent"]
        self.font = font_pair(self.meta.get("font_family") or self.style.font_family or DEFAULT_FONT)
        self.W = self.prs.slide_width / 914400
        self.H = self.prs.slide_height / 914400
        self.margin = 0.5 if self.W < 11 else 0.6
        self.page = 0

    # ── 템플릿 처리
    def _remove_slides(self):
        lst = self.prs.slides._sldIdLst
        for sld in list(lst):
            self.prs.part.drop_rel(sld.get(qn("r:id")))
            lst.remove(sld)

    def _layout(self, kind: str):
        layouts = self.prs.slide_layouts
        if not self.templated:
            return layouts[6]  # 빈 화면
        def has(layout, types):
            return any(ph.placeholder_format.type in types for ph in layout.placeholders)
        def only_title(layout):
            types = {ph.placeholder_format.type for ph in layout.placeholders}
            return bool(types & TITLE_TYPES) and not (types & {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE})
        if kind == "title":
            for l in layouts:
                if has(l, {PP_PLACEHOLDER.CENTER_TITLE}) or (has(l, TITLE_TYPES) and has(l, {PP_PLACEHOLDER.SUBTITLE})):
                    return l
        for l in layouts:
            if only_title(l):
                return l
        for l in layouts:
            if has(l, TITLE_TYPES):
                return l
        return layouts[len(layouts) - 1]

    def _new_slide(self, kind: str):
        slide = self.prs.slides.add_slide(self._layout(kind))
        self.page += 1
        return slide

    def _title_ph(self, slide):
        for ph in slide.placeholders:
            if ph.placeholder_format.type in TITLE_TYPES:
                return ph
        return None

    def _drop_empty_placeholders(self, slide):
        for ph in list(slide.placeholders):
            if ph.placeholder_format.type in (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER):
                continue
            if ph.has_text_frame and not ph.text_frame.text.strip():
                ph._element.getparent().remove(ph._element)

    # ── 공통 요소
    def _box(self, slide, x, y, w, h, fill: str | None = None, line: str | None = None, shape=MSO_SHAPE.RECTANGLE):
        s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill:
            s.fill.solid()
            s.fill.fore_color.rgb = _hex(fill)
        else:
            s.fill.background()
        if line:
            s.line.color.rgb = _hex(line)
            s.line.width = Pt(0.75)
        else:
            s.line.fill.background()
        s.shadow.inherit = False
        return s

    def _text(self, slide, x, y, w, h, paras: list, size: float, bold=False, color="body_text",
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fit=True, bullets=False, name=None, slide_no=None):
        """paras: [str | {"text", "level"}]"""
        items = [p if isinstance(p, dict) else {"text": str(p), "level": 0} for p in paras]
        texts = [("    " * it.get("level", 0)) + ("• " if bullets else "") + it["text"] for it in items]
        overflow = False
        if fit:
            size, overflow = _fit_size(texts, w - 0.2, h - 0.1, size)
            if overflow:
                self.warnings.append({"slide": slide_no or self.page, "issue": "글자 넘침 위험",
                                      "detail": f"{name or '텍스트'}: 최소 크기에서도 상자를 넘을 수 있음 — 내용을 줄일 것"})
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.08)
        col = self.style.rgb(color) if not color.startswith("#") else _hex(color)
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_after = Pt(size * 0.45)
            lvl = it.get("level", 0)
            prefix = ("• " if lvl == 0 else "– ") if bullets else ""
            r = p.add_run()
            r.text = prefix + it["text"]
            _set_font(r, self.font, size - (2 if lvl else 0), bold, col)
            if lvl:
                p.level = min(lvl, 4)
        return tb

    def _headline(self, slide, text: str):
        if len(text) > 45:
            self.warnings.append({"slide": self.page, "issue": "헤드라인이 김", "detail": f"{len(text)}자 — 45자 이내 권장"})
        ph = self._title_ph(slide)
        if ph is not None:
            # 템플릿 제목 크기(보통 40~44pt)를 그대로 두면 두 줄 제목이 본문과 겹친다 → 칸에 맞는 크기를 명시한다
            pw, phh = ph.width / 914400 - 0.2, ph.height / 914400 - 0.1
            size, over = _fit_size([text], pw, phh, 32 if self.W >= 11 else 30, minimum=18, spacing=1.2)
            if over:
                self.warnings.append({"slide": self.page, "issue": "헤드라인 넘침 위험",
                                      "detail": "템플릿 제목 칸에 18pt로도 들어가지 않음 — 헤드라인을 줄일 것"})
            ph.text_frame.text = ""
            ph.text_frame.word_wrap = True
            r = ph.text_frame.paragraphs[0].add_run()
            r.text = text
            _set_font(r, self.font, size, True, self.style.rgb("title_text"))
            top = ph.top / 914400 + ph.height / 914400 + 0.15
            return max(top, 1.2)
        self._text(slide, self.margin, 0.35, self.W - 2 * self.margin, 0.9, [text], 26, bold=True,
                   color="title_text", anchor=MSO_ANCHOR.MIDDLE, name="Headline")
        return 1.35

    def _footer(self, slide, sources: list[str]):
        y = self.H - 0.45
        if sources:
            label = "출처: " + "; ".join(self._source_label(s) for s in sources)
            self._text(slide, self.margin, y, self.W - 2 * self.margin - 0.8, 0.35, [label], 9,
                       color="muted", fit=False, name="Sources")
        self._text(slide, self.W - self.margin - 0.7, y, 0.7, 0.35, [str(self.page)], 9, color="muted",
                   align=PP_ALIGN.RIGHT, fit=False, name="PageNo")

    def _source_label(self, ref: str) -> str:
        sid, _, loc = ref.partition(" ")
        name = self.sources.get(sid)
        return f"{name} {loc}".strip() if name else ref

    def _notes(self, slide, s: dict):
        parts = []
        if s.get("notes"):
            parts.append(s["notes"])
        if s.get("sources"):
            parts.append("출처: " + "; ".join(self._source_label(x) for x in s["sources"]))
        if parts:
            slide.notes_slide.notes_text_frame.text = "\n\n".join(parts)

    def _takeaway(self, slide, text: str, y: float):
        h = 0.62
        self._box(slide, self.margin, y, self.W - 2 * self.margin, h, fill=self.style.colors["card_fill"],
                  line=self.style.colors["accent"])
        self._text(slide, self.margin + 0.15, y, self.W - 2 * self.margin - 0.3, h, [text], 15, bold=True,
                   color="accent_dark", anchor=MSO_ANCHOR.MIDDLE, name="Takeaway")

    def _content_box(self, top: float, takeaway: bool) -> tuple[float, float, float, float]:
        bottom = self.H - 0.6 - (0.8 if takeaway else 0)
        return self.margin, top, self.W - 2 * self.margin, bottom - top

    # ── 슬라이드 유형
    def title(self, s: dict):
        slide = self._new_slide("title")
        title = s.get("title") or self.meta.get("title", "")
        subtitle = s.get("subtitle") or self.meta.get("subtitle", "")
        sub_line = " · ".join(x for x in [self.meta.get("author"), self.meta.get("date") or date.today().isoformat()] if x)
        ph = self._title_ph(slide)
        if self.templated and ph is not None:
            ph.text_frame.text = ""
            ph.text_frame.word_wrap = True
            size, _ = _fit_size([title], ph.width / 914400 - 0.2, ph.height / 914400 - 0.1, 40, minimum=20, spacing=1.2)
            r = ph.text_frame.paragraphs[0].add_run(); r.text = title
            _set_font(r, self.font, size, True, None)
            for p in slide.placeholders:
                if p.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE:
                    p.text_frame.text = ""
                    r = p.text_frame.paragraphs[0].add_run(); r.text = "\n".join(x for x in [subtitle, sub_line] if x)
                    _set_font(r, self.font)
            self._drop_empty_placeholders(slide)
        else:
            self._box(slide, 0, 0, self.W, self.H, fill=self.style.colors["accent_dark"])
            self._text(slide, self.margin + 0.3, self.H * 0.30, self.W - 2 * self.margin - 0.6, 1.6, [title], 38,
                       bold=True, color="#FFFFFF", anchor=MSO_ANCHOR.BOTTOM, name="Headline")
            if subtitle:
                self._text(slide, self.margin + 0.3, self.H * 0.30 + 1.7, self.W - 2 * self.margin - 0.6, 0.8,
                           [subtitle], 18, color="#DDE6F5", name="Subtitle")
            self._text(slide, self.margin + 0.3, self.H - 1.1, self.W - 2 * self.margin - 0.6, 0.5, [sub_line], 12,
                       color="#C9D3E3", fit=False, name="Byline")
        logo = self.meta.get("logo")
        if logo and Path(logo).is_file():
            slide.shapes.add_picture(str(logo), Inches(self.W - self.margin - 1.6), Inches(0.4), height=Inches(0.8))
        self._notes(slide, s)

    def section(self, s: dict):
        slide = self._new_slide("section")
        self._box(slide, 0, 0, self.W, self.H, fill=self.style.colors["accent"])
        self._text(slide, self.margin + 0.3, self.H * 0.38, self.W - 2 * self.margin, 1.2, [s["headline"]], 32,
                   bold=True, color="#FFFFFF", name="Headline")
        if s.get("subtitle"):
            self._text(slide, self.margin + 0.3, self.H * 0.38 + 1.2, self.W - 2 * self.margin, 0.8,
                       [s["subtitle"]], 16, color="#FFFFFF", name="Subtitle")
        if self.templated:
            self._drop_empty_placeholders(slide)
        self._notes(slide, s)

    def agenda(self, s: dict):
        slide = self._new_slide("agenda")
        top = self._headline(slide, s.get("headline", "목차"))
        x, y, w, h = self._content_box(top, False)
        items = s.get("items", [])
        self._text(slide, x + 0.3, y + 0.1, w - 0.6, h, [f"{i}. {t}" for i, t in enumerate(items, 1)], 22,
                   color="body_text", name="Agenda")
        self._finish(slide, s)

    def bullets(self, s: dict):
        slide = self._new_slide("bullets")
        top = self._headline(slide, s["headline"])
        x, y, w, h = self._content_box(top, bool(s.get("takeaway")))
        items = s.get("bullets", [])
        if len([b for b in items if not isinstance(b, dict) or b.get("level", 0) == 0]) > 6:
            self.warnings.append({"slide": self.page, "issue": "항목 과다", "detail": "상위 항목 6개 이하 권장"})
        self._text(slide, x, y, w, h, items, 20, bullets=True, name="Body")
        self._finish(slide, s, y + h + 0.15)

    def kpi(self, s: dict):
        slide = self._new_slide("kpi")
        top = self._headline(slide, s["headline"])
        x, y, w, h = self._content_box(top, bool(s.get("takeaway")))
        kpis = s.get("kpis", [])[:4]
        n = max(1, len(kpis))
        gap = 0.3
        cw = (w - gap * (n - 1)) / n
        ch = min(h, 2.6)
        cy = y + 0.25
        for i, k in enumerate(kpis):
            cx = x + i * (cw + gap)
            self._box(slide, cx, cy, cw, ch, fill=self.style.colors["card_fill"], line=self.style.colors["border"])
            self._text(slide, cx + 0.15, cy + 0.15, cw - 0.3, 0.5, [k.get("label", "")], 14, color="muted", name=f"KPI{i+1}Label")
            self._text(slide, cx + 0.15, cy + 0.65, cw - 0.3, 1.0, [str(k.get("value", ""))], 34, bold=True,
                       color="accent_dark", name=f"KPI{i+1}Value")
            delta = str(k.get("delta", "") or "")
            if delta:
                up = delta.strip().startswith(("+", "▲", "↑"))
                down = delta.strip().startswith(("-", "▼", "↓", "−"))
                col = "success" if up else ("warning" if down else "muted")
                self._text(slide, cx + 0.15, cy + 1.65, cw - 0.3, 0.45, [delta], 14, bold=True, color=col, name=f"KPI{i+1}Delta")
            if k.get("note"):
                self._text(slide, cx + 0.15, cy + 2.05, cw - 0.3, 0.5, [k["note"]], 11, color="muted", name=f"KPI{i+1}Note")
        self._finish(slide, s, min(cy + ch + 0.4, y + h + 0.15))

    def chart(self, s: dict):
        slide = self._new_slide("chart")
        top = self._headline(slide, s["headline"])
        x, y, w, h = self._content_box(top, bool(s.get("takeaway")))
        c = s["chart"]
        if c.get("unit"):  # 단위는 차트 오른쪽 위에 가로로 쓴다(세로 축 제목은 한글이 눕혀져 읽기 어렵다)
            self._text(slide, x + w - 3.0, y - 0.05, 3.0, 0.35, [f"(단위: {c['unit']})"], 11, color="muted",
                       align=PP_ALIGN.RIGHT, fit=False, name="ChartUnit")
            y, h = y + 0.3, h - 0.3
        kind = c.get("kind", "column")
        if kind not in CHART_TYPES:
            raise ValueError(f"{self.page}번 슬라이드: 지원하지 않는 차트 종류 {kind} (가능: {', '.join(CHART_TYPES)})")
        data = CategoryChartData()
        data.categories = [str(v) for v in c["categories"]]
        for ser in c["series"]:
            if len(ser["values"]) != len(c["categories"]):
                raise ValueError(f"{self.page}번 슬라이드: '{ser['name']}' 값 개수가 categories와 다름")
            data.add_series(ser["name"], [float(v) if v is not None else None for v in ser["values"]],
                            number_format=c.get("number_format", "#,##0"))
        gf = slide.shapes.add_chart(CHART_TYPES[kind], Inches(x), Inches(y), Inches(w), Inches(h), data)
        gf.name = "Chart"
        chart = gf.chart
        chart.font.size = Pt(12)
        chart.font.name = self.font[1]  # 차트 라벨은 대부분 한글이므로 한글 글꼴
        multi = len(c["series"]) > 1 or kind in ("pie", "doughnut")
        chart.has_legend = multi
        if multi:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        # 단일 계열이면 PowerPoint가 계열 이름을 제목으로 자동 표시한다. 헤드라인이 있으므로 끈다.
        chart.has_title = bool(c.get("title"))
        if c.get("title"):
            chart.chart_title.text_frame.text = c["title"]
        plot = chart.plots[0]
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = c.get("number_format", "#,##0")
        dl.number_format_is_linked = False
        dl.font.size = Pt(11)
        if kind in ("pie", "doughnut"):
            dl.position = XL_LABEL_POSITION.OUTSIDE_END if kind == "pie" else XL_LABEL_POSITION.CENTER
            for i, pt in enumerate(plot.series[0].points):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = _hex(PALETTE[i % len(PALETTE)])
        else:
            for i, ser in enumerate(plot.series):
                col = _hex(self.style.colors["accent"] if i == 0 else PALETTE[i % len(PALETTE)])
                if kind == "line":
                    ser.format.line.color.rgb = col
                    ser.format.line.width = Pt(2.25)
                    ser.smooth = False
                else:
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = col
            try:
                chart.value_axis.has_major_gridlines = True
                chart.value_axis.major_gridlines.format.line.color.rgb = _hex(self.style.colors["border"])
                chart.value_axis.tick_labels.font.size = Pt(11)
                # 축 눈금은 값이 10 이상이면 정수로(1,200.0 같은 불필요한 소수점 제거)
                vals = [abs(float(v)) for ser in c["series"] for v in ser["values"] if v is not None]
                if vals and max(vals) >= 10:
                    chart.value_axis.tick_labels.number_format = "#,##0"
                    chart.value_axis.tick_labels.number_format_is_linked = False
                chart.category_axis.tick_labels.font.size = Pt(11)
            except (AttributeError, ValueError):
                pass
        self._finish(slide, s, y + h + 0.15)

    def table(self, s: dict):
        slide = self._new_slide("table")
        top = self._headline(slide, s["headline"])
        x, y, w, h = self._content_box(top, bool(s.get("takeaway")))
        t = s["table"]
        cols, rows = t["columns"], t["rows"]
        if len(rows) > 12:
            self.warnings.append({"slide": self.page, "issue": "표 행 과다", "detail": f"{len(rows)}행 — 12행 이하로 요약 권장"})
        n_r = len(rows) + 1
        row_h = min(0.5, h / n_r)
        size = 14 if n_r <= 6 else (12 if n_r <= 10 else 10)
        gf = slide.shapes.add_table(n_r, len(cols), Inches(x), Inches(y), Inches(w), Inches(row_h * n_r))
        gf.name = "Table"
        tbl = gf.table
        # 열 너비는 열마다 가장 긴 글자 폭에 비례해 나눈다(균등 분할은 긴 항목명을 줄바꿈시킨다)
        need = [max([_text_units(str(c))] + [_text_units("" if r[j] is None else str(r[j])) for r in rows if j < len(r)])
                for j, c in enumerate(cols)]
        need = [max(n, 3.0) for n in need]
        total = sum(need)
        widths = [max(n / total, 0.08) for n in need]
        scale = sum(widths)
        for j, frac in enumerate(widths):
            tbl.columns[j].width = Emu(int(Inches(w) * frac / scale))
        for j, name in enumerate(cols):
            cell = tbl.cell(0, j)
            cell.fill.solid(); cell.fill.fore_color.rgb = self.style.rgb("accent")
            cell.text_frame.text = ""
            r = cell.text_frame.paragraphs[0].add_run(); r.text = str(name)
            _set_font(r, self.font, size, True, RGBColor(0xFF, 0xFF, 0xFF))
        for i, row in enumerate(rows, 1):
            if len(row) != len(cols):
                raise ValueError(f"{self.page}번 슬라이드 표 {i}행: 열 개수가 columns와 다름")
            for j, v in enumerate(row):
                cell = tbl.cell(i, j)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _hex("#FFFFFF" if i % 2 else "#F3F6FA")
                cell.text_frame.text = ""
                p = cell.text_frame.paragraphs[0]
                txt = "" if v is None else str(v)
                if txt.replace(",", "").replace(".", "").replace("%", "").replace("-", "").strip().isdigit():
                    p.alignment = PP_ALIGN.RIGHT
                r = p.add_run(); r.text = txt
                _set_font(r, self.font, size, j == 0, self.style.rgb("body_text"))
        self._finish(slide, s, y + h + 0.15)

    def two_column(self, s: dict):
        slide = self._new_slide("two_column")
        top = self._headline(slide, s["headline"])
        x, y, w, h = self._content_box(top, bool(s.get("takeaway")))
        gap = 0.35
        cw = (w - gap) / 2
        # 상자 높이는 두 열 중 내용이 많은 쪽에 맞춘다(빈 상자가 화면 끝까지 늘어나지 않게)
        def body_lines(col):
            items = [b["text"] if isinstance(b, dict) else str(b) for b in col.get("bullets", [])]
            return _lines_needed(["• " + t for t in items] or [""], cw - 0.6, 16)
        need = max(body_lines(s.get("left", {})), body_lines(s.get("right", {})))
        bh = min(h, max(2.0, 0.95 + need * 16 * 1.45 / 72 + 0.3))
        for i, side in enumerate(("left", "right")):
            col = s.get(side, {})
            cx = x + i * (cw + gap)
            self._box(slide, cx, y, cw, bh, fill=self.style.colors["card_fill"], line=self.style.colors["border"])
            self._text(slide, cx + 0.2, y + 0.15, cw - 0.4, 0.55, [col.get("title", "")], 18, bold=True,
                       color="accent_dark", name=f"{side.title()}Title")
            self._text(slide, cx + 0.2, y + 0.75, cw - 0.4, bh - 0.9, col.get("bullets", []), 16, bullets=True,
                       name=f"{side.title()}Body")
        self._finish(slide, s, y + bh + 0.25)

    def closing(self, s: dict):
        slide = self._new_slide("closing")
        top = self._headline(slide, s.get("headline", "다음 단계"))
        x, y, w, h = self._content_box(top, False)
        items = s.get("bullets", [])
        self._text(slide, x, y, w, h - (0.6 if s.get("contact") else 0), items, 20, bullets=True, name="Body")
        if s.get("contact"):
            self._text(slide, x, y + h - 0.5, w, 0.5, [s["contact"]], 12, color="muted", fit=False, name="Contact")
        self._finish(slide, s)

    def _finish(self, slide, s: dict, takeaway_y: float | None = None):
        if s.get("takeaway") and takeaway_y is not None:
            self._takeaway(slide, s["takeaway"], takeaway_y)
        if self.templated:
            self._drop_empty_placeholders(slide)
        self._footer(slide, s.get("sources", []))
        self._notes(slide, s)

    def build(self) -> None:
        slides = self.o.get("slides", [])
        if not slides:
            raise ValueError("outline.slides가 비어 있습니다")
        for i, s in enumerate(slides, 1):
            t = s.get("type")
            if t not in SLIDE_TYPES:
                raise ValueError(f"{i}번 슬라이드: 알 수 없는 type {t!r} (가능: {', '.join(sorted(SLIDE_TYPES))})")
            if t not in ("title", "agenda", "closing") and not s.get("headline"):
                raise ValueError(f"{i}번 슬라이드({t}): headline이 없습니다 — 슬라이드마다 한 문장 메시지가 필요합니다")
            getattr(self, t)(s)


def load_sources(src_dir: str | None) -> dict[str, str]:
    if not src_dir:
        return {}
    p = Path(src_dir) / "sources.json"
    if not p.is_file():
        return {}
    return {e["id"]: e["file"] for e in json.loads(p.read_text(encoding="utf-8")) if e.get("id")}


def main(argv: list[str]) -> int:
    for st in (sys.stdout, sys.stderr):
        try:
            st.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
        except Exception:
            pass
    ap = argparse.ArgumentParser(description="outline.json → pptx")
    ap.add_argument("outline")
    ap.add_argument("--out", required=True)
    ap.add_argument("--template")
    ap.add_argument("--sources")
    a = ap.parse_args(argv)
    outline = json.loads(Path(a.outline).read_text(encoding="utf-8"))
    template = a.template or outline.get("meta", {}).get("template")
    if template and not Path(template).is_file():
        print(f"오류: 템플릿 파일이 없습니다: {template}")
        return 1
    out = Path(a.out)
    if out.suffix.lower() != ".pptx":
        print("오류: --out은 .pptx 파일이어야 합니다")
        return 1
    try:
        deck = Deck(outline, template, load_sources(a.sources))
        deck.build()
    except (ValueError, KeyError) as e:
        print(f"오류: outline 문제 — {e}")
        return 1
    out.parent.mkdir(parents=True, exist_ok=True)
    deck.prs.save(str(out))
    report = {"output": str(out), "slides": deck.page, "template": template, "font": list(deck.font),
              "size_in": [round(deck.W, 3), round(deck.H, 3)], "warnings": deck.warnings}
    (out.parent / "build_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"저장: {out} ({deck.page}장, 경고 {len(deck.warnings)}개)")
    for w in deck.warnings:
        print(f"  경고 {w['slide']}번: {w['issue']} — {w['detail']}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
