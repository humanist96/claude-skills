#!/usr/bin/env python3
"""보고용 PPT 품질 게이트 — 전달 전에 반드시 통과시킨다.

오류(전달 금지)
- number-provenance: 슬라이드(본문·표·차트 값)의 숫자가 출처 텍스트에도, outline의 derived(계산값)·given(사용자 제공값)에도 없다
- leftover-text: 템플릿 안내 문구가 남아 있다("텍스트를 입력하십시오", "Click to add" 등)
- empty-placeholder: 비어 있는 플레이스홀더가 있다
- missing-headline: 제목(헤드라인)이 없는 내용 슬라이드
- slide-count: outline 슬라이드 수와 다르다
- overflow: 글자가 상자를 넘칠 가능성이 높다(추정)
- unknown-source: 노트의 출처 ID가 sources.json에 없다
경고
- long-headline(45자 초과), many-bullets(상위 항목 7개 이상)

숫자 대조 규칙
- 콤마·공백을 무시하고 값으로 비교한다(1,234 = 1234, 12.0 = 12)
- 반올림 허용: 슬라이드 숫자의 소수 자릿수로 반올림했을 때 같은 출처 숫자가 있으면 통과
- 검사 제외: 9 이하 정수(개수·순번), 1990~2100 연도, 슬라이드 번호, 날짜(2026-03-16, 3/16, 10월 10일 형태)
- 출처에 없는 계산값(증감률, 합계, 단위 환산)은 outline.derived에 {"value", "formula", "inputs"}로 선언한다

사용법
  python verify_deck.py report.pptx --sources <sources 폴더> [--outline outline.json] [--json]
종료 코드: 오류 0개면 0
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

LEFTOVER = re.compile(r"(텍스트를 입력|제목을 입력|부제목을 입력|내용을 입력|Click to (add|edit)|Lorem ipsum|여기에 입력|\[여기|TODO|TBD|XXX)", re.I)
# 앞 경계는 영문자·숫자·마침표만 본다(\w는 한글도 포함해 "약100만"의 100을 놓친다)
NUM = re.compile(r"(?<![A-Za-z0-9.])[-+−]?\d{1,3}(?:,\d{3})+(?:\.\d+)?(?![0-9])|(?<![A-Za-z0-9.,])[-+−]?\d+(?:\.\d+)?")
DATE_PATTERNS = [
    re.compile(r"\b(19|20)\d{2}[-./]\s?\d{1,2}[-./]\s?\d{1,2}\.?"),
    re.compile(r"\b\d{1,2}/\d{1,2}\b"),
    re.compile(r"\d{1,2}\s?월\s?\d{1,2}\s?일"),
    re.compile(r"\d{1,2}:\d{2}"),
    re.compile(r"\b(19|20)\d{2}\s?년\s?\d{1,2}\s?월"),
    re.compile(r"\bQ[1-4]\b|\b[1-4]Q\b|[1-4]분기"),
    re.compile(r"\b\d{1,2}\.\d{1,2}\.?\s?\([월화수목금토일]\)"),  # 02.22(일)
    re.compile(r"\b\d{1,2}/\d{1,2}\s?\([월화수목금토일]\)"),
]
# 숫자 뒤 한국어 단위 → 배수. "5,893만 원"은 출처의 58,931,000과 같은 값이다
UNIT = re.compile(r"\s?(조|억|천만|백만|만|천)")
UNIT_MULT = {"조": 1e12, "억": 1e8, "천만": 1e7, "백만": 1e6, "만": 1e4, "천": 1e3}


def to_value(tok: str) -> float | None:
    t = tok.replace(",", "").replace("−", "-").replace("+", "").strip()
    try:
        return float(t)
    except ValueError:
        return None


def decimals(tok: str) -> int:
    t = tok.replace(",", "")
    return len(t.split(".")[1]) if "." in t else 0


def strip_dates(text: str) -> str:
    for p in DATE_PATTERNS:
        text = p.sub(" ", text)
    return text


def numbers_in(text: str) -> list[str]:
    return [m.group(0) for m in NUM.finditer(strip_dates(text))]


def numbers_with_units(text: str) -> list[tuple[str, float]]:
    """(숫자 토큰, 단위 배수). 단위가 없으면 배수 1."""
    t = strip_dates(text)
    out = []
    for m in NUM.finditer(t):
        u = UNIT.match(t, m.end())
        out.append((m.group(0), UNIT_MULT[u.group(1)] if u else 1.0))
    return out


def exempt(tok: str) -> bool:
    v = to_value(tok)
    if v is None:
        return True
    if decimals(tok) == 0 and abs(v) <= 9:
        return True
    if decimals(tok) == 0 and 1990 <= v <= 2100 and "," not in tok:
        return True
    return False


class SourceIndex:
    def __init__(self, values: list[float]):
        self.values = values
        self.exact = {round(v, 6) for v in values}

    def has(self, tok: str, mult: float = 1.0, pct_fraction: bool = False) -> bool:
        """출처에 같은 값이 있는가.

        - 그대로 일치하거나, 슬라이드 숫자의 소수 자릿수로 반올림해 일치
        - mult(단위 배수)가 있으면 출처값/배수를 반올림해 일치 (5,893만 ↔ 58,931,000)
        - pct_fraction이면 0.6 ↔ 60(%)도 인정 (차트가 비율을 소수로 저장하는 경우)
        """
        v = to_value(tok)
        if v is None:
            return True
        if round(abs(v), 6) in self.exact or round(v, 6) in self.exact:
            return True
        d = decimals(tok)
        for s in self.values:
            if round(abs(s), d) == round(abs(v), d):
                return True
            if mult != 1.0 and round(abs(s) / mult, d) == round(abs(v), d):
                return True
            if pct_fraction and round(abs(s), max(d - 2, 0)) == round(abs(v) * 100, max(d - 2, 0)):
                return True
        return False


def load_source_numbers(src_dir: Path) -> tuple[SourceIndex, set[str]]:
    manifest = json.loads((src_dir / "sources.json").read_text(encoding="utf-8"))
    vals: list[float] = []
    ids = set()
    for e in manifest:
        if e.get("id"):
            ids.add(e["id"])
        if e.get("md"):
            text = (src_dir / e["md"]).read_text(encoding="utf-8")
            vals += [v for v in (to_value(t) for t in NUM.findall(text)) if v is not None]
    extra = src_dir / "extra_numbers.txt"  # 웹·대화에서 얻은 수치를 출처와 함께 적어 둔 파일(선택)
    if extra.is_file():
        vals += [v for v in (to_value(t) for t in NUM.findall(extra.read_text(encoding="utf-8"))) if v is not None]
    return SourceIndex(vals), ids


def shape_texts(shape) -> list[tuple[str, str]]:
    """(위치 라벨, 텍스트) 목록. 표 셀·그룹 포함."""
    out = []
    if shape.shape_type == 6 and hasattr(shape, "shapes"):  # group
        for s in shape.shapes:
            out += shape_texts(s)
        return out
    if shape.has_text_frame:
        out.append((shape.name, shape.text_frame.text))
    if getattr(shape, "has_table", False) and shape.has_table:
        for r, row in enumerate(shape.table.rows):
            for c, cell in enumerate(row.cells):
                out.append((f"{shape.name}[{r},{c}]", cell.text))
    return out


def estimate_overflow(shape) -> bool:
    if not shape.has_text_frame or not shape.width or not shape.height:
        return False
    w_in, h_in = shape.width / 914400 - 0.2, shape.height / 914400 - 0.1
    if w_in <= 0 or h_in <= 0:
        return False
    # 글자 크기가 명시되지 않으면 템플릿에서 상속받는다. 제목 칸은 보통 40~44pt라 18pt로 가정하면 넘침을 놓친다
    default = 18
    if shape.is_placeholder:
        t = shape.placeholder_format.type
        default = 40 if t in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE) else (
            28 if t == PP_PLACEHOLDER.SUBTITLE else 24)
    total = 0.0
    for p in shape.text_frame.paragraphs:
        size = None
        for r in p.runs:
            if r.font.size:
                size = r.font.size.pt
                break
        size = size or default
        text = "".join(r.text for r in p.runs)
        units = sum(1.0 if ord(ch) > 0x2E7F else 0.55 for ch in text)
        lines = max(1, math.ceil(units / max(1.0, w_in * 72 / size)))
        total += lines * size * 1.35 / 72
    return total > h_in * 1.15


def verify(pptx: Path, src_dir: Path, outline: dict | None) -> dict:
    prs = Presentation(str(pptx))
    idx, ids = load_source_numbers(src_dir)
    declared: list[float] = []
    if outline:
        for d in outline.get("derived", []):
            v = to_value(str(d.get("value", "")).replace("%", "").strip())
            if v is not None:
                declared.append(v)
            if not d.get("formula"):
                pass
        for g in outline.get("given", []):
            declared += [v for v in (to_value(t) for t in NUM.findall(str(g))) if v is not None]
    decl = SourceIndex(declared)
    errors, warnings, unverified = [], [], []
    checked = 0
    if outline and "slides" in outline and len(outline["slides"]) != len(prs.slides):
        errors.append({"rule": "slide-count", "slide": None,
                       "detail": f"outline {len(outline['slides'])}장 ≠ pptx {len(prs.slides)}장"})
    for n, slide in enumerate(prs.slides, 1):
        title_text = ""
        for ph in slide.placeholders:
            if ph.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                title_text = ph.text_frame.text.strip()
            elif ph.has_text_frame and not ph.text_frame.text.strip() and ph.placeholder_format.type not in (
                    PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER):
                errors.append({"rule": "empty-placeholder", "slide": n, "detail": f"빈 칸: {ph.name}"})
        for sh in slide.shapes:
            if sh.name == "Headline" and sh.has_text_frame:
                title_text = title_text or sh.text_frame.text.strip()
        if n > 1 and not title_text:
            errors.append({"rule": "missing-headline", "slide": n, "detail": "제목이 없는 슬라이드"})
        if len(title_text) > 45:
            warnings.append({"rule": "long-headline", "slide": n, "detail": f"{len(title_text)}자"})
        for sh in slide.shapes:
            for label, text in shape_texts(sh):
                if LEFTOVER.search(text):
                    errors.append({"rule": "leftover-text", "slide": n, "detail": f"{label}: {LEFTOVER.search(text).group(0)}"})
                if sh.name in ("PageNo", "Sources", "Byline"):
                    continue
                for tok, mult in numbers_with_units(text):
                    if exempt(tok) and mult == 1.0:
                        continue
                    checked += 1
                    if not (idx.has(tok, mult) or decl.has(tok, mult)):
                        unverified.append({"slide": n, "where": label, "number": tok, "context": text.strip()[:80]})
            if sh.has_text_frame and estimate_overflow(sh):
                errors.append({"rule": "overflow", "slide": n, "detail": f"{sh.name}: 글자가 상자를 넘칠 가능성"})
            if sh.has_text_frame and sh.name in ("Body", "LeftBody", "RightBody"):
                top = [p for p in sh.text_frame.paragraphs if p.level == 0 and "".join(r.text for r in p.runs).strip()]
                if len(top) >= 7:
                    warnings.append({"rule": "many-bullets", "slide": n, "detail": f"상위 항목 {len(top)}개"})
            if getattr(sh, "has_chart", False) and sh.has_chart:
                for plot in sh.chart.plots:
                    for ser in plot.series:
                        for v in ser.values:
                            if v is None:
                                continue
                            tok = f"{v:.6f}".rstrip("0").rstrip(".")
                            if exempt(tok):
                                continue
                            checked += 1
                            if not (idx.has(tok, pct_fraction=True) or decl.has(tok, pct_fraction=True)):
                                unverified.append({"slide": n, "where": f"차트 '{ser.name}'", "number": tok, "context": ""})
        if slide.has_notes_slide:
            for sid in re.findall(r"\bS\d{2}\b", slide.notes_slide.notes_text_frame.text):
                if sid not in ids:
                    errors.append({"rule": "unknown-source", "slide": n, "detail": f"출처 ID {sid}가 sources.json에 없음"})
    for u in unverified:
        errors.append({"rule": "number-provenance", "slide": u["slide"],
                       "detail": f"{u['where']}: '{u['number']}'의 출처를 찾을 수 없음 — 출처 숫자를 쓰거나 outline.derived에 계산식을 선언"
                                 + (f" (문맥: {u['context']})" if u["context"] else "")})
    return {"ok": not errors, "slides": len(prs.slides), "numbers_checked": checked,
            "numbers_unverified": len(unverified), "errors": errors, "warnings": warnings}


def main(argv: list[str]) -> int:
    for st in (sys.stdout, sys.stderr):
        try:
            st.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
        except Exception:
            pass
    ap = argparse.ArgumentParser(description="보고용 PPT 품질 게이트")
    ap.add_argument("pptx")
    ap.add_argument("--sources", required=True)
    ap.add_argument("--outline")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args(argv)
    outline = json.loads(Path(a.outline).read_text(encoding="utf-8")) if a.outline else None
    res = verify(Path(a.pptx), Path(a.sources), outline)
    if a.json:
        print(json.dumps(res, ensure_ascii=False, indent=2))
    else:
        for e in res["errors"]:
            print(f"ERROR {e['rule']:<18} {('#' + str(e['slide'])) if e['slide'] else '':>4}  {e['detail']}")
        for w in res["warnings"]:
            print(f"WARN  {w['rule']:<18} #{w['slide']:>3}  {w['detail']}")
        print(f"슬라이드 {res['slides']}장, 숫자 {res['numbers_checked']}개 대조, 출처 미확인 {res['numbers_unverified']}개")
        print("DECK VERIFY OK" if res["ok"] else "DECK VERIFY FAILED")
    return 0 if res["ok"] else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
