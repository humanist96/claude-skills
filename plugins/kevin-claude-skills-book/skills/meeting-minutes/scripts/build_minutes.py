#!/usr/bin/env python3
"""minutes.json(구조화 회의록) → 용도·형식별 문서. 형식은 references/minutes-schema.md.

용도(--purpose, 기본은 minutes.json의 meta.purpose)
  team      팀 공유: 핵심 요약 → 결정 → 할 일 표 → 미결·다음 회의 → 확인 필요
  report    상위 보고: 결론(결정) 먼저 → 배경·안건별 논의 → 다음 단계(할 일) → 리스크·확인 필요
  personal  개인 기록: 안건별 상세 논의 + 결정 + 할 일 + 근거 대조표
  email     이메일: 제목 줄, 인사, 요약, 결정, 할 일, 맺음(근거 번호 없음)
형식(--format): md | docx | notion | text
근거(--evidence): appendix(문서 끝에 근거 발언 대조표, md·docx 기본) | none(email·notion·text 기본)

결정·할 일·미결이 비어 있으면 섹션을 지우지 않고 "없음"으로 쓴다(빠진 것인지 없는 것인지 구분).

사용법: python build_minutes.py <minutes.json> --out <파일> [--format md] [--purpose team] [--transcript transcript.json]
       [--evidence appendix|none] [--font "맑은 고딕"]
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

SKILL_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent / "_vendor"))

L = {
    "ko": {"date": "일시", "att": "참석자", "absent": "불참", "summary": "핵심 요약", "decisions": "결정 사항", "actions": "할 일",
           "task": "할 일", "owner": "담당자", "due": "기한", "open": "미결·다음 논의", "next": "다음 회의", "uncertain": "확인 필요",
           "agenda": "안건별 논의", "conclusion": "결론", "nextsteps": "다음 단계", "risks": "리스크", "evidence": "근거 발언",
           "none": "없음", "unknown": "미상", "hello": "안녕하세요.", "intro": "{title} 내용을 공유드립니다.",
           "bye": "수정하거나 보탤 내용이 있으면 회신 부탁드립니다.\n감사합니다.", "subject": "제목", "role": "작성 관점",
           "src": "원본", "evid_note": "회의록의 각 항목이 어느 발언에서 나왔는지 보여 준다."},
    "en": {"date": "Date", "att": "Attendees", "absent": "Absent", "summary": "Summary", "decisions": "Decisions", "actions": "Action items",
           "task": "Task", "owner": "Owner", "due": "Due", "open": "Open questions", "next": "Next meeting", "uncertain": "Needs confirmation",
           "agenda": "Discussion by topic", "conclusion": "Conclusion", "nextsteps": "Next steps", "risks": "Risks", "evidence": "Evidence",
           "none": "None", "unknown": "Unknown", "hello": "Hi all,", "intro": "Here is a summary of {title}.",
           "bye": "Please reply if anything needs to be corrected or added.\nThanks.", "subject": "Subject", "role": "Perspective",
           "src": "Source", "evid_note": "Where each item in these minutes came from."},
}


def txt(x) -> str:
    return x["text"] if isinstance(x, dict) else str(x)


def ev(x) -> list[str]:
    return x.get("evidence", []) if isinstance(x, dict) else []


def blocks(m: dict, purpose: str, lab: dict) -> list[tuple]:
    """문서를 형식 중립 블록 목록으로 만든다: (kind, payload)"""
    meta = m.get("meta", {})
    title = meta.get("title") or ("회의록" if lab is L["ko"] else "Meeting minutes")
    dec = m.get("decisions", [])
    act = m.get("action_items", [])
    opn = m.get("open_questions", [])
    unc = m.get("uncertain", [])
    nxt = m.get("next_meeting")
    rows = [[a.get("task", ""), a.get("owner") or "미정", a.get("due") or "미정"] for a in act]
    b: list[tuple] = []
    if purpose == "email":
        when = meta.get("date") if meta.get("date") and meta.get("date") not in ("미상", "Unknown", "unknown") else ""
        b.append(("para", f"{lab['subject']}: [{'회의록' if lab is L['ko'] else 'Minutes'}] {title}" + (f" ({when})" if when else "")))
        b.append(("para", lab["hello"]))
        b.append(("para", lab["intro"].format(title=title)))
    else:
        b.append(("title", title))
        info = [(lab["date"], " ".join(x for x in [meta.get("date") or lab["unknown"], meta.get("time") or ""] if x)),
                (lab["att"], ", ".join(meta.get("attendees", [])) or lab["unknown"])]
        if meta.get("absent"):
            info.append((lab["absent"], ", ".join(meta["absent"])))
        b.append(("info", info))
    summary = m.get("summary", [])

    def dec_block(h):
        b.append(("h", h))
        b.append(("bullets", [txt(d) for d in dec] or [lab["none"]]))
        for sec in m.get("sections", []):  # 역할별 추가 섹션(예: PM의 리스크 및 이슈, 영업의 고객 요청 사항)
            b.append(("h", sec.get("title", "")))
            b.append(("bullets", [txt(x) for x in sec.get("items", [])] or [lab["none"]]))

    def act_block(h):
        b.append(("h", h))
        if rows:
            b.append(("table", [lab["task"], lab["owner"], lab["due"]], rows))
        else:
            b.append(("bullets", [lab["none"]]))

    def tail():
        b.append(("h", lab["open"]))
        b.append(("bullets", [txt(o) for o in opn] or [lab["none"]]))
        b.append(("h", lab["next"]))
        b.append(("bullets", [txt(nxt) if nxt else lab["unknown"]]))
        if unc:
            b.append(("h", lab["uncertain"]))
            b.append(("bullets", [txt(u) for u in unc]))

    if purpose == "report":
        b.append(("h", lab["conclusion"]))
        b.append(("bullets", summary[:3] or [txt(d) for d in dec] or [lab["none"]]))  # 결론 = 핵심 요약(없으면 결정)
        if m.get("agenda"):
            b.append(("h", lab["agenda"]))
            for a in m["agenda"]:
                b.append(("h3", a.get("title", "")))
                b.append(("bullets", a.get("points", []) or [lab["none"]]))
        dec_block(lab["decisions"])
        act_block(lab["nextsteps"])
        if m.get("risks"):
            b.append(("h", lab["risks"]))
            b.append(("bullets", [txt(r) for r in m["risks"]]))
        tail()
    elif purpose == "personal":
        if summary:
            b.append(("h", lab["summary"]))
            b.append(("bullets", summary))
        b.append(("h", lab["agenda"]))
        for a in m.get("agenda", []) or [{"title": lab["none"], "points": []}]:
            b.append(("h3", a.get("title", "")))
            b.append(("bullets", a.get("points", []) or [lab["none"]]))
        dec_block(lab["decisions"])
        act_block(lab["actions"])
        tail()
    elif purpose == "email":
        if summary:
            b.append(("bullets", summary))
        dec_block(lab["decisions"])
        b.append(("h", lab["actions"]))
        b.append(("bullets", [f"{r[0]} — {r[1]}, {r[2]}" for r in rows] or [lab["none"]]))
        tail()
        b.append(("para", lab["bye"]))
    else:  # team
        b.append(("h", lab["summary"]))
        b.append(("bullets", summary or [lab["none"]]))
        dec_block(lab["decisions"])
        act_block(lab["actions"])
        tail()
    return b


def evidence_rows(m: dict, transcript: dict | None) -> list[list[str]]:
    by_id = {s["id"]: s for s in (transcript or {}).get("segments", [])}
    rows = []
    for kind, items in (("결정", m.get("decisions", [])), ("할 일", m.get("action_items", [])),
                        ("미결", m.get("open_questions", [])), ("다음 회의", [m["next_meeting"]] if m.get("next_meeting") else [])):
        for it in items:
            label = it.get("task") if kind == "할 일" else txt(it)
            quotes = []
            for i in ev(it):
                s = by_id.get(i)
                quotes.append(f"{i} {((s.get('speaker') or '') + ': ') if s and s.get('speaker') else ''}{s['text'][:90] if s else '(없는 번호)'}")
            rows.append([kind, label, " / ".join(quotes) or "-"])
    return rows


def render_md(b, notion: bool) -> str:
    out = []
    for k, *p in b:
        if k == "title":
            out.append(f"# {p[0]}\n")
        elif k == "info":
            out.append("\n".join(f"- **{a}**: {v}" for a, v in p[0]) + "\n")
        elif k == "h":
            out.append(f"## {p[0]}\n")
        elif k == "h3":
            out.append(f"### {p[0]}\n")
        elif k == "para":
            out.append(p[0] + "\n")
        elif k == "bullets":
            out.append("\n".join(f"- {x}" for x in p[0]) + "\n")
        elif k == "table":
            head, rows = p
            if notion:  # 노션 붙여넣기: 할 일은 체크박스가 가장 잘 살아남는다
                out.append("\n".join(f"- [ ] {r[0]} — {r[1]} · {r[2]}" for r in rows) + "\n")
            else:
                out.append("| " + " | ".join(head) + " |\n|" + "---|" * len(head) + "\n" +
                           "\n".join("| " + " | ".join(str(c).replace("|", "/") for c in r) + " |" for r in rows) + "\n")
    return "\n".join(out)


def render_text(b) -> str:
    out = []
    for k, *p in b:
        if k == "title":
            out.append(p[0])
        elif k == "info":
            out += [f"{a}: {v}" for a, v in p[0]]
        elif k in ("h", "h3"):
            out.append(f"\n[{p[0]}]")
        elif k == "para":
            out.append(p[0])
        elif k == "bullets":
            out += [f"- {x}" for x in p[0]]
        elif k == "table":
            out += [f"- {r[0]} ({r[1]}, {r[2]})" for r in p[1]]
    return "\n".join(out) + "\n"


def render_docx(b, path: Path, font: str):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    for sname in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
        st = doc.styles[sname]
        st.font.name = font
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rf)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            rf.set(qn(attr), font)
        rf.attrib.pop(qn("w:asciiTheme"), None)
        rf.attrib.pop(qn("w:eastAsiaTheme"), None)
        rf.attrib.pop(qn("w:hAnsiTheme"), None)
    doc.styles["Normal"].font.size = Pt(10.5)

    def shade(cell, color="DDE6F5"):
        tcpr = cell._tc.get_or_add_tcPr()
        sh = tcpr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): color})
        tcpr.append(sh)

    def table(head, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(head))
        t.style = "Table Grid"
        for i, h in enumerate(head):
            c = t.rows[0].cells[i]
            c.text = h
            c.paragraphs[0].runs[0].bold = True
            shade(c)
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = str(v)
        doc.add_paragraph()

    for k, *p in b:
        if k == "title":
            doc.add_heading(p[0], level=0)
        elif k == "info":
            table([a for a, _ in p[0]], [[v for _, v in p[0]]])
        elif k == "h":
            doc.add_heading(p[0], level=1)
        elif k == "h3":
            doc.add_heading(p[0], level=2)
        elif k == "para":
            for line in p[0].split("\n"):
                doc.add_paragraph(line)
        elif k == "bullets":
            for x in p[0]:
                doc.add_paragraph(x, style="List Bullet")
        elif k == "table":
            table(*p)
    doc.save(str(path))


def main(argv: list[str]) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    ap = argparse.ArgumentParser(description="minutes.json → 회의록 문서")
    ap.add_argument("minutes")
    ap.add_argument("--out", required=True)
    ap.add_argument("--format", choices=["md", "docx", "notion", "text"])
    ap.add_argument("--purpose", choices=["team", "report", "personal", "email"])
    ap.add_argument("--transcript")
    ap.add_argument("--evidence", choices=["appendix", "none"])
    ap.add_argument("--font", default="맑은 고딕")
    a = ap.parse_args(argv)
    m = json.loads(Path(a.minutes).read_text(encoding="utf-8"))
    out = Path(a.out)
    fmt = a.format or {".docx": "docx", ".txt": "text"}.get(out.suffix.lower(), "md")
    purpose = a.purpose or m.get("meta", {}).get("purpose") or "team"
    lab = L["en"] if str(m.get("meta", {}).get("language", "ko")).startswith("en") else L["ko"]
    evid = a.evidence or ("appendix" if fmt in ("md", "docx") and purpose != "email" else "none")
    tpath = a.transcript or m.get("meta", {}).get("transcript")
    transcript = None
    if tpath and Path(tpath).is_file():
        transcript = json.loads(Path(tpath).read_text(encoding="utf-8"))
    b = blocks(m, purpose, lab)
    if evid == "appendix":
        b.append(("h", lab["evidence"]))
        b.append(("para", lab["evid_note"]))
        b.append(("table", ["구분" if lab is L["ko"] else "Type", "항목" if lab is L["ko"] else "Item", lab["evidence"]], evidence_rows(m, transcript)))
    out.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "docx":
        render_docx(b, out, a.font)
    elif fmt == "text":
        out.write_text(render_text(b), encoding="utf-8")
    else:
        out.write_text(render_md(b, notion=(fmt == "notion")), encoding="utf-8")
    n_dec, n_act = len(m.get("decisions", [])), len(m.get("action_items", []))
    print(f"출력: {out} — 용도 {purpose}, 형식 {fmt}, 결정 {n_dec}개, 할 일 {n_act}개, 근거 {'부록' if evid == 'appendix' else '없음'}")
    print(f"BUILD DONE format={fmt} purpose={purpose} decisions={n_dec} actions={n_act}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
