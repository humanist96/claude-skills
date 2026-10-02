#!/usr/bin/env python3
"""엑셀·CSV 구조 파악 — 어떤 기능(정리·분석·취합)이 필요한지, 어떤 열이 무엇인지 판단하는 근거.

시트마다: 행·열 수, 머리글 행 위치, 열별 자료형·빈 칸·고유값 비율, 의미 추정(전화·이메일·날짜·금액·코드),
키 후보, 일련번호 열, 완전 중복 행 수(일련번호 제외), 수식·병합 셀 수, 시트 간 공통 열.

사용법: python excel_profile.py <파일.xlsx|.csv> [--json]
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from excel_utils import email_issue, normalize_date, normalize_phone  # noqa: E402

PHONE_HINT = re.compile(r"(전화|연락처|휴대|핸드폰|phone|tel|mobile)", re.I)
DATE_HINT = re.compile(r"(일자|날짜|일$|date|가입일|주문일|입고일|출시일)", re.I)
EMAIL_HINT = re.compile(r"(메일|email|e-mail)", re.I)


def detect_header_row(rows: list[tuple]) -> int:
    """앞 10행 중 '대부분이 문자열이고 비어 있지 않은' 첫 행을 머리글로 본다(제목 행·빈 행이 위에 있는 양식 대응)."""
    for i, r in enumerate(rows[:10]):
        vals = [v for v in r if v is not None and str(v).strip() != ""]
        if len(vals) >= max(2, int(len(r) * 0.6)) and sum(isinstance(v, str) for v in vals) >= len(vals) * 0.8:
            return i
    return 0


def col_semantics(name: str, values: list) -> str:
    vals = [v for v in values if v is not None and str(v).strip() != ""]
    if not vals:
        return "empty"
    n = len(vals)
    if PHONE_HINT.search(name) or sum(normalize_phone(v)[0] is not None for v in vals[:200]) >= 0.6 * min(n, 200):
        if sum(re.sub(r"\D", "", str(v)).__len__() >= 8 for v in vals[:200]) >= 0.6 * min(n, 200):
            return "phone"
    if EMAIL_HINT.search(name) or sum("@" in str(v) for v in vals[:200]) >= 0.6 * min(n, 200):
        return "email"
    if DATE_HINT.search(name) and sum(normalize_date(v)[0] is not None or "모호" in normalize_date(v)[1] for v in vals[:200]) >= 0.5 * min(n, 200):
        return "date"
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals):
        return "number"
    if sum(bool(re.fullmatch(r"[A-Z]{1,5}[-_]?[A-Z0-9-]{2,}", str(v))) for v in vals[:200]) >= 0.8 * min(n, 200):
        return "code"
    return "text"


def is_sequence(values: list) -> bool:
    ints = [v for v in values if isinstance(v, int) and not isinstance(v, bool)]
    return len(ints) == len(values) and len(ints) >= 3 and sorted(ints) == list(range(min(ints), min(ints) + len(ints)))


def profile_sheet(title: str, rows: list[tuple], formulas: int = 0, merged: int = 0) -> dict:
    if not rows:
        return {"sheet": title, "rows": 0, "columns": [], "note": "빈 시트"}
    h = detect_header_row(rows)
    header = [str(c).strip() if c is not None else f"(열{i + 1})" for i, c in enumerate(rows[h])]
    data = [r for r in rows[h + 1:] if any(v is not None and str(v).strip() != "" for v in r)]
    cols = []
    seq_cols = []
    for j, name in enumerate(header):
        values = [r[j] if j < len(r) else None for r in data]
        nonnull = [v for v in values if v is not None and str(v).strip() != ""]
        uniq = len({str(v) for v in nonnull})
        sem = col_semantics(name, values)
        if is_sequence(values):
            seq_cols.append(name)
        cols.append({"name": name, "semantic": sem, "nulls": len(values) - len(nonnull),
                     "unique_ratio": round(uniq / len(nonnull), 3) if nonnull else 0,
                     "sample": [str(v) for v in nonnull[:3]]})
    keys = [c["name"] for c in cols if c["unique_ratio"] >= 0.95 and c["nulls"] == 0 and c["semantic"] in ("code", "text")
            and c["name"] not in seq_cols]
    idx = [j for j, n in enumerate(header) if n not in seq_cols]
    seen, dups = set(), 0
    for r in data:
        t = tuple(r[j] if j < len(r) else None for j in idx)
        dups += t in seen
        seen.add(t)
    return {"sheet": title, "header_row": h + 1, "rows": len(data), "columns": cols, "key_candidates": keys,
            "sequence_columns": seq_cols, "exact_duplicates_excluding_sequence": dups,
            "formulas": formulas, "merged_ranges": merged}


def profile(path: Path) -> dict:
    sheets = []
    if path.suffix.lower() in (".csv", ".tsv"):
        import csv
        for enc in ("utf-8-sig", "cp949"):
            try:
                with open(path, encoding=enc, newline="") as f:
                    rows = [tuple(None if c == "" else c for c in r) for r in csv.reader(f, delimiter="\t" if path.suffix == ".tsv" else ",")]
                break
            except UnicodeDecodeError:
                continue
        sheets.append(profile_sheet(path.stem, rows))
    else:
        from openpyxl import load_workbook
        wb = load_workbook(path)
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            f = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
            sheets.append(profile_sheet(ws.title, rows, f, len(ws.merged_cells.ranges)))
    common = {}
    for s in sheets:
        for c in s.get("columns", []):
            common.setdefault(c["name"], []).append(s["sheet"])
    shared = {k: v for k, v in common.items() if len(v) > 1}
    hint = "consolidate" if len([s for s in sheets if s.get("rows")]) >= 3 and shared else (
        "clean" if any(c["semantic"] in ("phone", "email") for s in sheets for c in s.get("columns", [])) else "analyze")
    return {"file": str(path), "sheets": sheets, "shared_columns": shared, "suggested_function": hint}


def main(argv: list[str]) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args(argv)
    p = profile(Path(a.input))
    if a.json:
        print(json.dumps(p, ensure_ascii=False, indent=2))
        return 0
    for s in p["sheets"]:
        print(f"[{s['sheet']}] {s.get('rows', 0)}행 (머리글 {s.get('header_row')}행), 키 후보 {s.get('key_candidates')}, "
              f"일련번호 {s.get('sequence_columns')}, 완전중복 {s.get('exact_duplicates_excluding_sequence')}, 수식 {s.get('formulas')}")
        for c in s.get("columns", []):
            print(f"   - {c['name']:<16} {c['semantic']:<7} 빈칸 {c['nulls']:>4}  고유 {c['unique_ratio']:.2f}  예: {', '.join(c['sample'])}")
    print(f"공통 열(시트 간): {list(p['shared_columns'])[:10]}")
    print(f"추천 기능: {p['suggested_function']}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
