#!/usr/bin/env python3
"""수식 재계산 + 오류 셀 점검.

openpyxl로 쓴 수식에는 계산된 값이 없다. 그대로 두면 (1) 파일을 연 사람이 보기 전까지 값이 비어 있고,
(2) 수식 오류(#N/A, #REF! 등)가 있어도 알 수 없다. 그래서 저장 후 실제 계산 엔진으로 한 번 계산해 저장한다.

엔진(있는 것 순서대로)
1. Windows Microsoft Excel (COM 자동화, PowerShell)
2. LibreOffice(soffice) headless 변환
엔진이 없으면(또는 환경변수 CLAUDE_SKILLS_NO_RECALC=1) 'RECALC UNAVAILABLE'을 출력한다. 그때는 사용자에게 파일을 Excel로 열어 저장하면 값이 채워진다고 안내한다.

출력 마지막 줄: RECALC OK formulas=<수식 수> errors=<오류 셀 수> (<엔진>)
사용법: python recalc.py <파일.xlsx> [--json]
"""
from __future__ import annotations

import json
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ERRORS = {"#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}


def via_excel(path: Path) -> bool:
    if platform.system() != "Windows" or not shutil.which("powershell"):
        return False
    p = str(path.resolve()).replace("'", "''")
    ps = f"""
$ErrorActionPreference = 'Stop'
$xl = New-Object -ComObject Excel.Application
$xl.Visible = $false; $xl.DisplayAlerts = $false
try {{
  $wb = $xl.Workbooks.Open('{p}')
  $xl.CalculateFull()
  $wb.Save(); $wb.Close($true)
}} finally {{ $xl.Quit(); [System.Runtime.Interopservices.Marshal]::ReleaseComObject($xl) | Out-Null }}
"""
    r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps], capture_output=True, text=True, timeout=300)
    return r.returncode == 0


def via_soffice(path: Path) -> bool:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return False
    with tempfile.TemporaryDirectory() as td:
        r = subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", td, str(path)],
                           capture_output=True, text=True, timeout=300)
        res = Path(td) / path.name
        if r.returncode != 0 or not res.is_file():
            return False
        shutil.copy2(res, path)
    return True


def scan(path: Path) -> dict:
    wf = load_workbook(path)
    wv = load_workbook(path, data_only=True)
    formulas, errors, uncached, cached, blank_ok = 0, [], 0, 0, 0
    for ws in wf.worksheets:
        vs = wv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas += 1
                    v = vs[c.coordinate].value
                    if v is None:
                        # Excel은 빈 문자열 결과(IFERROR(...,""))를 값 없이 저장한다
                        if re.search(r',\s*""\s*\)\s*$', c.value):
                            blank_ok += 1
                        else:
                            uncached += 1
                        continue
                    cached += 1
                    if isinstance(v, str) and v.strip() in ERRORS:
                        errors.append(f"{ws.title}!{c.coordinate}={v}")
                elif isinstance(c.value, str) and c.value.strip() in ERRORS:
                    errors.append(f"{ws.title}!{c.coordinate}={c.value}")
    if cached == 0:
        uncached += blank_ok
    return {"formulas": formulas, "errors": len(errors), "error_cells": errors[:20], "uncached_formulas": uncached}


def main(argv: list[str]) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    if not argv:
        print(__doc__)
        return 2
    path = Path(argv[0])
    if not path.is_file():
        print(f"오류: 파일이 없습니다: {path}")
        return 1
    engine = None
    engines = () if os.environ.get("CLAUDE_SKILLS_NO_RECALC") else (("Excel", via_excel), ("LibreOffice", via_soffice))
    for name, fn in engines:
        try:
            if fn(path):
                engine = name
                break
        except (subprocess.SubprocessError, OSError):
            continue
    res = scan(path)
    res["engine"] = engine
    if "--json" in argv:
        print(json.dumps(res, ensure_ascii=False, indent=2))
    for e in res["error_cells"]:
        print("ERROR CELL", e)
    if engine is None:
        print(f"RECALC UNAVAILABLE formulas={res['formulas']} (Excel·LibreOffice 없음 — 파일을 Excel로 열어 저장하면 값이 채워진다)")
        return 2
    print(f"RECALC OK formulas={res['formulas']} errors={res['errors']} ({engine})")
    return 0 if res["errors"] == 0 else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
