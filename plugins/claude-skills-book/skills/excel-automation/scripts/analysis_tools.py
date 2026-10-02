#!/usr/bin/env python3
"""분석·시각화 도구 — 표 쓰기, 엑셀 기본 차트, 히트맵 서식, 근거가 붙은 인사이트 시트.

왜 엑셀 기본 차트인가: 그림(PNG)으로 넣은 차트는 받는 사람이 값을 고치거나 확인할 수 없다.
openpyxl 차트는 시트의 셀을 참조하므로 클릭하면 근거 숫자가 보이고, 데이터를 고치면 차트도 바뀐다.

라이브러리로 쓰기(대화마다 다른 질문에 맞춰 조합)
    from analysis_tools import load_table, write_df, add_chart, add_heatmap, write_insights
    df = load_table("매출.xlsx")                       # 머리글 행 자동 감지
    ws = wb.create_sheet("카테고리별"); write_df(ws, table)
    add_chart(ws, "col", title="카테고리별 매출", cats_col=1, val_cols=[2], anchor="H2")
    write_insights(wb, [("1위 카테고리는 전자기기(31,730,550원)", "카테고리별!B2")])

표준 보고서(기간·금액 열이 있는 거래 데이터의 기본 분석)
    python analysis_tools.py report <입력.xlsx> --out <결과.xlsx> --date-col 주문일 --value-col 결제금액
          [--category-col 카테고리] [--item-col 상품명] [--sheet 이름]
    → 원본 시트 보존 + 요약통계·월별추이·카테고리별·월x카테고리·상위항목·인사이트 시트와 차트
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="DDE6F5")
INSIGHT_SHEET = "인사이트"
REPORT_SHEETS = ("요약통계", "월별추이", "카테고리별", "월x카테고리", "상위항목", INSIGHT_SHEET)


def load_table(path: str | Path, sheet: str | None = None) -> pd.DataFrame:
    """제목 행이 위에 있어도 머리글 행을 찾아 읽는다."""
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from excel_profile import profile
    p = profile(Path(path))
    sp = next((s for s in p["sheets"] if s["sheet"] == sheet), None) if sheet else p["sheets"][0]
    if sp is None:
        raise SystemExit(f"오류: 시트 '{sheet}'가 없습니다. 있는 시트: {[s['sheet'] for s in p['sheets']]}")
    return pd.read_excel(path, sheet_name=sp["sheet"], header=sp["header_row"] - 1)


def _plain(v):
    if hasattr(v, "item"):
        v = v.item()
    if isinstance(v, float) and v != v:
        return None
    return v


def write_df(ws, df: pd.DataFrame, start_row: int = 1, start_col: int = 1, number_format: str = "#,##0") -> tuple[int, int]:
    """DataFrame을 머리글 서식과 함께 쓴다. (마지막 행, 마지막 열)을 돌려준다."""
    for j, name in enumerate(df.columns, start_col):
        c = ws.cell(start_row, j, str(name))
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center")
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(row, start_col):
            v = _plain(v)
            c = ws.cell(i, j, v)
            if isinstance(v, float) and not float(v).is_integer():
                c.number_format = "#,##0.0"
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                c.number_format = number_format
    for j, name in enumerate(df.columns, start_col):
        col = [_plain(x) for x in df.iloc[:200, j - start_col]]
        width = max([len(str(name))] + [len(f"{x:,}" if isinstance(x, (int, float)) else str(x)) for x in col])
        ws.column_dimensions[get_column_letter(j)].width = min(40, max(9, width * 1.5))
    ws.freeze_panes = ws.cell(start_row + 1, start_col)
    return start_row + len(df), start_col + len(df.columns) - 1


def add_chart(ws, kind: str, title: str, cats_col: int, val_cols: list[int], anchor: str,
              header_row: int = 1, last_row: int | None = None, y_title: str | None = None, x_title: str | None = None):
    """시트 데이터를 참조하는 엑셀 기본 차트. kind: bar(가로 막대) | col(세로 막대) | line | pie"""
    last = last_row or ws.max_row
    chart = {"bar": BarChart, "col": BarChart, "line": LineChart, "pie": PieChart}[kind]()
    if kind in ("bar", "col"):
        chart.type = kind
    chart.title = title
    for vc in val_cols:
        chart.add_data(Reference(ws, min_col=vc, min_row=header_row, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=cats_col, min_row=header_row + 1, max_row=last))
    if kind != "pie":
        chart.y_axis.title = y_title
        chart.x_axis.title = x_title
        chart.y_axis.numFmt = "#,##0"
        chart.y_axis.delete = False
        chart.x_axis.delete = False
        if len(val_cols) == 1:
            chart.legend = None
    chart.width, chart.height = 18, 9
    ws.add_chart(chart, anchor)
    return chart


def add_heatmap(ws, cell_range: str):
    """색 단계 조건부 서식 히트맵(값이 클수록 진한 색). 그림이 아니라 셀이라 값이 그대로 보인다."""
    ws.conditional_formatting.add(cell_range, ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                             mid_type="percentile", mid_value=50, mid_color="9DC3E6",
                                                             end_type="max", end_color="2E75B6"))


def write_insights(wb, items: list[tuple[str, str]]):
    """인사이트(문장, 근거 셀) 목록을 '인사이트' 시트에 쓴다.
    문장 속 숫자는 통합 문서 안의 셀 값이어야 한다(verify_excel analysis가 대조한다)."""
    ws = wb[INSIGHT_SHEET] if INSIGHT_SHEET in wb.sheetnames else wb.create_sheet(INSIGHT_SHEET)
    ws.delete_rows(1, ws.max_row)
    ws.append(["번호", "인사이트", "근거"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
    for i, (text, ref) in enumerate(items, 1):
        ws.append([i, text, ref])
        ws.cell(i + 1, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 26
    return ws


def won(v: float) -> str:
    return f"{int(round(v)):,}원"


def report(src: Path, out: Path, date_col: str, value_col: str, category_col: str | None, item_col: str | None,
           sheet: str | None) -> dict:
    df = load_table(src, sheet)
    for c in [date_col, value_col] + [x for x in (category_col, item_col) if x]:
        if c not in df.columns:
            raise SystemExit(f"오류: 열 '{c}'가 없습니다. 가능한 열: {list(df.columns)}")
    df[date_col] = pd.to_datetime(df[date_col], errors="coerce")
    bad = int(df[date_col].isna().sum())
    df = df.dropna(subset=[date_col]).copy()
    df["월"] = df[date_col].dt.strftime("%Y-%m")
    wb = load_workbook(src)  # 원본 시트를 보존한 채 분석 시트를 더한다
    for name in REPORT_SHEETS:
        if name in wb.sheetnames:
            raise SystemExit(f"오류: '{name}' 시트가 이미 있습니다. 결과는 새 파일에 만드세요.")
    v = df[value_col]
    total = int(v.sum())
    m = df.groupby("월")[value_col].agg(["sum", "count"]).reset_index()
    m.columns = ["월", "합계", "건수"]
    m["전월대비"] = m["합계"].diff()
    m["전월대비(%)"] = (m["합계"].pct_change() * 100).round(1)
    m["합계"] = m["합계"].astype(int)
    month_avg = int(round(float(m["합계"].mean())))
    summ = pd.DataFrame({"지표": ["행 수", "합계", "평균", "중앙값", "최대", "최소", "월평균", "시작일", "종료일", "날짜 해석 실패 행"],
                         "값": [len(df), total, round(float(v.mean()), 1), float(v.median()), int(v.max()), int(v.min()), month_avg,
                                df[date_col].min().strftime("%Y-%m-%d"), df[date_col].max().strftime("%Y-%m-%d"), bad]})
    ws = wb.create_sheet("요약통계")
    write_df(ws, summ)

    ws = wb.create_sheet("월별추이")
    last, _ = write_df(ws, m)
    add_chart(ws, "line", f"월별 {value_col} 추이", 1, [2], "H2", last_row=last, y_title=value_col)

    insights: list[tuple[str, str]] = [
        (f"기간 전체 {value_col} 합계는 {won(total)}이고 월평균은 {won(month_avg)}이다.", "요약통계!B3, 요약통계!B8")]
    ib, iw = int(m["합계"].idxmax()), int(m["합계"].idxmin())
    best, worst = m.loc[ib], m.loc[iw]
    insights.append((f"가장 높은 달은 {best['월']}({won(best['합계'])}), 가장 낮은 달은 {worst['월']}({won(worst['합계'])})이다.",
                     f"월별추이!B{ib + 2}, 월별추이!B{iw + 2}"))
    if len(m) >= 2:
        lm = m.iloc[-1]
        sign = "늘었다" if lm["전월대비"] >= 0 else "줄었다"
        insights.append((f"최근 달 {lm['월']}은 전월보다 {won(abs(lm['전월대비']))}({abs(lm['전월대비(%)']):.1f}%) {sign}.",
                         f"월별추이!D{len(m) + 1}, 월별추이!E{len(m) + 1}"))
    if category_col:
        c = df.groupby(category_col)[value_col].agg(["sum", "count", "mean"]).reset_index().sort_values("sum", ascending=False)
        c.columns = [category_col, "합계", "건수", "평균"]
        c["비중(%)"] = (c["합계"] / total * 100).round(1)
        c["합계"] = c["합계"].astype(int)
        c["평균"] = c["평균"].round(0).astype(int)
        ws = wb.create_sheet("카테고리별")
        last, _ = write_df(ws, c)
        add_chart(ws, "col", f"{category_col}별 {value_col}", 1, [2], "H2", last_row=last, y_title=value_col)
        add_chart(ws, "pie", f"{category_col}별 비중", 1, [2], "H22", last_row=last)
        top = c.iloc[0]
        insights.append((f"{category_col} 1위는 {top[category_col]}로 {won(top['합계'])}이며 전체의 {top['비중(%)']:.1f}%다.",
                         "카테고리별!B2, 카테고리별!E2"))
        pv = df.pivot_table(index="월", columns=category_col, values=value_col, aggfunc="sum", fill_value=0).reset_index()
        pv.columns = [str(x) for x in pv.columns]
        ws = wb.create_sheet("월x카테고리")
        last, lc = write_df(ws, pv)
        add_heatmap(ws, f"B2:{get_column_letter(lc)}{last}")
    if item_col:
        it = df.groupby(item_col)[value_col].sum().sort_values(ascending=False).head(10).reset_index()
        it.columns = [item_col, "합계"]
        it["합계"] = it["합계"].astype(int)
        ws = wb.create_sheet("상위항목")
        last, _ = write_df(ws, it)
        add_chart(ws, "bar", f"{item_col} 상위 10", 1, [2], "E2", last_row=last, y_title=value_col)
        insights.append((f"{item_col} 1위는 {it.iloc[0][item_col]}({won(it.iloc[0]['합계'])})이다.", "상위항목!B2"))
    write_insights(wb, insights[:5])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {"out": str(out), "rows": len(df), "total": total, "insights": [t for t, _ in insights[:5]], "unparsed_dates": bad}


def main(argv: list[str]) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    ap = argparse.ArgumentParser(description="표준 분석 보고서")
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("report")
    r.add_argument("input")
    r.add_argument("--out", required=True)
    r.add_argument("--date-col", required=True)
    r.add_argument("--value-col", required=True)
    r.add_argument("--category-col")
    r.add_argument("--item-col")
    r.add_argument("--sheet")
    a = ap.parse_args(argv)
    src, out = Path(a.input), Path(a.out)
    if out.resolve() == src.resolve():
        print("오류: 결과 파일이 원본과 같습니다. 원본은 덮어쓰지 않습니다.")
        return 1
    res = report(src, out, a.date_col, a.value_col, a.category_col, a.item_col, a.sheet)
    print(f"분석 완료: {res['out']} — {res['rows']}행, 합계 {res['total']:,}")
    if res["unparsed_dates"]:
        print(f"  주의: 날짜를 해석하지 못한 행 {res['unparsed_dates']}개는 제외했다(요약통계 시트에 기록)")
    for t in res["insights"]:
        print("  -", t)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
