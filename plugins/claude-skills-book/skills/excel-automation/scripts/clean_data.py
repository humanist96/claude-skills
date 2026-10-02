#!/usr/bin/env python3
"""데이터 정리 엔진 — 규칙이 정해진 정리는 매번 코드로 새로 쓰지 않고 이 스크립트로 한다.

하는 일(원본은 절대 수정하지 않는다)
1. 완전 중복 행 제거(일련번호 열은 비교에서 제외, 첫 행 유지)
2. 전화번호·날짜 형식 통일 — 안전하게 바꿀 수 있는 값만. 나머지는 원래 값 유지 + 후보 시트
3. 이메일 오타·형식 오류, 숫자 극단값, 범주 표기 흔들림 → 고치지 않고 후보 시트
4. 빈 칸 → 채우지 않고 현황만 보고
5. 유사 중복(키 열이 같고 다른 값이 다른 행) → 지우지 않고 후보 시트
6. 모든 셀 변경을 '원본 행 번호·열·이전·이후·사유'로 기록(조용한 변경 금지)

출력 xlsx 시트: 정리완료, 변경리포트, 변경내역, 유사중복_후보, 오타_이상값_후보, 빈칸현황
출력 json(<출력>.report.json): verify_excel.py가 읽는 기계용 기록

사용법
  python clean_data.py <입력.xlsx> --out <결과.xlsx> [--sheet 이름] [--near-key 이름,연락처]
       [--phone-cols 연락처] [--date-cols 가입일] [--email-cols 이메일] [--number-cols 누적구매액]
       [--category-cols 지역] [--date-order mdy|dmy] [--outlier-k 3] [--keep-duplicates] [--no-overrides]
  열을 지정하지 않으면 excel_profile의 추정을 쓴다.
  사용자 맞춤(오버라이드 폴더): settings.yaml의 date_order·outlier_k·near_key·keep_duplicates,
  references/category-map.json의 범주 표기 사전 {"*": {"서울시": "서울"}, "등급": {"vip": "VIP"}}
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

SKILL_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent / "_vendor"))
from excel_profile import profile  # noqa: E402
from excel_utils import category_variants, email_issue, normalize_date, normalize_phone, outliers  # noqa: E402

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

HEAD_FILL = PatternFill("solid", fgColor="DDE6F5")


def sha256(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def split(s: str | None) -> list[str]:
    return [x.strip() for x in s.split(",") if x.strip()] if s else []


def user_rules() -> tuple[dict, dict | None, str | None]:
    """오버라이드 settings와 범주 표기 사전(references/category-map.json). 없으면 빈 값."""
    try:
        import overrides  # scripts/_vendor/overrides.py
    except ImportError:
        return {}, None, None
    st = overrides.settings("excel-automation", SKILL_DIR)["settings"]
    found = overrides.resolve("excel-automation", "references/category-map.json", SKILL_DIR)
    cmap = json.loads(Path(found["path"]).read_text(encoding="utf-8")) if found["path"] else None
    return st, cmap, found["source"]


def write_table(ws, header: list, rows: list[list]):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
    for r in rows:
        ws.append(r)
    for j, h in enumerate(header, 1):
        width = max([len(str(h))] + [len(str(r[j - 1])) for r in rows[:200] if j - 1 < len(r) and r[j - 1] is not None])
        ws.column_dimensions[get_column_letter(j)].width = min(45, max(8, width * 1.6))
    ws.freeze_panes = "A2"


def main(argv: list[str]) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    ap = argparse.ArgumentParser(description="엑셀 데이터 정리 엔진")
    ap.add_argument("input")
    ap.add_argument("--out", required=True)
    ap.add_argument("--sheet")
    ap.add_argument("--near-key")
    ap.add_argument("--phone-cols")
    ap.add_argument("--date-cols")
    ap.add_argument("--email-cols")
    ap.add_argument("--number-cols")
    ap.add_argument("--category-cols")
    ap.add_argument("--date-order", choices=["mdy", "dmy"])
    ap.add_argument("--outlier-k", type=float)
    ap.add_argument("--keep-duplicates", action="store_true")
    ap.add_argument("--no-overrides", action="store_true", help="사용자 맞춤(settings·범주 사전)을 무시")
    a = ap.parse_args(argv)
    st, cmap, cmap_src = ({}, None, None) if a.no_overrides else user_rules()
    # 우선순위: 명령 인자 > 사용자 맞춤 settings > 기본값
    a.date_order = a.date_order or (st.get("date_order") if st.get("date_order") in ("mdy", "dmy") else None)
    a.near_key = a.near_key or st.get("near_key")
    outlier_k = a.outlier_k or float(st.get("outlier_k", 3.0))
    a.keep_duplicates = a.keep_duplicates or bool(st.get("keep_duplicates", False))
    src, out = Path(a.input), Path(a.out)
    if out.resolve() == src.resolve():
        print("오류: 결과 파일이 원본과 같습니다. 원본은 덮어쓰지 않습니다.")
        return 1
    before = sha256(src)
    prof = profile(src)
    wb = load_workbook(src, data_only=True)
    ws = wb[a.sheet] if a.sheet else wb.worksheets[0]
    sp = next(s for s in prof["sheets"] if s["sheet"] == ws.title)
    rows = list(ws.iter_rows(values_only=True))
    h = sp["header_row"] - 1
    header = [str(c).strip() if c is not None else f"(열{i + 1})" for i, c in enumerate(rows[h])]
    sem = {c["name"]: c["semantic"] for c in sp["columns"]}
    pick = lambda arg, kind: split(arg) or [n for n, s in sem.items() if s == kind]  # noqa: E731
    phone_cols, date_cols, email_cols = pick(a.phone_cols, "phone"), pick(a.date_cols, "date"), pick(a.email_cols, "email")
    number_cols = split(a.number_cols) or [n for n, s in sem.items() if s == "number" and n not in sp["sequence_columns"]]
    cat_cols = split(a.category_cols) or [c["name"] for c in sp["columns"] if c["semantic"] == "text" and 0 < c["unique_ratio"] <= 0.5]
    for c in phone_cols + date_cols + email_cols + number_cols + cat_cols + split(a.near_key):
        if c not in header:
            print(f"오류: 열 '{c}'가 없습니다. 가능한 열: {header}")
            return 1
    col = {n: i for i, n in enumerate(header)}
    excel_row0 = h + 2  # 데이터 첫 행의 엑셀 행 번호
    data = []
    for k, r in enumerate(rows[h + 1:]):
        if any(v is not None and str(v).strip() != "" for v in r):
            data.append({"src_row": excel_row0 + k, "values": list(r) + [None] * (len(header) - len(r))})
    seq = set(sp["sequence_columns"])
    cmp_idx = [i for i, n in enumerate(header) if n not in seq]

    # 1. 완전 중복
    kept, removed, seen = [], [], {}
    for d in data:
        t = tuple(d["values"][i] for i in cmp_idx)
        if t in seen and not a.keep_duplicates:
            removed.append({"src_row": d["src_row"], "duplicate_of": seen[t]})
        else:
            seen.setdefault(t, d["src_row"])
            kept.append(d)

    changes, candidates = [], []
    # 2. 전화·날짜
    for d in kept:
        for cname in phone_cols:
            v = d["values"][col[cname]]
            if v is None or str(v).strip() == "":
                continue
            new, why = normalize_phone(v)
            if new is None:
                candidates.append([d["src_row"], cname, str(v), "전화번호 " + why, ""])
            elif new != v:
                changes.append({"src_row": d["src_row"], "column": cname, "before": v, "after": new, "reason": "전화번호 형식 통일(" + why + ")"})
                d["values"][col[cname]] = new
        for cname in date_cols:
            v = d["values"][col[cname]]
            if v is None or str(v).strip() == "":
                continue
            new, why = normalize_date(v, a.date_order)
            if new is None:
                candidates.append([d["src_row"], cname, str(v), "날짜 " + why, ""])
            elif new != v:
                changes.append({"src_row": d["src_row"], "column": cname, "before": str(v), "after": new, "reason": "날짜 형식 통일(" + why + ")"})
                d["values"][col[cname]] = new
        for cname in email_cols:
            kind, sug = email_issue(d["values"][col[cname]])
            if kind:
                candidates.append([d["src_row"], cname, str(d["values"][col[cname]]), "이메일 " + kind, sug or ""])
    # 3. 범주 표기 사전(사용자가 정한 규칙이므로 바꾸고 기록) → 극단값·범주 흔들림(후보만)
    if cmap:
        for cname in cat_cols:
            rules = {**cmap.get("*", {}), **cmap.get(cname, {})}
            for d in kept:
                v = d["values"][col[cname]]
                if v is not None and str(v) in rules and rules[str(v)] != v:
                    changes.append({"src_row": d["src_row"], "column": cname, "before": v, "after": rules[str(v)],
                                    "reason": f"범주 표기 통일(사용자 사전: {cmap_src})"})
                    d["values"][col[cname]] = rules[str(v)]
    for cname in number_cols:
        vals = [d["values"][col[cname]] if isinstance(d["values"][col[cname]], (int, float)) else None for d in kept]
        for i in outliers(vals, outlier_k):
            candidates.append([kept[i]["src_row"], cname, str(vals[i]), f"극단값(사분위 범위 {outlier_k:g}배 밖) — 입력 오류인지 확인", ""])
    for cname in cat_cols:
        for group in category_variants([d["values"][col[cname]] for d in kept]):
            rows_in = [str(d["src_row"]) for d in kept if str(d["values"][col[cname]]) in group]
            candidates.append(["/".join(rows_in[:10]), cname, " | ".join(group), "같은 값의 다른 표기 의심 — 하나로 통일할지 결정 필요", group[0]])
    # 4. 빈 칸
    blanks = []
    for cname in header:
        if cname in seq:
            continue
        rws = [d["src_row"] for d in kept if d["values"][col[cname]] is None or str(d["values"][col[cname]]).strip() == ""]
        if rws and len(rws) < len(kept):
            blanks.append([cname, len(rws), f"{len(rws) / len(kept):.0%}", ", ".join(map(str, rws[:20])),
                           "필수 정보면 원본 확인 후 입력, 아니면 그대로 두기"])
    # 5. 유사 중복
    near = []
    keys = split(a.near_key)
    if keys:
        groups: dict[tuple, list] = {}
        for d in kept:
            groups.setdefault(tuple(d["values"][col[k]] for k in keys), []).append(d)
        for k, ds in groups.items():
            if len(ds) > 1 and any(x is not None for x in k):
                for d in ds:
                    near.append([f"{' / '.join(map(str, k))}", d["src_row"]] + d["values"])

    # 출력
    owb = Workbook()
    o = owb.active
    o.title = "정리완료"
    write_table(o, header, [d["values"] for d in kept])
    rep = owb.create_sheet("변경리포트")
    summary = [["입력 행 수", len(data)], ["완전 중복 제거", len(removed)], ["정리 후 행 수", len(kept)],
               ["전화번호 형식 통일", sum(1 for c in changes if c["reason"].startswith("전화"))],
               ["날짜 형식 통일", sum(1 for c in changes if c["reason"].startswith("날짜"))],
               ["범주 표기 통일(사용자 사전)", sum(1 for c in changes if c["reason"].startswith("범주"))],
               ["확인 필요 후보(고치지 않음)", len(candidates)], ["유사 중복 후보 행", len(near)],
               ["빈 칸이 있는 열", len(blanks)], ["원본 파일", src.name], ["원본 변경 여부", "변경 없음(새 파일로 저장)"]]
    write_table(rep, ["항목", "값"], summary)
    det = owb.create_sheet("변경내역")
    write_table(det, ["원본 행", "열", "이전 값", "이후 값", "사유"],
                [[c["src_row"], c["column"], str(c["before"]), c["after"], c["reason"]] for c in changes]
                + [[r["src_row"], "(행 전체)", "완전 중복", "삭제", f"{r['duplicate_of']}행과 같음"] for r in removed])
    nd = owb.create_sheet("유사중복_후보")
    write_table(nd, ["키", "원본 행"] + header, near) if near else write_table(nd, ["안내"], [["유사 중복 후보 없음(또는 --near-key 미지정)"]])
    cd = owb.create_sheet("오타_이상값_후보")
    write_table(cd, ["원본 행", "열", "값", "사유", "제안"], candidates) if candidates else write_table(cd, ["안내"], [["후보 없음"]])
    bl = owb.create_sheet("빈칸현황")
    write_table(bl, ["열", "빈 칸 수", "비율", "원본 행(앞 20개)", "제안"], blanks) if blanks else write_table(bl, ["안내"], [["빈 칸 없음"]])
    out.parent.mkdir(parents=True, exist_ok=True)
    owb.save(out)
    after = sha256(src)
    report = {"input": str(src.resolve()), "input_sha256_before": before, "input_sha256_after": after, "sheet": ws.title,
              "header": header, "header_row": h + 1, "sequence_columns": sorted(seq), "rows_in": len(data),
              "rows_out": len(kept), "kept_src_rows": [d["src_row"] for d in kept], "removed": removed,
              "changes": [{**c, "before": str(c["before"])} for c in changes],
              "candidates": [{"src_row": c[0], "column": c[1], "value": c[2], "reason": c[3], "suggestion": c[4]} for c in candidates],
              "blanks": [{"column": b[0], "count": b[1]} for b in blanks], "near_duplicate_rows": len(near),
              "columns_used": {"phone": phone_cols, "date": date_cols, "email": email_cols, "number": number_cols, "category": cat_cols,
                               "near_key": keys},
              "rules": {"date_order": a.date_order, "outlier_k": outlier_k, "keep_duplicates": a.keep_duplicates,
                        "category_map": cmap_src, "settings": st}}
    Path(str(out) + ".report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(f"정리 완료: {out}")
    for k, v in summary[:9]:
        print(f"  {k}: {v}")
    print(f"  사용한 열: {report['columns_used']}")
    if before != after:
        print("경고: 원본 파일 해시가 바뀌었습니다!")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
