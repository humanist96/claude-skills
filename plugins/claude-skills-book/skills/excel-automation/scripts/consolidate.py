#!/usr/bin/env python3
"""멀티탭 취합 엔진 — 원본 탭을 모두 보존하고, 수식으로 연결된 '통합관리' 시트를 만든다.

왜 수식인가: 값을 계산해 붙여 넣으면 원본 탭을 고쳐도 통합 시트가 따라오지 않는다.
INDEX/MATCH 수식으로 원본 탭을 직접 참조하면 원본이 바뀔 때 통합 시트도 바뀐다.
(XLOOKUP·FILTER 같은 최신 함수는 Excel 2016·LibreOffice에서 동작하지 않아 쓰지 않는다)

spec.json 예
{
  "key": "상품코드",
  "sheets": [
    {"sheet": "상품마스터", "columns": {"상품명": "상품명", "카테고리": "카테고리"}},
    {"sheet": "스마트스토어", "columns": {"3월판매량": "스마트스토어_판매량", "3월매출": "스마트스토어_매출"}},
    {"sheet": "쿠팡", "columns": {"3월판매량": "쿠팡_판매량", "3월매출": "쿠팡_매출"}}
  ],
  "sums": {"전체_판매량": ["스마트스토어_판매량", "쿠팡_판매량"]},
  "count_nonblank": {"등록채널수": ["스마트스토어_판매량", "쿠팡_판매량"]},
  "missing_label": ""
}
- key 열은 각 시트에서 이름으로 찾는다(시트마다 위치가 달라도 된다)
- 통합 키 목록 = 지정한 모든 시트의 키 합집합(일부 탭에만 있는 항목도 빠지지 않는다)
- 해당 탭에 없는 항목은 missing_label(spec > 사용자 맞춤 settings.yaml > 기본 빈 칸)로 표시된다

사용법: python consolidate.py <입력.xlsx> --spec spec.json --out <결과.xlsx> [--sheet-name 통합관리] [--no-recalc]
결과 끝에 recalc.py로 재계산해 수식 값을 채운다(엔진이 있으면).
"""
from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname

HEAD_FILL = PatternFill("solid", fgColor="DDE6F5")


def sha256(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def user_settings() -> dict:
    """오버라이드 settings.yaml(missing_label 등). 없으면 빈 dict."""
    try:
        sys.path.insert(0, str(Path(__file__).resolve().parent / "_vendor"))
        import overrides
        return overrides.settings("excel-automation", Path(__file__).resolve().parents[1])["settings"]
    except Exception:
        return {}


def sheetnames(p: Path) -> list[str]:
    wb = load_workbook(p, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()  # 읽기 전용 모드는 파일을 열어 둔다(Windows에서 삭제·이동이 막힘)


def header_map(ws) -> tuple[int, dict[str, int]]:
    """머리글 행 번호와 {열 이름: 열 번호(1부터)}. 앞 10행에서 문자열이 가장 많은 행을 머리글로 본다."""
    best, best_n = 1, -1
    for r in range(1, min(ws.max_row, 10) + 1):
        n = sum(isinstance(ws.cell(r, c).value, str) for c in range(1, ws.max_column + 1))
        if n > best_n:
            best, best_n = r, n
    return best, {str(ws.cell(best, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(best, c).value is not None}


def main(argv: list[str]) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    ap = argparse.ArgumentParser(description="멀티탭 취합 엔진")
    ap.add_argument("input")
    ap.add_argument("--spec", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--sheet-name", default="통합관리")
    ap.add_argument("--no-recalc", action="store_true")
    a = ap.parse_args(argv)
    src, out = Path(a.input), Path(a.out)
    if out.resolve() == src.resolve():
        print("오류: 결과 파일이 원본과 같습니다. 원본은 덮어쓰지 않습니다.")
        return 1
    spec = json.loads(Path(a.spec).read_text(encoding="utf-8"))
    before = sha256(src)
    wb = load_workbook(src)  # 원본 탭을 그대로 품은 채로 작업한다
    key = spec["key"]
    if a.sheet_name in wb.sheetnames:
        print(f"오류: '{a.sheet_name}' 시트가 이미 있습니다. --sheet-name으로 다른 이름을 지정하세요.")
        return 1
    keys: list = []
    seen = set()
    plan = []
    for s in spec["sheets"]:
        if s["sheet"] not in wb.sheetnames:
            print(f"오류: 시트 '{s['sheet']}'가 없습니다. 있는 시트: {wb.sheetnames}")
            return 1
        ws = wb[s["sheet"]]
        hr, hm = header_map(ws)
        if key not in hm:
            print(f"오류: '{s['sheet']}' 시트에 키 열 '{key}'가 없습니다. 열: {list(hm)}")
            return 1
        for src_col in s["columns"]:
            if src_col not in hm:
                print(f"오류: '{s['sheet']}' 시트에 열 '{src_col}'가 없습니다. 열: {list(hm)}")
                return 1
        kc = hm[key]
        count = 0
        for r in range(hr + 1, ws.max_row + 1):
            v = ws.cell(r, kc).value
            if v is None or str(v).strip() == "":
                continue
            count += 1
            if v not in seen:
                seen.add(v)
                keys.append(v)
        plan.append({"sheet": s["sheet"], "header_row": hr, "key_col": kc, "cols": {k: (hm[k], v) for k, v in s["columns"].items()},
                     "rows": count})
    keys.sort(key=lambda x: str(x))
    t = wb.create_sheet(a.sheet_name, 0)
    out_cols = [key] + [v for p in plan for _, (_, v) in p["cols"].items()]
    sums = spec.get("sums", {})
    counts = spec.get("count_nonblank", {})
    out_cols += list(sums) + list(counts)
    for j, name in enumerate(out_cols, 1):
        c = t.cell(1, j, name)
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        t.column_dimensions[get_column_letter(j)].width = max(10, len(name) * 1.8)
    pos = {name: j for j, name in enumerate(out_cols, 1)}
    missing = spec.get("missing_label", user_settings().get("missing_label", ""))
    miss = '""' if missing == "" else f'"{missing}"'
    for i, k in enumerate(keys, 2):
        t.cell(i, 1, k)
        for p in plan:
            sh = quote_sheetname(p["sheet"])
            kcol = get_column_letter(p["key_col"])
            for _, (sc, outname) in p["cols"].items():
                vcol = get_column_letter(sc)
                f = f'=IFERROR(INDEX({sh}!${vcol}:${vcol},MATCH($A{i},{sh}!${kcol}:${kcol},0)),{miss})'
                t.cell(i, pos[outname], f)
        for name, parts in sums.items():
            refs = ",".join(f"{get_column_letter(pos[p])}{i}" for p in parts)
            t.cell(i, pos[name], f"=SUM({refs})")
        for name, parts in counts.items():
            refs = ",".join(f"{get_column_letter(pos[p])}{i}" for p in parts)
            t.cell(i, pos[name], f'=SUMPRODUCT(--(LEN(CHOOSE({{{",".join(str(n) for n in range(1, len(parts) + 1))}}},{refs}))>0))')
    t.freeze_panes = "B2"
    rep = wb.create_sheet("취합리포트", 1)
    rep.append(["항목", "내용"])
    for c in rep[1]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
    rep.append(["통합 기준 열", key])
    rep.append(["통합 항목 수(키 합집합)", len(keys)])
    rep.append(["통합 방식", "INDEX/MATCH 수식으로 원본 탭 참조 — 원본 탭을 고치면 자동 반영"])
    rep.append(["해당 탭에 없는 항목 표시", missing or "(빈 칸)"])
    for p in plan:
        rep.append([f"{p['sheet']} 탭 항목 수", p["rows"]])
    rep.append(["원본 탭", ", ".join(s for s in wb.sheetnames if s not in (a.sheet_name, "취합리포트"))])
    rep.column_dimensions["A"].width = 28
    rep.column_dimensions["B"].width = 60
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    report = {"input": str(src.resolve()), "input_sha256_before": before, "input_sha256_after": sha256(src),
              "sheet": a.sheet_name, "key": key, "keys": len(keys), "spec": spec,
              "plan": [{"sheet": p["sheet"], "rows": p["rows"], "columns": {k: v[1] for k, v in p["cols"].items()}} for p in plan],
              "original_sheets": sheetnames(src)}
    rc = None
    if not a.no_recalc:
        r = subprocess.run([sys.executable, str(Path(__file__).resolve().parent / "recalc.py"), str(out)],
                           capture_output=True, text=True, encoding="utf-8", errors="replace")
        rc = r.stdout.strip().splitlines()[-1] if r.stdout.strip() else r.stderr[-300:]
        report["recalc"] = rc
    Path(str(out) + ".report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"취합 완료: {out} — 항목 {len(keys)}개, 열 {len(out_cols)}개, 원본 탭 {len(report['original_sheets'])}개 보존")
    if rc:
        print(f"재계산: {rc}")
    if report["input_sha256_before"] != report["input_sha256_after"]:
        print("경고: 원본 파일이 바뀌었습니다!")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
