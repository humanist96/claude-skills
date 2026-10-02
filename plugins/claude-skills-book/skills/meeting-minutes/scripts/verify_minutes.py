#!/usr/bin/env python3
"""회의록 품질 게이트 — minutes.json을 번호 붙은 녹취(transcript.json)와 대조한다. 전달 전에 반드시 통과시킨다.

오류(전달 금지)
- schema: meta·decisions·action_items가 없다
- evidence-missing: 결정·할 일·다음 회의에 근거 발언 번호가 없다
- evidence-unknown: 근거 번호가 녹취에 없다
- owner-unsupported: 담당자가 녹취(발언·화자·머리글)에 없는 이름이다("미정"은 허용)
- due-unsupported: 기한의 숫자·요일·상대 표현(내일, 다음 주…)이 근거 발언에 없다("미정"은 허용)
- due-date-mismatch: 기한에 붙인 날짜(예: 수요일(3/11))가 회의일 기준 그 요일이 아니다
- number-provenance: 회의록의 숫자가 녹취에 없다(단위 환산·반올림 허용, 9 이하 정수·근거 번호 제외)
- attendee-unsupported: 참석자 이름이 녹취에 없다
- correction-unsupported: 선언한 STT 교정(meta.corrections)의 들린 말(heard)이 근거 발언에 없다
- promo-text: 결과물과 무관한 스킬 홍보 문구
- rendered-missing: 렌더링한 문서(--rendered)에 결정·할 일이 빠졌다
- leftover: 렌더링한 문서에 {자리표시}·TODO가 남았다
경고
- low-confidence-evidence: STT가 확신하지 못한 발언((?) 표시)을 근거로 썼는데 uncertain에 없다
- stt-low-accuracy: 한국어를 tiny·base 모델로 변환했다(전달 시 정확도 한계를 알릴 것)
- empty-section: 결정 또는 할 일이 0개

STT 오인식 교정(예: '김 피암' → '김PM')은 meta.corrections에 [{"heard", "corrected", "evidence"}]로 선언한다.
교정된 말은 녹취에 있는 것으로 본다.

사용법: python verify_minutes.py <minutes.json> --transcript <transcript.json> [--rendered <md|docx|txt> ...] [--json]
마지막 줄: MINUTES VERIFY OK 또는 MINUTES VERIFY FAILED
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, timedelta
from pathlib import Path

NUM = re.compile(r"(?<![A-Za-z0-9.])[-+]?\d{1,3}(?:,\d{3})+(?:\.\d+)?(?![0-9])|(?<![A-Za-z0-9.,])[-+]?\d+(?:\.\d+)?")
UNIT = re.compile(r"\s?(조|억|천만|백만|만|천)")
UNIT_MULT = {"조": 1e12, "억": 1e8, "천만": 1e7, "백만": 1e6, "만": 1e4, "천": 1e3}
TID = re.compile(r"\bT\d{3,5}\b")
UNDECIDED = re.compile(r"^(미정|미상|tbd|tbc|n/?a|-|없음)\b", re.I)
PROMO = re.compile(r"(이 스킬은|이 스킬로|다음에 녹음 파일|회의록 외에도 활용|강의 노트 정리, 인터뷰)")
LEFTOVER = re.compile(r"(\{[a-z_가-힣 ]+\}|TODO|XXX|\[여기|\(작성\))", re.I)
WEEKDAYS = {"월요일": 0, "화요일": 1, "수요일": 2, "목요일": 3, "금요일": 4, "토요일": 5, "일요일": 6,
            "monday": 0, "tuesday": 1, "wednesday": 2, "thursday": 3, "friday": 4, "saturday": 5, "sunday": 6}
RELATIVE = ["오늘", "내일", "모레", "이번 주", "다음 주", "다음주", "이번주", "월말", "월 말", "주말", "오전", "오후", "today", "tomorrow",
            "next week", "this week", "end of day", "eod", "next monday", "next friday", "morning", "afternoon"]
WEEKDAY_SHORT = {"월": 0, "화": 1, "수": 2, "목": 3, "금": 4, "토": 5, "일": 6}


def strip_tid(s: str) -> str:
    return TID.sub(" ", s)


def numbers(text: str) -> list[tuple[str, float]]:
    t = strip_tid(text)
    out = []
    for m in NUM.finditer(t):
        u = UNIT.match(t, m.end())
        out.append((m.group(0), UNIT_MULT[u.group(1)] if u else 1.0))
    return out


def val(tok: str) -> float | None:
    try:
        return float(tok.replace(",", "").replace("+", ""))
    except ValueError:
        return None


def norm(s: str) -> str:
    return re.sub(r"\s+", "", s or "").lower()


def walk_texts(m: dict) -> list[tuple[str, str]]:
    """(위치, 문장) — 숫자·홍보 문구 검사 대상"""
    out = []
    for i, s in enumerate(m.get("summary", [])):
        out.append((f"summary[{i}]", str(s)))
    for i, a in enumerate(m.get("agenda", [])):
        out.append((f"agenda[{i}].title", a.get("title", "")))
        out += [(f"agenda[{i}].points[{j}]", str(p)) for j, p in enumerate(a.get("points", []))]
    for key in ("decisions", "open_questions", "uncertain", "risks"):
        out += [(f"{key}[{i}]", x.get("text", "") if isinstance(x, dict) else str(x)) for i, x in enumerate(m.get(key, []))]
    for i, sec in enumerate(m.get("sections", [])):
        out += [(f"sections[{i}].items[{j}]", x.get("text", "") if isinstance(x, dict) else str(x)) for j, x in enumerate(sec.get("items", []))]
    for i, a in enumerate(m.get("action_items", [])):
        out += [(f"action_items[{i}].task", a.get("task", "")), (f"action_items[{i}].due", str(a.get("due", "")))]
    nx = m.get("next_meeting")
    if nx:
        out.append(("next_meeting", nx.get("text", "") if isinstance(nx, dict) else str(nx)))
    return out


def rendered_text(p: Path) -> str:
    if p.suffix.lower() == ".docx":
        from docx import Document
        d = Document(str(p))
        parts = [x.text for x in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                parts += [c.text for c in r.cells]
        return "\n".join(parts)
    return p.read_text(encoding="utf-8", errors="replace")


def due_dates(due: str, meeting: date | None) -> list[str]:
    """기한에 붙은 날짜가 요일과 맞는지. 예: '수요일(3/11)' 회의일 2026-03-09 → 3/11은 수요일인가"""
    probs = []
    if not meeting:
        return probs
    wd = next((v for k, v in WEEKDAYS.items() if k in due.lower()), None)
    if wd is None:
        m = re.search(r"\(([월화수목금토일])\)", due)
        wd = WEEKDAY_SHORT.get(m.group(1)) if m else None
    for m in re.finditer(r"(\d{1,2})\s*[/.월]\s*(\d{1,2})", due):
        mo, d = int(m.group(1)), int(m.group(2))
        try:
            dd = date(meeting.year + (1 if mo < meeting.month - 6 else 0), mo, d)
        except ValueError:
            probs.append(f"없는 날짜 {m.group(0)}")
            continue
        if not (meeting - timedelta(days=1) <= dd <= meeting + timedelta(days=120)):
            probs.append(f"{m.group(0)}이 회의일 {meeting} 기준으로 맞지 않는 시점")
        elif wd is not None and dd.weekday() != wd:
            probs.append(f"{m.group(0)}은 {'월화수목금토일'[dd.weekday()]}요일인데 기한 요일과 다름")
    return probs


def verify(m: dict, tr: dict, rendered: list[Path]) -> tuple[list, list]:
    issues, warns = [], []
    for k in ("meta", "decisions", "action_items"):
        if k not in m:
            issues.append({"check": "schema", "detail": f"'{k}' 키가 없다(없으면 빈 목록으로 둔다)"})
    if issues:
        return issues, warns
    meta = m["meta"]
    segs = {s["id"]: s for s in tr.get("segments", [])}
    tmeta = tr.get("meta", {})
    corrections = meta.get("corrections", [])
    corpus = " ".join([s.get("text", "") for s in segs.values()] + [s.get("speaker") or "" for s in segs.values()] + tmeta.get("header", []))
    for c in corrections:
        found = any(norm(c.get("heard", "")) in norm(segs[i]["text"]) for i in c.get("evidence", []) if i in segs)
        if not found:
            issues.append({"check": "correction-unsupported", "where": "meta.corrections",
                           "detail": f"'{c.get('heard')}'가 근거 발언 {c.get('evidence')}에 없다"})
        corpus += " " + c.get("corrected", "")
    ncorpus = norm(corpus)

    def check_ev(where, item, required=True):
        ids = item.get("evidence", []) if isinstance(item, dict) else []
        if not ids:
            if required:
                issues.append({"check": "evidence-missing", "where": where, "detail": "근거 발언 번호가 없다"})
            return []
        bad = [i for i in ids if i not in segs]
        if bad:
            issues.append({"check": "evidence-unknown", "where": where, "detail": f"녹취에 없는 번호 {bad}"})
        good = [segs[i] for i in ids if i in segs]
        low = [s["id"] for s in good if s.get("low_confidence")]
        if low and not any(i in json.dumps(m.get("uncertain", []), ensure_ascii=False) for i in low):
            warns.append({"check": "low-confidence-evidence", "where": where, "detail": f"STT 확신 낮은 발언 {low}을 근거로 씀 — uncertain에 적을 것"})
        return good

    for i, d in enumerate(m["decisions"]):
        check_ev(f"decisions[{i}]", d)
    for i, q in enumerate(m.get("open_questions", [])):
        check_ev(f"open_questions[{i}]", q, required=False)
    meeting = None
    try:
        y, mo, d = (int(x) for x in str(meta.get("date", ""))[:10].split("-"))
        meeting = date(y, mo, d)
    except (ValueError, TypeError):
        pass
    for i, a in enumerate(m["action_items"]):
        where = f"action_items[{i}]"
        good = check_ev(where, a)
        etext = " ".join(s["text"] for s in good) + " " + " ".join(c.get("corrected", "") for c in corrections)
        owner = str(a.get("owner") or "미정").strip()
        if not UNDECIDED.match(owner):
            core = re.sub(r"\s*(님|씨|팀장|과장|대리|부장|사원|매니저|\(.*?\))$", "", owner)
            names = [x.strip() for x in re.split(r"[,/·&]| and ", core) if x.strip()]
            for nm in names:
                variants = {nm, nm[1:] if len(nm) == 3 and re.match(r"[가-힣]{3}$", nm) else nm}  # '이준호' ↔ '준호'
                if not any(norm(v) in ncorpus for v in variants):
                    issues.append({"check": "owner-unsupported", "where": where, "detail": f"담당자 '{nm}'가 녹취에 없다(모르면 '미정')"})
        due = str(a.get("due") or "미정").strip()
        if not UNDECIDED.match(due):
            base_due = re.sub(r"\(.*?\)", "", due)
            tokens = [t for t in re.findall(r"\d+", base_due)]
            words = [w for w in list(WEEKDAYS) + RELATIVE if w in base_due.lower()]
            short = re.findall(r"([월화수목금토일])요?일?까지", base_due)
            if not tokens and not words and not short:
                pass  # "오픈 전"처럼 사건 기준 기한: 숫자·요일이 없으면 근거 발언만 확인
            missing = [t for t in tokens if t not in etext] + [w for w in words if w not in etext.lower() and w.replace(" ", "") not in etext.replace(" ", "").lower()]
            if missing:
                issues.append({"check": "due-unsupported", "where": where, "detail": f"기한 '{due}'의 {missing}가 근거 발언에 없다"})
            for p in due_dates(due, meeting):
                issues.append({"check": "due-date-mismatch", "where": where, "detail": p})
    nx = m.get("next_meeting")
    if nx:
        check_ev("next_meeting", nx if isinstance(nx, dict) else {})
    for att in meta.get("attendees", []):
        core = re.sub(r"\s*\(.*?\)$", "", att).strip()
        if core and norm(core) not in ncorpus and not (len(core) == 3 and norm(core[1:]) in ncorpus):
            issues.append({"check": "attendee-unsupported", "where": "meta.attendees", "detail": f"'{att}'가 녹취에 없다"})
    # 숫자 출처
    tvals = set()
    for tok, mult in numbers(corpus):
        v = val(tok)
        if v is not None:
            tvals.add(round(v, 4))
            tvals.add(round(v * mult, 4))
    for where, text in walk_texts(m):
        if where.endswith(".due") and meeting:
            # 요일 옆에 붙인 계산 날짜(수요일(3/11))는 due-date-mismatch가 회의일 기준으로 따로 검사한다
            text = re.sub(r"\(\s*\d{1,2}\s*[/.월]\s*\d{1,2}\s*일?\s*(?:\([월화수목금토일]\))?\s*\)", " ", text)
        if PROMO.search(text):
            issues.append({"check": "promo-text", "where": where, "detail": text[:60]})
        for tok, mult in numbers(text):
            v = val(tok)
            if v is None or (float(v).is_integer() and abs(v) <= 9 and "." not in tok):
                continue
            if 1990 <= v <= 2100 and "," not in tok:
                continue
            if not ({round(v, 4), round(v * mult, 4)} & tvals) and not any(abs(v - x) < 0.051 for x in tvals if abs(x) < 1000):
                issues.append({"check": "number-provenance", "where": where, "detail": f"'{tok}'가 녹취에 없다: {text[:70]}"})
    if not m["decisions"]:
        warns.append({"check": "empty-section", "detail": "결정 사항 0개 — 정말 없었다면 그대로 두고 문서에 '없음'으로 표시"})
    if not m["action_items"]:
        warns.append({"check": "empty-section", "detail": "할 일 0개"})
    if tmeta.get("kind") == "audio" and str(tmeta.get("language", "")).startswith("ko") and tmeta.get("model") in ("tiny", "base"):
        warns.append({"check": "stt-low-accuracy", "detail": f"한국어를 {tmeta.get('model')} 모델로 변환 — 고유명사·숫자 오인식 가능성을 전달 시 알릴 것"})
    for p in rendered:
        if not p.is_file():
            issues.append({"check": "rendered-missing", "where": str(p), "detail": "파일 없음"})
            continue
        rt = norm(rendered_text(p))
        raw = rendered_text(p)
        for i, d in enumerate(m["decisions"]):
            if norm(d.get("text", ""))[:40] not in rt:
                issues.append({"check": "rendered-missing", "where": f"{p.name}", "detail": f"결정 {i + 1} 누락: {d.get('text', '')[:40]}"})
        for i, a in enumerate(m["action_items"]):
            if norm(a.get("task", ""))[:40] not in rt:
                issues.append({"check": "rendered-missing", "where": f"{p.name}", "detail": f"할 일 {i + 1} 누락: {a.get('task', '')[:40]}"})
        for mm in LEFTOVER.finditer(raw):
            issues.append({"check": "leftover", "where": p.name, "detail": mm.group(0)})
        if PROMO.search(raw):
            issues.append({"check": "promo-text", "where": p.name, "detail": PROMO.search(raw).group(0)})
    return issues, warns


def main(argv: list[str]) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass
    ap = argparse.ArgumentParser(description="회의록 품질 게이트")
    ap.add_argument("minutes")
    ap.add_argument("--transcript", required=True)
    ap.add_argument("--rendered", nargs="*", default=[])
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args(argv)
    try:
        m = json.loads(Path(a.minutes).read_text(encoding="utf-8"))
        tr = json.loads(Path(a.transcript).read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as e:
        print(f"오류: 파일을 읽지 못했다: {e}")
        print("MINUTES VERIFY FAILED")
        return 1
    issues, warns = verify(m, tr, [Path(p) for p in a.rendered])
    if a.json:
        print(json.dumps({"errors": issues, "warnings": warns}, ensure_ascii=False, indent=2))
    else:
        for i in issues:
            print(f"ERROR [{i['check']}] {i.get('where', '')} {i['detail']}")
        for w in warns:
            print(f"WARN  [{w['check']}] {w.get('where', '')} {w['detail']}")
    n_ev = len(m.get("decisions", [])) + len(m.get("action_items", []))
    print(f"MINUTES VERIFY {'OK' if not issues else 'FAILED'} items={n_ev} errors={len(issues)} warnings={len(warns)}")
    return 0 if not issues else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
